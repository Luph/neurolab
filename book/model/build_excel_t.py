"""Build model/Case_T_Model.xlsx: the Case T workbook with live formulas (FAST-consistent).

Reads inputs_case_t.json (assumptions) and outputs_case_t.json (only for the locked-financing
block: the banking-case debt profiles that every non-sizing scenario uses, exactly as a modeler
pastes a sized debt profile before running sensitivities; the Checks sheet proves that the live
sizing reproduces them when Scenario = 2).

Layout: columns A-C blank, D label, E unit, F constant, G row total, H-I blank, J = first model
period (2015-05-27 to 2015-06-30) to DB = 2059H1. Inputs blue on pale yellow; links from other
sheets green; calculations black. Circularity: iterative calculation on (1,000 iterations,
maximum change 0.0000001) with the circuit breaker Inputs!F(Circ): 0 = financing feedback cut
(locked financing used), 1 = live.
"""
import json
import os
import re
import datetime as dt

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties
from openpyxl.formatting.rule import CellIsRule
from openpyxl.workbook.defined_name import DefinedName

HERE = os.path.dirname(os.path.abspath(__file__))
INP = json.load(open(os.path.join(HERE, "inputs_case_t.json")))
OUT = json.load(open(os.path.join(HERE, "outputs_case_t.json")))
T = 97
FIRST_COL = 10  # J
LAST_COL = FIRST_COL + T - 1
LASTL = get_column_letter(LAST_COL)

BLUE = Font(color="0000FF")
GREEN = Font(color="008000")
BOLD = Font(bold=True)
TITLE = Font(bold=True, size=14)
INFILL = PatternFill("solid", fgColor="FFF2CC")
HDRFILL = PatternFill("solid", fgColor="D9E1F2")
REDFILL = PatternFill("solid", fgColor="FF9999")

FMT = {"ARD m": '#,##0.0;(#,##0.0);"–"', "x": '0.00"x"', "%": '0.00%', "date": 'yyyy-mm-dd', "flag": '0',
       "days": '0', "k/day": '0.0', "factor": '0.0000', "ARD/km": '0.0000', "year": '0', "#": '0',
       "%pa": '0.00', "ARD m4": '#,##0.0000;(#,##0.0000);"–"', "code": '0', "text": '@', "km": '0.0'}

SHEETS = ["Cover", "Inputs", "Time", "Construction", "Operations", "Tax", "Funding", "Debt", "Reserves",
          "Waterfall", "Financials", "Ratios", "Returns", "Checks", "Outputs"]


class Row:
    def __init__(self, sheet, name, label, unit, kind, f=None, value=None, total=False, ncols=None,
                 values=None, fmt=None, note=None):
        self.sheet, self.name, self.label, self.unit, self.kind = sheet, name, label, unit, kind
        self.f, self.value, self.total, self.ncols, self.values = f, value, total, ncols, values
        self.fmt, self.note = fmt, note
        self.r = None


class Builder:
    def __init__(self):
        self.rows = {}
        self.order = {s: [] for s in SHEETS}
        self.next = {s: 5 for s in SHEETS}

    def _add(self, row):
        assert row.name not in self.rows, row.name
        row.r = self.next[row.sheet]
        self.next[row.sheet] += 1
        self.rows[row.name] = row
        self.order[row.sheet].append(row)
        return row

    def head(self, sheet, text):
        r = Row(sheet, "_h%d_%s" % (len(self.rows), sheet), text, "", "head")
        self._add(r)
        self.next[sheet] += 0

    def gap(self, sheet):
        self.next[sheet] += 1

    def inp(self, sheet, name, label, unit, value, note=None, fmt=None):
        """Scalar input (column F, blue on yellow)."""
        return self._add(Row(sheet, name, label, unit, "inp", value=value, note=note, fmt=fmt))

    def sc(self, sheet, name, label, unit, f, note=None, fmt=None):
        """Scalar calculation in column F."""
        return self._add(Row(sheet, name, label, unit, "sc", f=f, note=note, fmt=fmt))

    def tr(self, sheet, name, label, unit, f, total=False, note=None, fmt=None):
        """Timeline row: one formula copied across J..DB."""
        return self._add(Row(sheet, name, label, unit, "tr", f=f, total=total, note=note, fmt=fmt))

    def tin(self, sheet, name, label, unit, values, total=False, note=None, fmt=None):
        """Timeline-aligned input row (values for J..)."""
        return self._add(Row(sheet, name, label, unit, "tin", values=values, total=total, note=note, fmt=fmt))

    def table(self, sheet, name, label, unit, values, note=None, fmt=None):
        """Input table row starting in column J (scenario, case, year or code tables)."""
        return self._add(Row(sheet, name, label, unit, "table", values=values, ncols=len(values), note=note, fmt=fmt))

    # ---------------------------------------------------------------- formula resolution
    def ref(self, name, mod, col, cur_sheet):
        if name == "col":
            return get_column_letter(col)
        row = self.rows[name]
        pre = "" if row.sheet == cur_sheet else row.sheet + "!"
        r = row.r
        if mod in ("c", None) and row.kind in ("inp", "sc"):
            return f"{pre}$F${r}"
        if mod == "r":
            if row.kind == "table":
                return f"{pre}$J${r}:${get_column_letter(FIRST_COL + row.ncols - 1)}${r}"
            return f"{pre}$J${r}:${LASTL}${r}"
        if mod == "rt":
            return f"{pre}{get_column_letter(col)}${r}:${LASTL}${r}"
        if mod == "n6":
            return f"{pre}{get_column_letter(col + 1)}{r}:{get_column_letter(col + 6)}{r}"
        if mod == "p":
            return f"{pre}{get_column_letter(col - 1)}{r}"
        if mod == "n":
            return f"{pre}{get_column_letter(col + 1)}{r}"
        if mod is None:
            return f"{pre}{get_column_letter(col)}{r}"
        raise ValueError((name, mod))

    def resolve(self, f, col, sheet):
        def sub(m):
            return self.ref(m.group(1), m.group(2), col, sheet)
        return re.sub(r"\{([A-Za-z0-9_]+)(?:\.([a-z0-9]+))?\}", sub, f)

    # ---------------------------------------------------------------- writing
    def write(self, path, scenario=2, contribution_option=1):
        wb = Workbook()
        wb.remove(wb.active)
        ws_by = {s: wb.create_sheet(s) for s in SHEETS}
        for s in SHEETS:
            ws = ws_by[s]
            ws.column_dimensions["A"].width = 2
            ws.column_dimensions["B"].width = 2
            ws.column_dimensions["C"].width = 2
            ws.column_dimensions["D"].width = 58
            ws.column_dimensions["E"].width = 10
            ws.column_dimensions["F"].width = 13
            ws.column_dimensions["G"].width = 12
            ws.column_dimensions["H"].width = 2
            ws.column_dimensions["I"].width = 2
            for c in range(FIRST_COL, LAST_COL + 1):
                ws.column_dimensions[get_column_letter(c)].width = 11
            ws["D1"] = f"Case T: Merrick Link toll road PPP  |  {s}"
            ws["D1"].font = TITLE
            if s not in ("Cover", "Outputs", "Checks"):
                ws["D3"] = "Period end"
                ws["E3"] = "date"
                ws["D2"] = "Units: ARD millions unless stated. Inputs blue on yellow; links green; calculations black."
                for c in range(FIRST_COL, LAST_COL + 1):
                    cell = ws.cell(3, c)
                    if s == "Time":
                        cell.value = f"={get_column_letter(c)}{self.rows['End'].r}"
                    else:
                        cell.value = f"=Time!{get_column_letter(c)}{self.rows['End'].r}"
                        cell.font = GREEN
                    cell.number_format = FMT["date"]
                    cell.fill = HDRFILL
                    cell.font = BOLD
                ws.freeze_panes = "J4"
        for s in SHEETS:
            ws = ws_by[s]
            for row in self.order[s]:
                r = row.r
                if row.kind == "head":
                    ws.cell(r, 4, row.label).font = BOLD
                    continue
                ws.cell(r, 4, row.label)
                ws.cell(r, 5, row.unit)
                fmt = row.fmt or FMT.get(row.unit, '#,##0.00')
                if row.note:
                    ws.cell(r, 7 if row.kind in ("inp", "sc") else 7).value = None
                if row.kind == "inp":
                    c = ws.cell(r, 6, row.value)
                    c.font, c.fill, c.number_format = BLUE, INFILL, fmt
                elif row.kind == "sc":
                    c = ws.cell(r, 6, "=" + self.resolve(row.f, 6, s))
                    c.number_format = fmt
                    if re.fullmatch(r"[A-Za-z]+![$A-Z]+\$?[0-9]+", self.resolve(row.f, 6, s)):
                        c.font = GREEN
                elif row.kind in ("tr",):
                    for k in range(T):
                        col = FIRST_COL + k
                        f = self.resolve(row.f, col, s)
                        c = ws.cell(r, col, "=" + f)
                        c.number_format = fmt
                        if re.fullmatch(r"[A-Za-z]+![A-Z]+[0-9]+", f):
                            c.font = GREEN
                    if row.total:
                        ws.cell(r, 7, f"=SUM(J{r}:{LASTL}{r})").number_format = fmt
                elif row.kind in ("tin", "table"):
                    for k, v in enumerate(row.values):
                        c = ws.cell(r, FIRST_COL + k, v)
                        c.font, c.fill, c.number_format = BLUE, INFILL, fmt
                    if row.total:
                        ws.cell(r, 7, f"=SUM(J{r}:{LASTL}{r})").number_format = fmt
                if row.note:
                    ws.cell(r, 3 if False else 4).comment = None
        # scenario selectors
        ws_by["Inputs"].cell(self.rows["Scenario"].r, 6).value = scenario
        ws_by["Inputs"].cell(self.rows["ContribOpt"].r, 6).value = contribution_option
        wb.defined_names["Scenario"] = DefinedName("Scenario", attr_text=f"Inputs!$F${self.rows['Scenario'].r}")
        # conditional formatting on checks
        ws = ws_by["Checks"]
        for row in self.order["Checks"]:
            if row.kind == "sc":
                ws.conditional_formatting.add(f"F{row.r}", CellIsRule(operator="notEqual", formula=["0"], fill=REDFILL))
        self.cover(ws_by["Cover"], scenario)
        wb.calculation = CalcProperties(iterate=True, iterateCount=1000, iterateDelta=1e-7, fullCalcOnLoad=True)
        wb.save(path)

    def cover(self, ws, scenario):
        lines = [
            "Case T reference model: the Merrick Link toll road PPP, State of Brannock, Commonwealth of Ardmore.",
            "Model version 1.0; story as of October 3, 2026. Currency: Ardmore dollar (ARD), millions, nominal.",
            "Python mirror: model/case_t.py (source of truth); verification: model/case_t_verification.md.",
            "",
            "How to use: choose the scenario in Inputs!F (named range Scenario):",
        ]
        for k, v in SCENARIO_NAMES.items():
            lines.append(f"   {k} = {v}")
        lines += [
            "",
            "Contribution option (Inputs): 1 = bid ARD 287.4 million; 2 = contribution solved by the Python mirror for an",
            "   11.4% equity IRR on the banking case (run with Scenario 2; Returns shows the IRR it delivers).",
            "",
            "Circularity: financing costs during construction, the funding requirement, debt sizing, tax and the sculpting",
            "   divisors are circular. The file is saved with iterative calculation on (File > Options > Formulas >",
            "   Enable iterative calculation, 1,000 iterations, maximum change 0.0000001). Circuit breaker: Inputs Circ.",
            "   Set Circ to 0 if errors ever propagate (the locked financing is then used and the loop is cut), recalculate,",
            "   then set it back to 1.",
            "",
            "Sizing: Scenario 2 sizes the debt live. Every other scenario uses the locked banking-case financing on the",
            "   Inputs sheet (as a lender's model does once the debt is sized). Checks shows live = locked in Scenario 2.",
            "   Scenario 5 sculpts the Restructured Senior Notes and the restructured NILO loan live; Scenario 4 uses",
            "   those profiles locked.",
            "",
            "Sheets: Inputs, Time, Construction, Operations, Tax, Funding, Debt, Reserves, Waterfall, Financials, Ratios,",
            "   Returns (equity, project, PSC and value for money, termination, restructuring), Checks, Outputs.",
            "Conventions: time across columns from column J; one formula per row; flags are 1/0; costs positive and",
            "   subtracted; tables show negatives in parentheses.",
        ]
        for i, l in enumerate(lines):
            ws.cell(3 + i, 4, l)
        ws.column_dimensions["D"].width = 120


SCENARIO_NAMES = {
    1: "Bid base (Pellow traffic), locked financing",
    2: "Banking case (Ridgeway traffic), live debt sizing",
    3: "Downside traffic, locked financing",
    4: "Actual history to 2026H1, then Ridgeway 2023 case; restructuring at 2023-12-31",
    5: "Restructuring case: actual history to 2023H2, Ridgeway 2023 case from 2024 (live notes sculpting)",
    6: "Retender valuation at 2022-06-30 (actual to 2022H1, retender traffic case after)",
    7: "Sensitivity: traffic -10%",
    8: "Sensitivity: ramp-up one year slower",
    9: "Sensitivity: toll escalation CPI only",
    10: "Sensitivity: opex +10%",
    11: "Sensitivity: lifecycle +20%",
    12: "Sensitivity: interest +100 bp on refinancing",
    13: "Sensitivity: heavy-vehicle share -2 points",
}


def xdate(s):
    return dt.date.fromisoformat(s)


def build(path, scenario=2, contribution_option=1):
    B = Builder()
    I = "Inputs"
    fin = INP["financing_at_close"]
    cap = INP["capex_ard_m"]
    ev = INP["events"]["restructuring"]
    MA = INP["modeler_assumptions"]
    tolls = INP["tolls"]
    traf = INP["traffic_k_trips_per_day"]
    ops = INP["opex_ard_m_2015_prices"]
    D = dt.date

    # ===================================================================== INPUTS
    B.head(I, "Switches")
    B.inp(I, "Scenario", "Scenario (1-13, see Cover)", "#", scenario)
    B.inp(I, "ContribOpt", "Contribution option (1 = bid, 2 = solved for 11.4% on banking)", "#", contribution_option)
    B.inp(I, "Circ", "Circuit breaker (1 = live circular links, 0 = cut)", "flag", 1)
    B.inp(I, "Tol", "Test tolerance", "ARD m", 0.000001, fmt="0.000000")
    B.inp(I, "Big", "Large number for MIN helper rows", "#", 1e12, fmt="0.0E+00")
    B.gap(I)
    B.head(I, "Dates")
    B.inp(I, "FC_Date", "Financial close", "date", D(2015, 5, 27))
    B.inp(I, "FirstEnd", "End of first model period", "date", D(2015, 6, 30))
    B.inp(I, "Con_End", "Scheduled completion and end of construction funding", "date", D(2019, 3, 31))
    B.inp(I, "Open_Sched", "Scheduled opening", "date", D(2019, 3, 31))
    B.inp(I, "Open_Act", "Actual opening", "date", D(2019, 5, 6))
    B.inp(I, "Conc_Orig", "Concession expiry (original)", "date", D(2053, 5, 26))
    B.inp(I, "Conc_Ext", "Concession expiry (extended 2023)", "date", D(2059, 5, 26))
    B.inp(I, "Restr_Date", "Restructuring (model date)", "date", D(2023, 12, 31))
    B.inp(I, "First_Rep", "Senior first repayment", "date", D(2021, 6, 30))
    B.inp(I, "Final_Rep", "Senior final repayment", "date", D(2048, 12, 31))
    B.inp(I, "Margin_Step", "Bank margin step-up date", "date", D(2020, 5, 27))
    B.inp(I, "MPM", "Bank mini-perm maturity (refinancing or A&E date)", "date", D(2022, 5, 27))
    B.inp(I, "Hedge_End", "Swap end", "date", D(2035, 6, 30))
    B.inp(I, "Nilo_CapEnd", "NILO interest capitalized to", "date", D(2024, 3, 31))
    B.inp(I, "Nilo_First", "NILO first repayment (original)", "date", D(2024, 6, 30))
    B.inp(I, "Nilo_Final", "NILO final repayment (original)", "date", D(2052, 12, 31))
    B.inp(I, "NiloR_First", "NILO first repayment (restructured)", "date", D(2031, 6, 30))
    B.inp(I, "NiloR_Final", "NILO final repayment (restructured)", "date", D(2058, 12, 31))
    B.inp(I, "PIK_End", "NILO PIK period end (restructured)", "date", D(2030, 12, 31))
    B.inp(I, "Notes_Final", "Restructured notes final repayment", "date", D(2052, 12, 31))
    B.inp(I, "Sweep_End", "Notes 50% cash sweep end", "date", D(2030, 12, 31))
    B.inp(I, "First_Test", "First DSCR test date", "date", D(2019, 12, 31))
    B.inp(I, "Default_Test", "First default DSCR test date", "date", D(2020, 12, 31))
    B.inp(I, "FV_Date", "Termination comparison and fair value date", "date", D(2022, 6, 30))
    B.inp(I, "Ins2_Date", "Construction insurance second instalment", "date", D(2017, 6, 30))
    B.inp(I, "RC_Start", "Restructuring costs: first period end", "date", D(2022, 6, 30))
    B.inp(I, "Upg_First", "Interchange upgrade: first period end", "date", D(2024, 6, 30))
    B.inp(I, "Upg_Last", "Interchange upgrade: last period end", "date", D(2025, 12, 31))
    B.inp(I, "PSC_Date", "PSC present value date", "date", D(2012, 12, 31))
    B.inp(I, "Contrib_Date", "PSC: contribution payment date", "date", D(2019, 3, 31))
    B.gap(I)
    B.head(I, "Conventions")
    B.inp(I, "DaysYear", "Days per year (ACT/365)", "days", 365)
    B.inp(I, "Days360Y", "Days per year (30/360)", "days", 360)
    B.inp(I, "DaysMonth", "Average days per month (period length in months)", "days", 30.4, fmt="0.0")
    B.inp(I, "RetMonths", "Defects period", "months", 24, fmt="0")
    B.inp(I, "HB_Months", "Handback funding window", "months", 60, fmt="0")
    B.inp(I, "Last6Months", "Bridge draws: last construction months", "months", 18, fmt="0")
    B.inp(I, "EscMonth", "Toll escalation month (July 1)", "#", 7)
    B.inp(I, "Toll_dp", "Toll rounding (decimal places)", "#", 4)
    B.inp(I, "Esc_Floor_Last", "Last year of the 3% escalation floor (July 1)", "year", 2029)
    B.inp(I, "Restr_Esc_From", "CPI-only escalation from (restructured)", "year", 2024)
    B.inp(I, "Band1_End", "Traffic growth band 1 ends", "year", 2030)
    B.inp(I, "Band2_End", "Traffic growth band 2 ends", "year", 2040)
    B.inp(I, "Ramp_First", "First ramp-up year", "year", 2019)
    B.inp(I, "Ramp_Last", "Last ramp-up year in table", "year", 2023)
    B.inp(I, "Uplift_From", "Heavy-vehicle uplift from (restructured case)", "year", 2024)
    B.inp(I, "OY_Base", "Operating year base (OY = year - base)", "year", 2018)
    B.inp(I, "CPI_From", "CPI product for revenue-share threshold: from year", "year", 2016)
    B.inp(I, "CPI_To", "CPI product for revenue-share threshold: to year (2023 prices)", "year", 2023)
    B.inp(I, "HalfCode", "Half-year code multiplier (year x 10 + half)", "#", 10)
    B.inp(I, "KTrips", "Thousand trips", "#", 1000)
    B.inp(I, "Million", "ARD per ARD million", "#", 1000000)
    B.gap(I)
    B.head(I, "Construction costs")
    B.inp(I, "DC_Price", "D&C contract price (lump sum)", "ARD m", cap["dc_contract_price"])
    B.inp(I, "Ret1", "Retention released at completion", "%", cap["dc_retention_release"]["at_completion_pct"] / 100)
    B.inp(I, "Ret2", "Retention released at end of defects period", "%", cap["dc_retention_release"]["end_of_defects_period_24_months_pct"] / 100)
    B.inp(I, "Dev", "Development and bid costs reimbursed (at close)", "ARD m", cap["development_and_bid_costs_reimbursed"])
    B.inp(I, "SPVC", "Project company costs during construction (flat by day)", "ARD m", cap["spv_costs_during_construction"])
    B.inp(I, "InsC", "Insurance during construction", "ARD m", cap["insurance_during_construction"])
    B.inp(I, "InsC1", "Insurance share paid at close", "%", 0.8)
    B.inp(I, "Cert", "Independent certifier (flat by day)", "ARD m", cap["independent_certifier_spv_share"])
    B.inp(I, "Adv", "Lenders' advisors and legal (flat by day)", "ARD m", cap["lenders_advisors_and_legal"])
    B.inp(I, "Cont", "Contingency (pro rata with D&C)", "ARD m", cap["spv_contingency"])
    B.inp(I, "LD_Act", "Delay LDs received (actual)", "ARD m", cap["dc_terms"]["actual_delay_lds_ard_m"])
    B.gap(I)
    B.head(I, "State contribution and bridge")
    B.inp(I, "Contrib_Bid", "State construction contribution (bid)", "ARD m", INP["state_contribution"]["amount_ard_m"])
    B.inp(I, "Contrib_Solved", "Contribution for 11.4% on banking (solved by Python mirror)", "ARD m",
          round(OUT["contribution_solved"]["banking_case_at_11.4pct"], 6))
    B.inp(I, "Bridge_Margin", "Bridge margin over ABBR", "%", 0.016)
    B.inp(I, "Bridge_Fee", "Bridge upfront fee", "%", 0.01)
    B.gap(I)
    B.head(I, "Equity and sizing")
    B.inp(I, "EqShare", "Equity share of funding requirement net of contribution", "%", 0.22)
    B.inp(I, "ShareCap", "Share capital share of equity", "%", 0.15)
    B.inp(I, "SHL_Rate", "Shareholder loan rate", "%", 0.1025)
    B.inp(I, "Target_IRR", "Bid target equity IRR (nominal post-tax)", "%", 0.114)
    B.inp(I, "DSCR_T", "Senior sculpting / minimum DSCR, banking", "x", 1.50)
    B.inp(I, "DSCR_Down", "Minimum DSCR, downside", "x", 1.15)
    B.inp(I, "LLCR_T", "Minimum LLCR, banking", "x", 1.55)
    B.inp(I, "GearCap", "Senior cap, % of funding net of contribution", "%", 0.55)
    B.inp(I, "RIA_DSCR", "Ramp-up interest account target DSCR", "x", MA["ramp_up_interest_account"]["target_dscr_x"])
    B.inp(I, "Bank_Split", "Senior split: bank", "%", 0.45)
    B.inp(I, "Bond_Split", "Senior split: bonds", "%", 0.55)
    B.gap(I)
    B.head(I, "Senior debt terms")
    B.inp(I, "Swap", "Bank swap fixed rate", "%pa", 3.48)
    B.inp(I, "SwapMkt", "Market swap rate at 2023-12-31", "%pa", 4.36)
    B.inp(I, "M1", "Bank margin to 2020-05-27", "%pa", 2.35)
    B.inp(I, "M2", "Bank margin 2020-05-27 to 2022-05-27", "%pa", 2.60)
    B.inp(I, "M3_Refi", "Base-case refinancing margin from 2022-05-27", "%pa", 2.25)
    B.inp(I, "M3_AE", "Amend-and-extend margin from 2022-05-27 (actual)", "%pa", 3.25)
    B.inp(I, "Refi_Fee", "Base-case refinancing fee", "%", 0.0125)
    B.inp(I, "AE_Fee", "Amend-and-extend fee", "%", 0.005)
    B.inp(I, "Bank_Up", "Bank upfront fee", "%", 0.0185)
    B.inp(I, "Bank_Commit", "Bank commitment fee", "%", 0.0095)
    B.inp(I, "BondCpn", "Bond coupon", "%", 0.0485)
    B.inp(I, "Escrow_Rate", "Bond escrow earnings rate", "%", 0.021)
    B.inp(I, "BondCost", "Bond issue costs", "%", 0.012)
    B.inp(I, "LockUp", "Lock-up DSCR (historic)", "x", 1.20)
    B.inp(I, "DefDSCR", "Event of default DSCR (historic)", "x", 1.05)
    B.inp(I, "DSRA_Months", "DSRA months of debt service", "months", 6, fmt="0")
    B.gap(I)
    B.head(I, "NILO loan")
    B.inp(I, "Nilo_Rate", "NILO fixed rate", "%", 0.0306)
    B.inp(I, "Nilo_PIK", "NILO PIK rate to 2030 (restructured)", "%", 0.01)
    B.inp(I, "Nilo_Max", "NILO maximum share of eligible costs", "%", 0.33)
    B.inp(I, "Nilo_Fee", "NILO application fee", "%", 0.001)
    B.inp(I, "Nilo_Comb", "Combined senior plus NILO minimum DSCR (banking)", "x", 1.25)
    B.gap(I)
    B.head(I, "Tolls, traffic and revenue")
    B.inp(I, "Toll2014", "Car toll per km, June 2014 prices", "ARD/km", tolls["car_toll_ard_per_km_2014"])
    B.inp(I, "Esc_Floor", "Escalation floor to 2030", "%", 0.03)
    B.inp(I, "TripKm", "Average trip length", "km", tolls["average_trip_length_km"])
    B.inp(I, "Leak", "Revenue leakage", "%", tolls["revenue_leakage_pct"] / 100)
    B.inp(I, "Mix_Car", "Vehicle mix: car (original)", "%pa", tolls["vehicle_mix_pct_original"]["car"])
    B.inp(I, "Mix_LCV", "Vehicle mix: light commercial (original)", "%pa", tolls["vehicle_mix_pct_original"]["light_commercial"])
    B.inp(I, "Mix_HV", "Vehicle mix: heavy (original)", "%pa", tolls["vehicle_mix_pct_original"]["heavy_vehicle"])
    B.inp(I, "MixN_Car", "Vehicle mix: car (from 2024)", "%pa", tolls["vehicle_mix_pct_from_2024"]["car"])
    B.inp(I, "MixN_LCV", "Vehicle mix: light commercial (from 2024)", "%pa", tolls["vehicle_mix_pct_from_2024"]["light_commercial"])
    B.inp(I, "MixN_HV", "Vehicle mix: heavy (from 2024)", "%pa", tolls["vehicle_mix_pct_from_2024"]["heavy_vehicle"])
    B.inp(I, "Mult_LCV", "Class multiplier: light commercial", "factor", 1.6)
    B.inp(I, "Mult_HV", "Class multiplier: heavy (original)", "factor", 2.85)
    B.inp(I, "MultN_HV", "Class multiplier: heavy (from 2024)", "factor", 2.40)
    B.inp(I, "Mult_Car", "Class multiplier: car", "factor", 1.0)
    B.gap(I)
    B.head(I, "Operating costs (2015 prices) and lifecycle")
    B.inp(I, "OM", "Corvus fixed O&M", "ARD m", ops["om_fixed_corvus_pa"])
    B.inp(I, "OM_Cut", "O&M fee cut from 2024 (restructured)", "%", 0.08)
    B.inp(I, "BO_Pct", "Tolling back office, % of toll revenue", "%", 0.062)
    B.inp(I, "BO_Fix", "Tolling back office fixed", "ARD m", ops["tolling_back_office_fixed_pa"])
    B.inp(I, "Ins", "Insurance", "ARD m", ops["insurance_pa"])
    B.inp(I, "SPV", "Project company costs", "ARD m", ops["spv_costs_pa"])
    B.inp(I, "LC_Pav", "Pavement resurfacing (OY12, 24, 36)", "ARD m", 42.6)
    B.inp(I, "Cyc_Pav", "Pavement cycle", "years", 12, fmt="0")
    B.inp(I, "LC_Tun", "Tunnel M&E (OY15, 30)", "ARD m", 31.9)
    B.inp(I, "Cyc_Tun", "Tunnel cycle", "years", 15, fmt="0")
    B.inp(I, "LC_ITS", "Tolling and ITS replacement (every 8 years from OY8)", "ARD m", 14.7)
    B.inp(I, "Cyc_ITS", "Tolling and ITS cycle", "years", 8, fmt="0")
    B.inp(I, "MMRA_N", "Lifecycle reserve accumulation periods", "#", 6)
    B.inp(I, "HB_Est", "Handback works estimate", "ARD m", MA["handback_works_estimate_ard_m_2015"])
    B.inp(I, "HB_N", "Handback funding periods", "#", MA["handback_funding_periods"])
    B.inp(I, "RecDays", "Toll receivable days", "days", 14)
    B.inp(I, "PayDays", "Payable days", "days", 30)
    B.inp(I, "TaxRate", "Corporate tax rate", "%", 0.30)
    B.gap(I)
    B.head(I, "Events and restructuring")
    B.inp(I, "RC_Total", "Restructuring costs 2022-2023", "ARD m", ops["restructuring_costs_2022_2023"])
    B.inp(I, "RC_N", "Restructuring cost periods", "#", 4)
    B.inp(I, "NotesShare", "Restructured notes, % of net claims", "%", 0.76)
    B.inp(I, "Cancel", "Claims cancelled", "%", 0.14)
    B.inp(I, "Convert", "Claims converted to equity", "%", 0.10)
    B.inp(I, "NotesCpn", "Restructured notes coupon", "%", 0.051)
    B.inp(I, "NotesDSCR", "Notes minimum DSCR (sculpting floor)", "x", 1.30)
    B.inp(I, "NotesSweep", "Notes excess cash sweep to 2030", "%", 0.5)
    B.inp(I, "CredEq", "Senior creditors' share of new equity", "%", 0.85)
    B.inp(I, "StateEq", "State share of new equity", "%", 0.15)
    B.inp(I, "StateMoney", "State new money", "ARD m", ev["state_new_money_ard_m"])
    B.inp(I, "Upgrade", "Holloway Junction interchange upgrade", "ARD m", ev["state_new_money_use"]["holloway_junction_interchange_upgrade_2024_2025"])
    B.inp(I, "Upg_N", "Upgrade spending periods", "#", 4)
    B.inp(I, "StateRes", "Reserve top-up (to cash at restructuring)", "ARD m", ev["state_new_money_use"]["reserve_top_up"])
    B.inp(I, "Warrant", "Warrants over new equity (original sponsors)", "%", 0.03)
    B.inp(I, "RevShare", "State share of toll revenue above threshold", "%", 0.30)
    B.inp(I, "RevThresh", "Revenue share threshold (2023 prices, a year)", "ARD m", 260.0)
    B.inp(I, "EqV_Rate", "Plan valuation: new equity discount rate", "%", MA["plan_valuation"]["equity_discount_rate_pct"] / 100)
    B.inp(I, "Notes_Yield", "Plan valuation: notes market yield", "%", MA["plan_valuation"]["notes_market_yield_pct"] / 100)
    B.inp(I, "FV_Rate", "Fair value 2022: pre-tax discount rate", "%", MA["fair_value_2022"]["pre_tax_discount_rate_pct"] / 100)
    B.inp(I, "Retender_Cost", "Retendering costs", "ARD m", MA["fair_value_2022"]["retendering_costs_ard_m"])
    B.gap(I)
    B.head(I, "Public sector comparator (PV at 2012-12-31)")
    psc = INP["psc_and_vfm_ard_m_pv_2012"]
    B.inp(I, "PSC_Rate", "PSC discount rate (nominal)", "%", psc["discount_rate_nominal_pct"] / 100)
    B.inp(I, "PSC_Capex", "Raw capital cost", "ARD m", psc["psc"]["raw_capex"])
    B.inp(I, "PSC_OM", "O&M and lifecycle", "ARD m", psc["psc"]["om_and_lifecycle"])
    B.inp(I, "PSC_Toll", "Toll revenue retained by the state", "ARD m", psc["psc"]["toll_revenue_retained"])
    B.inp(I, "PSC_ConRisk", "Construction risk", "ARD m", psc["psc"]["construction_risk"])
    B.inp(I, "PSC_TrafRisk", "Traffic revenue risk", "ARD m", psc["psc"]["traffic_revenue_risk"])
    B.inp(I, "PSC_OpRisk", "Operating risk", "ARD m", psc["psc"]["operating_risk"])
    B.inp(I, "PSC_CN", "Competitive neutrality", "ARD m", psc["psc"]["competitive_neutrality"])
    B.inp(I, "PSC_Ref", "Reference (shadow bid) contribution", "ARD m", psc_ref := 410.0)
    B.inp(I, "PSC_Retained", "PPP retained risks", "ARD m", psc["ppp_shadow_bid"]["retained_risks"])
    B.inp(I, "PSC_CM", "PPP contract management", "ARD m", psc["ppp_shadow_bid"]["contract_management"])
    B.gap(I)
    # ---------------- scenario table
    B.head(I, "Scenario table (columns J to V = scenarios 1 to 13)")
    import case_t as M
    sc = M.SCENARIOS
    cols = list(range(1, 14))
    B.table(I, "ST_No", "Scenario number", "#", cols)
    B.table(I, "ST_Hist", "History run (actual events)", "flag", [1 if s in M.HISTORY_SCEN else 0 for s in cols])
    B.table(I, "ST_HistEnd", "Actual traffic used to", "date", [sc[s][2] if sc[s][2] else D(1900, 1, 1) for s in cols])
    B.table(I, "ST_Ext", "Restructured (extension and new terms)", "flag", [sc[s][3] for s in cols])
    B.table(I, "ST_CPIPath", "CPI path (1 bid 2.5%, 2 actual)", "#", [sc[s][4] for s in cols])
    B.table(I, "ST_TCase", "Traffic case (1 Pellow, 2 Ridgeway, 3 downside, 4 R2023, 5 retender)", "#", [sc[s][1] for s in cols])
    B.table(I, "ST_TFac", "Traffic factor", "factor", [sc[s][5] for s in cols])
    B.table(I, "ST_Lag", "Ramp-up lag", "years", [sc[s][6] for s in cols], fmt="0")
    B.table(I, "ST_CPIOnly", "Toll escalation CPI only", "flag", [sc[s][7] for s in cols])
    B.table(I, "ST_OpexF", "Opex factor", "factor", [sc[s][8] for s in cols])
    B.table(I, "ST_LCF", "Lifecycle factor", "factor", [sc[s][9] for s in cols])
    B.table(I, "ST_RefiAdd", "Refinancing margin add", "%pa", [sc[s][10] for s in cols])
    B.table(I, "ST_HVD", "Heavy-vehicle share change (points)", "%pa", [sc[s][11] for s in cols])
    B.table(I, "ST_Live", "Live senior sizing", "flag", [sc[s][12] for s in cols])
    B.table(I, "ST_LiveR", "Live restructured sculpting", "flag", [sc[s][13] for s in cols])
    B.head(I, "Selected scenario values")
    for nm, lab, u in (("Hist", "History run", "flag"), ("HistEnd", "Actual traffic used to", "date"),
                       ("Ext", "Restructured", "flag"), ("CPIPath", "CPI path", "#"), ("TCase", "Traffic case", "#"),
                       ("TFac", "Traffic factor", "factor"), ("Lag", "Ramp-up lag", "years"),
                       ("CPIOnly", "CPI-only escalation", "flag"), ("OpexF", "Opex factor", "factor"),
                       ("LCF", "Lifecycle factor", "factor"), ("RefiAdd", "Refinancing margin add", "%pa"),
                       ("HVD", "Heavy-vehicle share change", "%pa"), ("Live", "Live senior sizing", "flag"),
                       ("LiveR", "Live restructured sculpting", "flag")):
        B.sc(I, nm, lab + " (selected)", u, f"INDEX({{ST_{nm}.r}},{{Scenario.c}})")
    B.sc(I, "Contrib", "State contribution used", "ARD m", "IF({ContribOpt.c}=2,{Contrib_Solved.c},{Contrib_Bid.c})")
    B.gap(I)
    # ---------------- traffic cases
    B.head(I, "Traffic cases (columns J to N = cases 1 to 5)")
    tc = M.TRAFFIC_CASES
    B.table(I, "TC_L0", "Base level (k trips/day)", "k/day", [tc[c][0] for c in range(1, 6)])
    B.table(I, "TC_BY", "Base year", "year", [tc[c][1] for c in range(1, 6)])
    B.table(I, "TC_G1", "Growth to band 1 end", "%pa", [tc[c][2][0] for c in range(1, 6)])
    B.table(I, "TC_G2", "Growth band 2", "%pa", [tc[c][2][1] for c in range(1, 6)])
    B.table(I, "TC_G3", "Growth after band 2", "%pa", [tc[c][2][2] for c in range(1, 6)])
    for k in range(5):
        B.table(I, f"TC_R{k}", f"Ramp-up factor {2019 + k}", "factor", [tc[c][3][k] for c in range(1, 6)])
    B.table(I, "TC_Upl", "Heavy-vehicle uplift in total trips", "%pa", [tc[c][4] for c in range(1, 6)])
    for nm in ("L0", "BY", "G1", "G2", "G3", "Upl"):
        B.sc(I, "T" + nm, f"Selected traffic case: {nm}", "#", f"INDEX({{TC_{nm}.r}},{{TCase.c}})", fmt="0.000")
    B.sc(I, "TRamp_dummy", "Selected ramp-up factors follow (2019 to 2023)", "", "0")
    for k in range(5):
        B.sc(I, f"TRamp{k}", f"Selected ramp-up factor {2019 + k}", "factor", f"INDEX({{TC_R{k}.r}},{{TCase.c}})")
    B.gap(I)
    # ---------------- CPI, ABBR, actual traffic, support tables
    B.head(I, "Annual and half-year tables")
    years = list(range(2012, 2060))
    B.table(I, "CPI_Year", "Year", "year", years)
    B.table(I, "CPI_Act", "Ardmore CPI, annual average change (actual path)", "%pa", [M.cpi_pct(y, 2) for y in years])
    B.table(I, "CPI_Bid", "Ardmore CPI, bid-date path", "%pa", [M.BID_CPI for y in years])
    codes = [y * 10 + h for y in range(2015, 2027) for h in (1, 2)]
    B.table(I, "ABBR_Code", "ABBR fixing code (year x 10 + half)", "code", codes)
    B.table(I, "ABBR_Val", "6-month ABBR fixing", "%pa", [M.ABBR[(c // 10, c % 10)] for c in codes])
    B.inp(I, "ABBR_Long", "ABBR from 2027", "%pa", M.ABBR_LONG)
    tcodes = sorted(y * 10 + h for (y, h) in M.ACT_TRAFFIC)
    B.table(I, "AT_Code", "Actual traffic code", "code", tcodes)
    B.table(I, "AT_Val", "Actual traffic (k trips/day)", "k/day", [M.ACT_TRAFFIC[(c // 10, c % 10)] for c in tcodes])
    B.table(I, "SUP_Code", "Sponsor support code", "code", [20211, 20212])
    B.table(I, "SUP_Val", "Sponsor support (subordinated shareholder loans)", "ARD m", [22.5, 22.5])
    B.gap(I)
    # ---------------- timeline-aligned inputs
    B.head(I, "Time-based inputs (aligned with the model timeline, column J = first period)")
    prof = [v / 100 for v in cap["dc_payment_profile_pct_by_quarter"].values()] + [0.0] * (T - 16)
    B.tin(I, "DC_Prof", "D&C payment profile (excludes 5% retention)", "%", prof, total=True, fmt="0.0000%")
    L = OUT["locked"]
    B.head(I, "Locked financing from the banking run (Scenario 2); Checks confirms live = locked")
    B.inp(I, "Locked_D", "Senior debt (bank plus bonds)", "ARD m", L["D"], fmt='#,##0.000000')
    B.inp(I, "Locked_N", "NILO loan", "ARD m", L["N"], fmt='#,##0.000000')
    B.inp(I, "Locked_E", "Equity", "ARD m", L["E"], fmt='#,##0.000000')
    B.tin(I, "Locked_P", "Scheduled senior principal", "ARD m", L["P"], total=True, fmt='#,##0.0000')
    B.tin(I, "Locked_RIA", "Ramp-up interest account releases", "ARD m", L["ria"], total=True, fmt='#,##0.0000')
    B.tin(I, "Locked_nP", "Scheduled NILO principal (original)", "ARD m", L["nP"], total=True, fmt='#,##0.0000')
    B.tin(I, "Locked_DownCF", "Downside CFADS (from the downside run)", "ARD m", L["down_cfads"], total=True, fmt='#,##0.0000')
    B.head(I, "Locked restructured profiles from Scenario 5")
    B.tin(I, "Locked_NotesP", "Scheduled restructured notes principal", "ARD m", L["notesP"], total=True, fmt='#,##0.0000')
    B.tin(I, "Locked_nPr", "Scheduled NILO principal (restructured)", "ARD m", L["nPr"], total=True, fmt='#,##0.0000')

    # ===================================================================== TIME
    S = "Time"
    B.head(S, "Timeline")
    B.tr(S, "Period", "Period number", "#", "{Period.p}+1")
    B.tr(S, "Start", "Period start", "date", "IF({Period}=1,{FC_Date.c},{End.p})")
    B.tr(S, "End", "Period end", "date", "IF({Period}=1,{FirstEnd.c},EOMONTH({End.p},IF({End.p}<={Con_End.c},3,6)))")
    B.tr(S, "Days", "Days in period", "days", "{End}-{Start}")
    B.tr(S, "D360", "Days in period (30/360)", "days",
         "{Days360Y.c}*(YEAR({End})-YEAR({Start}))+30*(MONTH({End})-MONTH({Start}))+MIN(DAY({End}),30)-MIN(DAY({Start}),30)")
    B.tr(S, "Year", "Calendar year", "year", "YEAR({End})")
    B.tr(S, "Half", "Half (1 = Jan-Jun, 2 = Jul-Dec)", "#", "IF(MONTH({End})<=6,1,2)")
    B.tr(S, "Code", "Half-year code", "code", "{Year}*{HalfCode.c}+{Half}")
    B.tr(S, "Months", "Months in period", "#", "ROUND({Days}/{DaysMonth.c},0)")
    B.head(S, "Key dates for the selected scenario")
    B.sc(S, "Opening", "Opening date", "date", "IF({Hist.c}=1,{Open_Act.c},{Open_Sched.c})")
    B.sc(S, "ConcFinal", "Concession expiry", "date", "IF({Ext.c}=1,{Conc_Ext.c},{Conc_Orig.c})")
    B.sc(S, "ConDays", "Construction days (close to scheduled completion)", "days", "{Con_End.c}-{FC_Date.c}")
    B.head(S, "Flags")
    B.tr(S, "Flag_Construction", "Flag_Construction", "flag", "IF({End}<={Con_End.c},1,0)")
    B.tr(S, "Flag_FirstPost", "Flag_FirstPostConstruction", "flag", "IF(AND({Flag_Construction.p}=1,{Flag_Construction}=0),1,0)")
    B.tr(S, "Flag_Last6", "Flag_BridgeDraw (last six construction quarters)", "flag",
         "IF(AND({Flag_Construction}=1,{End}>EDATE({Con_End.c},-{Last6Months.c})),1,0)")
    B.tr(S, "Post", "Flag_PostRestructuring", "flag", "IF(AND({Ext.c}=1,{End}>{Restr_Date.c}),1,0)")
    B.tr(S, "ConcT", "Applicable concession expiry", "date", "IF({Post}=1,{Conc_Ext.c},{Conc_Orig.c})")
    B.tr(S, "OpsDays", "Operating days", "days", "MAX(0,MIN({End},{ConcFinal.c})-MAX({Start},{Opening.c}))")
    B.tr(S, "Flag_Ops", "Flag_Operations", "flag", "IF({OpsDays}>0,1,0)")
    B.tr(S, "Flag_OpenPer", "Flag_OpeningPeriod", "flag", "IF(AND({Start}<{Opening.c},{Opening.c}<={End}),1,0)")
    B.tr(S, "Flag_Final", "Flag_FinalPeriod (concession expiry)", "flag", "IF(AND({Start}<{ConcFinal.c},{ConcFinal.c}<={End}),1,0)")
    B.tr(S, "RemDays", "Remaining concession days from period start", "days", "MAX(0,{ConcT}-MAX({Start},{Opening.c}))")
    B.tr(S, "Flag_Rep", "Flag_Repayment (senior, original)", "flag", "IF(AND({End}>={First_Rep.c},{End}<={Final_Rep.c}),1,0)")
    B.tr(S, "Flag_FirstRep", "Flag_FirstRepayment", "flag", "IF(AND({Flag_Rep}=1,{Flag_Rep.p}=0),1,0)")
    B.tr(S, "Flag_RampUp", "Flag_RampUp (operations before first repayment)", "flag", "IF(AND({Flag_Ops}=1,{End}<{First_Rep.c}),1,0)")
    B.tr(S, "Flag_PreFirst", "Flag_BeforeFirstRepayment", "flag", "IF({End}<{First_Rep.c},1,0)")
    B.tr(S, "Flag_Standstill", "Flag_Standstill", "flag", "IF(AND({Hist.c}=1,{End}>={First_Rep.c},{End}<={Restr_Date.c}),1,0)")
    B.tr(S, "Flag_Restr", "Flag_Restructuring", "flag", "IF(AND({Ext.c}=1,{End}={Restr_Date.c}),1,0)")
    B.tr(S, "Flag_NotesFirst", "Flag_FirstNotesPeriod", "flag", "IF(AND({Post}=1,{Flag_Restr.p}=1),1,0)")
    B.tr(S, "Flag_NotesRep", "Flag_NotesRepayment", "flag", "IF(AND({Post}=1,{End}<={Notes_Final.c}),1,0)")
    B.tr(S, "Flag_NotesSweep", "Flag_NotesSweep", "flag", "IF(AND({Post}=1,{End}<={Sweep_End.c}),1,0)")
    B.tr(S, "Flag_History", "Flag_ActualTraffic", "flag", "IF(AND({Hist.c}=1,{Flag_Ops}=1,{End}<={HistEnd.c}),1,0)")
    B.tr(S, "F1", "Share of period before margin step", "factor", "MIN(1,MAX(0,(MIN({End},{Margin_Step.c})-{Start})/{Days}))")
    B.tr(S, "F2", "Share of period before mini-perm maturity", "factor", "MIN(1,MAX(0,(MIN({End},{MPM.c})-{Start})/{Days}))")
    B.tr(S, "Flag_MPM", "Flag_MiniPermMaturity", "flag", "IF(AND({Start}<{MPM.c},{MPM.c}<={End}),1,0)")
    B.tr(S, "Flag_Hedge", "Flag_Hedged", "flag", "IF({End}<={Hedge_End.c},1,0)")
    B.tr(S, "Flag_TestDate", "Flag_DSCRTest", "flag", "IF(AND({Flag_Ops}=1,{End}>={First_Test.c}),1,0)")
    B.tr(S, "Flag_EoDTest", "Flag_DefaultTest", "flag", "IF({End}>={Default_Test.c},1,0)")
    B.tr(S, "Flag_RetRel", "Flag_RetentionRelease", "flag",
         "IF(AND({Start}<EDATE({Opening.c},{RetMonths.c}),EDATE({Opening.c},{RetMonths.c})<={End}),1,0)")
    B.tr(S, "Flag_NiloRep", "Flag_NILORepayment", "flag",
         "IF({Ext.c}=1,IF(AND({End}>={NiloR_First.c},{End}<={NiloR_Final.c}),1,0),IF(AND({End}>={Nilo_First.c},{End}<={Nilo_Final.c}),1,0))")
    B.tr(S, "Flag_NiloFirst", "Flag_NILOFirstRepayment", "flag", "IF(AND({Flag_NiloRep}=1,{Flag_NiloRep.p}=0),1,0)")
    B.tr(S, "Flag_NiloStart", "Flag_NILOScheduleStart", "flag", "IF({Ext.c}=1,{Flag_NotesFirst},{Flag_FirstPost})")
    B.tr(S, "NRate", "NILO interest rate", "%", "IF(AND({Post}=1,{End}<={PIK_End.c}),{Nilo_PIK.c},{Nilo_Rate.c})")
    B.tr(S, "NCash", "NILO interest paid in cash (share)", "factor",
         "IF(AND({Post}=1,{End}<={PIK_End.c}),0,1-MIN(1,MAX(0,(MIN({End},{Nilo_CapEnd.c})-{Start})/{Days})))")
    B.tr(S, "Flag_SwapMTM", "Flag_SwapRemaining (after 2023-12-31)", "flag", "IF(AND({End}>{Restr_Date.c},{Flag_Hedge}=1),1,0)")
    B.head(S, "Macro")
    B.tr(S, "ABBR", "6-month ABBR", "%pa", "IFERROR(INDEX({ABBR_Val.r},MATCH({Code},{ABBR_Code.r},0)),{ABBR_Long.c})")
    B.tr(S, "CPI", "Ardmore CPI change for the year (selected path)", "%pa",
         "IF({CPIPath.c}=1,INDEX({CPI_Bid.r},MATCH({Year},{CPI_Year.r},0)),INDEX({CPI_Act.r},MATCH({Year},{CPI_Year.r},0)))")
    B.tr(S, "CPIf", "CPI factor (2015 = 1.0)", "factor", "IF({Period}=1,1,{CPIf.p}*IF({Year}>{Year.p},1+{CPI}/100,1))", fmt="0.000000")
    B.sc(S, "CPIf2023", "CPI factor 2023 (2015 = 1.0)", "factor",
         "IF({CPIPath.c}=1,EXP(SUMPRODUCT(({CPI_Year.r}>={CPI_From.c})*({CPI_Year.r}<={CPI_To.c})*LN(1+{CPI_Bid.r}/100))),"
         "EXP(SUMPRODUCT(({CPI_Year.r}>={CPI_From.c})*({CPI_Year.r}<={CPI_To.c})*LN(1+{CPI_Act.r}/100))))", fmt="0.000000")
    B.tr(S, "Flag_July", "Flag_TollEscalation (contains July 1)", "flag",
         "IF(AND({Start}<DATE({Year},{EscMonth.c},1),DATE({Year},{EscMonth.c},1)<={End}),1,0)")
    B.tr(S, "Esc", "Toll escalation factor", "factor",
         "IF(OR({CPIOnly.c}=1,AND({Ext.c}=1,{Year}>={Restr_Esc_From.c}),{Year}>{Esc_Floor_Last.c}),1+{CPI}/100,MAX(1+{CPI}/100,1+{Esc_Floor.c}))",
         fmt="0.000000")
    B.tr(S, "Toll", "Car toll per km", "ARD/km",
         "IF({Period}=1,IF({Flag_July}=1,ROUND({Toll2014.c}*{Esc},{Toll_dp.c}),{Toll2014.c}),IF({Flag_July}=1,ROUND({Toll.p}*{Esc},{Toll_dp.c}),{Toll.p}))")

    # ===================================================================== CONSTRUCTION
    S = "Construction"
    B.head(S, "Construction costs (paid at period end)")
    B.tr(S, "DC", "D&C contract payments (incl. 5% retention at completion)", "ARD m",
         "{DC_Price.c}*{DC_Prof}+IF({End}={Con_End.c},{DC_Price.c}*({Ret1.c}+{Ret2.c}),0)", total=True)
    B.tr(S, "DevC", "Development and bid costs", "ARD m", "IF({Period}=1,{Dev.c},0)", total=True)
    B.tr(S, "SPVCc", "Project company costs", "ARD m", "{SPVC.c}*{Days}/{ConDays.c}*{Flag_Construction}", total=True)
    B.tr(S, "InsCc", "Insurance during construction", "ARD m",
         "IF({Period}=1,{InsC.c}*{InsC1.c},IF({End}={Ins2_Date.c},{InsC.c}*(1-{InsC1.c}),0))", total=True)
    B.tr(S, "CertC", "Independent certifier", "ARD m", "{Cert.c}*{Days}/{ConDays.c}*{Flag_Construction}", total=True)
    B.tr(S, "AdvC", "Lenders' advisors and legal", "ARD m", "{Adv.c}*{Days}/{ConDays.c}*{Flag_Construction}", total=True)
    B.tr(S, "ContC", "Contingency", "ARD m", "{Cont.c}*{DC_Prof}/SUM({DC_Prof.r})", total=True)
    B.tr(S, "Capex", "Total construction costs before financing", "ARD m",
         "{DC}+{DevC}+{SPVCc}+{InsCc}+{CertC}+{AdvC}+{ContC}", total=True)
    B.tr(S, "RetDep", "Retention deposited at completion (released after defects period)", "ARD m",
         "IF({End}={Con_End.c},{DC_Price.c}*{Ret2.c},0)", total=True)

    # ===================================================================== OPERATIONS
    S = "Operations"
    B.head(S, "Traffic")
    B.tr(S, "G", "Growth rate for the year", "%pa", "IF({Year}<={Band1_End.c},{TG1.c},IF({Year}<={Band2_End.c},{TG2.c},{TG3.c}))")
    B.tr(S, "Level", "Underlying traffic level", "k/day",
         "IF({Year}<={TBY.c},{TL0.c}/(1+{TG1.c}/100)^({TBY.c}-{Year}),{Level.p}*IF({Year}>{Year.p},1+{G}/100,1))", fmt="0.0000")
    B.tr(S, "RampYear", "Ramp-up table year", "year", "MAX({Ramp_First.c},MIN({Ramp_Last.c},{Year}-{Lag.c}))")
    B.tr(S, "Ramp", "Ramp-up factor", "factor",
         "CHOOSE({RampYear}-{Ramp_First.c}+1,{TRamp0.c},{TRamp1.c},{TRamp2.c},{TRamp3.c},{TRamp4.c})")
    B.tr(S, "Uplift", "Heavy-vehicle uplift factor", "factor", "IF({Year}>={Uplift_From.c},1+{TUpl.c}/100,1)")
    B.tr(S, "TrafficFc", "Forecast traffic", "k/day", "{Level}*{Ramp}*{Uplift}*{TFac.c}", fmt="0.0000")
    B.tr(S, "TrafficAct", "Actual traffic", "k/day", "IFERROR(INDEX({AT_Val.r},MATCH({Code},{AT_Code.r},0)),0)")
    B.tr(S, "Traffic", "Traffic used (average daily trips)", "k/day", "IF({Flag_History}=1,{TrafficAct},{TrafficFc})*{Flag_Ops}", fmt="0.0000")
    B.head(S, "Toll revenue")
    B.sc(S, "W_Orig", "Weighted class multiplier, original", "factor",
         "(({Mix_Car.c}-{HVD.c})*{Mult_Car.c}+{Mix_LCV.c}*{Mult_LCV.c}+({Mix_HV.c}+{HVD.c})*{Mult_HV.c})/100", fmt="0.000000")
    B.sc(S, "W_New", "Weighted class multiplier, from 2024 (restructured)", "factor",
         "({MixN_Car.c}*{Mult_Car.c}+{MixN_LCV.c}*{Mult_LCV.c}+{MixN_HV.c}*{MultN_HV.c})/100", fmt="0.000000")
    B.tr(S, "WMult", "Weighted class multiplier", "factor", "IF({Post}=1,{W_New.c},{W_Orig.c})", fmt="0.000000")
    B.tr(S, "Gross", "Gross toll revenue", "ARD m",
         "{Traffic}*{KTrips.c}*{OpsDays}*{TripKm.c}*{Toll}*{WMult}/{Million.c}", total=True)
    B.tr(S, "LeakR", "Revenue leakage", "ARD m", "{Gross}*{Leak.c}", total=True)
    B.tr(S, "NetRev", "Net toll revenue", "ARD m", "{Gross}-{LeakR}", total=True)
    B.tr(S, "LD", "Delay liquidated damages received", "ARD m", "IF(AND({Hist.c}=1,{Flag_OpenPer}=1),{LD_Act.c},0)", total=True)
    B.head(S, "Operating costs")
    B.tr(S, "OpxF", "Cost factor (CPI x operating days / 365 x opex factor)", "factor",
         "{CPIf}*{OpsDays}/{DaysYear.c}*{OpexF.c}", fmt="0.000000")
    B.tr(S, "OMc", "Corvus fixed O&M", "ARD m", "{OM.c}*{OpxF}*IF({Post}=1,1-{OM_Cut.c},1)", total=True)
    B.tr(S, "BOFix", "Tolling back office, fixed", "ARD m", "{BO_Fix.c}*{OpxF}", total=True)
    B.tr(S, "BOVar", "Tolling back office, variable", "ARD m", "{BO_Pct.c}*{NetRev}*{OpexF.c}", total=True)
    B.tr(S, "InsO", "Insurance", "ARD m", "{Ins.c}*{OpxF}", total=True)
    B.tr(S, "SPVO", "Project company costs", "ARD m", "{SPV.c}*{OpxF}", total=True)
    B.tr(S, "Opex", "Total operating costs", "ARD m", "{OMc}+{BOFix}+{BOVar}+{InsO}+{SPVO}", total=True)
    B.tr(S, "RevSh", "State revenue share", "ARD m",
         "IF({Post}=1,{RevShare.c}*MAX(0,{NetRev}-{RevThresh.c}*{CPIf}/{CPIf2023.c}*{OpsDays}/{DaysYear.c}),0)", total=True)
    B.tr(S, "RestrCost", "Restructuring costs", "ARD m",
         "IF(AND({Flag_History}=1,{End}>={RC_Start.c},{End}<={Restr_Date.c}),{RC_Total.c}/{RC_N.c},0)", total=True)
    B.tr(S, "EBITDA", "EBITDA", "ARD m", "{NetRev}+{LD}-{Opex}-{RevSh}-{RestrCost}", total=True)
    B.head(S, "Lifecycle and handback")
    B.tr(S, "OY", "Operating year (H2 periods)", "#", "IF(AND({Half}=2,{Flag_Ops}=1,{End}<={ConcFinal.c}),{Year}-{OY_Base.c},0)")
    B.tr(S, "LC2015", "Lifecycle works (2015 prices)", "ARD m",
         "IF({OY}>0,{LC_Pav.c}*IF(MOD({OY},{Cyc_Pav.c})=0,1,0)+{LC_Tun.c}*IF(MOD({OY},{Cyc_Tun.c})=0,1,0)+{LC_ITS.c}*IF(MOD({OY},{Cyc_ITS.c})=0,1,0),0)", total=True)
    B.tr(S, "LC", "Lifecycle works (nominal), paid from the MMRA", "ARD m", "{LC2015}*{CPIf}*{LCF.c}", total=True)
    B.tr(S, "MMRAc", "Lifecycle reserve contribution (next six periods / 6)", "ARD m", "SUM({LC.n6})/{MMRA_N.c}", total=True)
    B.tr(S, "HBFlag", "Flag_HandbackFunding", "flag", "IF(AND({Flag_Ops}=1,{Start}>=EDATE({ConcFinal.c},-{HB_Months.c})),1,0)")
    B.tr(S, "HBc", "Handback reserve contribution", "ARD m", "{HB_Est.c}*{CPIf}*{LCF.c}/{HB_N.c}*{HBFlag}", total=True)
    B.head(S, "Working capital")
    B.tr(S, "Rec", "Toll receivables", "ARD m", "IF(AND({Flag_Ops}=1,{Flag_Final}=0),{NetRev}/{OpsDays}*{RecDays.c},0)")
    B.tr(S, "Pay", "Operating payables", "ARD m", "IF(AND({Flag_Ops}=1,{Flag_Final}=0),{Opex}/{OpsDays}*{PayDays.c},0)")
    B.tr(S, "NWC", "Net working capital", "ARD m", "{Rec}-{Pay}")
    B.tr(S, "dWC", "Increase in working capital", "ARD m", "{NWC}-{NWC.p}", total=True)
    B.tr(S, "UpgCapex", "Interchange upgrade capex (from the upgrade account)", "ARD m",
         "IF(AND({Ext.c}=1,{End}>={Upg_First.c},{End}<={Upg_Last.c}),{Upgrade.c}/{Upg_N.c},0)", total=True)
    B.head(S, "Cash flow available for debt service")
    B.tr(S, "TaxL", "Tax paid", "ARD m", "{Tax}", total=True)
    B.tr(S, "CFADS", "CFADS", "ARD m", "({EBITDA}-{dWC}-{TaxL}-{MMRAc}-{HBc})*(1-{Flag_Construction})", total=True)

    # ===================================================================== FUNDING
    S = "Funding"
    B.head(S, "Commitments")
    B.sc(S, "BankC", "Bank mini-perm commitment", "ARD m", "{Bank_Split.c}*{Dsen.c}")
    B.sc(S, "BondF", "Bond face value (fully funded at close into escrow)", "ARD m", "{Bond_Split.c}*{Dsen.c}")
    B.head(S, "Uses of funds")
    B.tr(S, "CapexL", "Construction costs", "ARD m", "{Capex}", total=True)
    B.tr(S, "Fees", "Upfront fees (bank, bond issue, NILO, bridge)", "ARD m",
         "IF({Period}=1,{Bank_Up.c}*{BankC.c}+{BondCost.c}*{BondF.c}+{Nilo_Fee.c}*{Nsub.c}+{Bridge_Fee.c}*{Contrib.c},0)", total=True)
    B.tr(S, "BankIntC", "Bank interest during construction", "ARD m", "{BankInt}*{Flag_Construction}", total=True)
    B.tr(S, "Commit", "Bank commitment fee", "ARD m",
         "{Bank_Commit.c}*({BankC.c}-{BankOpen})*{Days}/{DaysYear.c}*{Flag_Construction}", total=True)
    B.tr(S, "BondIntC", "Bond interest during construction", "ARD m", "{BondInt}*{Flag_Construction}", total=True)
    B.tr(S, "EscrowOpen", "Bond escrow: opening", "ARD m", "IF({Period}=1,{BondF.c},{EscrowClose.p})")
    B.tr(S, "EscrowInt", "Bond escrow earnings (deducted)", "ARD m", "{EscrowOpen}*{Escrow_Rate.c}*{Days}/{DaysYear.c}*{Flag_Construction}", total=True)
    B.tr(S, "EscrowClose", "Bond escrow: closing", "ARD m", "{EscrowOpen}-{BondDraw}")
    B.tr(S, "BridgeOpen", "Contribution bridge: opening", "ARD m", "{BridgeClose.p}")
    B.tr(S, "BridgeInt", "Bridge interest", "ARD m",
         "{BridgeOpen}*({ABBR}/100+{Bridge_Margin.c})*MAX(0,MIN({End},{Opening.c})-{Start})/{DaysYear.c}", total=True)
    B.tr(S, "BridgeIntC", "Bridge interest during construction", "ARD m", "{BridgeInt}*{Flag_Construction}", total=True)
    B.tr(S, "BridgeIntOps", "Bridge interest after completion (paid from operations)", "ARD m", "{BridgeInt}*(1-{Flag_Construction})", total=True)
    B.tr(S, "DSRAInit", "DSRA initial funding", "ARD m", "IF({End}={Con_End.c},{DSRATarget},0)", total=True)
    B.tr(S, "RIAInit", "Ramp-up interest account funding", "ARD m", "IF({End}={Con_End.c},SUM({RIA.r}),0)", total=True)
    B.tr(S, "Uses", "Total uses", "ARD m",
         "({CapexL}+{Fees}+{BankIntC}+{Commit}+{BondIntC}-{EscrowInt}+{BridgeIntC}+{DSRAInit}+{RIAInit})*{Flag_Construction}", total=True)
    B.sc(S, "TotalUses", "Total uses (eligible costs)", "ARD m", "SUM({Uses.r})")
    B.sc(S, "Fnet", "Funding requirement net of the state contribution", "ARD m", "{TotalUses.c}-{Contrib.c}")
    B.head(S, "Sources of funds (equity first, then debt pro rata; bridge in the last six quarters)")
    B.tr(S, "EqDraw", "Equity drawn", "ARD m", "MIN({Uses},MAX(0,{Eq.c}-{CumEq.p}))*{Flag_Construction}", total=True)
    B.tr(S, "CumEq", "Cumulative equity", "ARD m", "{CumEq.p}+{EqDraw}")
    B.tr(S, "Need", "Funding need after equity", "ARD m", "{Uses}-{EqDraw}", total=True)
    B.sc(S, "R6", "Need in the bridge draw quarters", "ARD m", "SUMPRODUCT({Need.r},{Flag_Last6.r})")
    B.tr(S, "BridgeDraw", "Bridge drawn", "ARD m", "IF({Flag_Last6}=1,{Contrib.c}*{Need}/{R6.c},0)", total=True)
    B.tr(S, "DebtDraw", "Debt drawn (senior plus NILO)", "ARD m", "{Need}-{BridgeDraw}", total=True)
    B.tr(S, "SenDraw", "Senior drawn", "ARD m", "{DebtDraw}*{Dsen.c}/({Dsen.c}+{Nsub.c})", total=True)
    B.tr(S, "NiloDraw", "NILO drawn", "ARD m", "{DebtDraw}-{SenDraw}", total=True)
    B.tr(S, "BankDraw", "Bank drawn", "ARD m", "{Bank_Split.c}*{SenDraw}", total=True)
    B.tr(S, "BondDraw", "Bond escrow released", "ARD m", "{Bond_Split.c}*{SenDraw}", total=True)
    B.tr(S, "ShareCapDraw", "Equity: share capital", "ARD m", "{ShareCap.c}*{EqDraw}", total=True)
    B.tr(S, "SHLDraw", "Equity: shareholder loans", "ARD m", "{EqDraw}-{ShareCapDraw}", total=True)
    B.tr(S, "Contribution", "State contribution received (repays the bridge)", "ARD m", "{Contrib.c}*{Flag_OpenPer}", total=True)
    B.tr(S, "BridgeClose", "Contribution bridge: closing", "ARD m", "{BridgeOpen}+{BridgeDraw}-{Contribution}")
    B.tr(S, "CapCost", "Capitalized cost (uses less reserves plus capitalized NILO and SHL interest)", "ARD m",
         "({Uses}-{DSRAInit}-{RIAInit}+{NiloAccr}+{SHLInt})*{Flag_Construction}", total=True)
    B.tr(S, "SUCheck", "Sources less uses", "ARD m", "{EqDraw}+{BridgeDraw}+{DebtDraw}-{Uses}", total=True)

    # ===================================================================== DEBT
    S = "Debt"
    B.head(S, "Rates")
    B.sc(S, "M3", "Bank margin after the mini-perm maturity", "%pa", "IF({Hist.c}=1,{M3_AE.c},{M3_Refi.c}+{RefiAdd.c})")
    B.tr(S, "Margin", "Bank margin (day-weighted)", "%pa", "{M1.c}*{F1}+{M2.c}*({F2}-{F1})+{M3.c}*(1-{F2})")
    B.tr(S, "BankRate", "Bank all-in rate", "%", "(IF({Flag_Hedge}=1,{Swap.c},{ABBR})+{Margin})/100", fmt="0.0000%")
    B.tr(S, "RR", "Senior blended period rate", "%",
         "{Bank_Split.c}*{BankRate}*{Days}/{DaysYear.c}+{Bond_Split.c}*{BondCpn.c}*{D360}/{Days360Y.c}", fmt="0.0000%")
    B.tr(S, "RRN", "Notes period rate", "%", "{NotesCpn.c}*{D360}/{Days360Y.c}", fmt="0.0000%")
    B.head(S, "Senior sizing (live in Scenario 2)")
    B.tr(S, "DF", "Discount factor from 2021-01-01 at the senior rate", "factor",
         "IF({Flag_Rep}=1,IF({Flag_Rep.p}=1,{DF.p},1)/(1+{RR}),0)", fmt="0.000000")
    B.tr(S, "CFADSL", "CFADS", "ARD m", "{CFADS}", total=True)
    B.sc(S, "PVCF", "PV of CFADS over the repayment periods", "ARD m", "SUMPRODUCT({CFADSL.r},{DF.r},{Flag_Rep.r})")
    B.sc(S, "D_Gear", "Capacity: 55% of funding requirement net of contribution", "ARD m", "{GearCap.c}*{Fnet.c}")
    B.sc(S, "D_Sculpt", "Capacity: sculpted at 1.50x (PV / 1.50)", "ARD m", "{PVCF.c}/{DSCR_T.c}")
    B.sc(S, "D_LLCR", "Capacity: LLCR 1.55x", "ARD m", "{PVCF.c}/{LLCR_T.c}")
    B.tr(S, "IOCap", "Interest-only capacity at 1.50x", "ARD m", "IF({Flag_Rep}=1,{CFADSL}/({DSCR_T.c}*{RR}),{Big.c})", fmt='#,##0.0;(#,##0.0);"–"')
    B.sc(S, "D_IO", "Capacity: interest cover 1.50x in every repayment period", "ARD m", "MIN({IOCap.r})")
    B.tr(S, "IODown", "Downside interest-only capacity at 1.15x", "ARD m",
         "IF({Flag_Rep}=1,{Locked_DownCF}/({DSCR_Down.c}*{RR}),{Big.c})", fmt='#,##0.0;(#,##0.0);"–"')
    B.sc(S, "D_Down", "Capacity: downside interest cover 1.15x", "ARD m", "MIN({IODown.r})")
    B.sc(S, "D_Live", "Senior debt, live sizing (lesser of the constraints)", "ARD m",
         "MIN({D_Gear.c},{D_Sculpt.c},{D_LLCR.c},{D_IO.c},{D_Down.c})")
    B.sc(S, "Binding", "Binding constraint (1 gearing, 2 sculpt, 3 LLCR, 4 interest cover, 5 downside)", "#",
         "IF({D_Live.c}={D_Gear.c},1,IF({D_Live.c}={D_Sculpt.c},2,IF({D_Live.c}={D_LLCR.c},3,IF({D_Live.c}={D_IO.c},4,5))))")
    B.sc(S, "LiveOn", "Live sizing active (Live and Circ)", "flag", "IF(AND({Circ.c}=1,{Live.c}=1),1,0)")
    B.sc(S, "LiveROn", "Live restructured sculpting active", "flag", "IF(AND({Circ.c}=1,{LiveR.c}=1),1,0)")
    B.sc(S, "Dsen", "Senior debt applied", "ARD m", "IF({LiveOn.c}=1,{D_Live.c},{Locked_D.c})", fmt='#,##0.000')
    B.sc(S, "N_Live", "NILO, live (remainder, capped at 33% of eligible costs)", "ARD m",
         "MIN({Fnet.c}-{EqShare.c}*{Fnet.c}-{D_Live.c},{Nilo_Max.c}*{TotalUses.c})")
    B.sc(S, "Nsub", "NILO applied", "ARD m", "IF({LiveOn.c}=1,{N_Live.c},{Locked_N.c})", fmt='#,##0.000')
    B.sc(S, "Eq", "Equity applied", "ARD m", "IF({LiveOn.c}=1,{Fnet.c}-{Dsen.c}-{Nsub.c},{Locked_E.c})", fmt='#,##0.000')
    B.head(S, "Senior sculpting: DS = MAX(interest, CFADS / s) from June 2021")
    B.tr(S, "SBOpen", "Scheduled balance: opening", "ARD m", "IF({Flag_FirstPost}=1,{Dsen.c},{SBClose.p})")
    B.tr(S, "SInt", "Scheduled interest", "ARD m", "{SBOpen}*{RR}", total=True)
    B.tr(S, "SA", "Flag_InterestOnly (interest above CFADS / s)", "flag", "IF(AND({Flag_Rep}=1,{SInt}>{CFADSL}/{Sdiv.c}),1,0)")
    B.tr(S, "SDS", "Scheduled debt service", "ARD m", "IF({Flag_Rep}=1,MAX({SInt},{CFADSL}/{Sdiv.c}),0)", total=True)
    B.tr(S, "SP", "Scheduled principal (live)", "ARD m", "IF({Flag_Rep}=1,{SDS}-{SInt},0)", total=True)
    B.tr(S, "SBClose", "Scheduled balance: closing", "ARD m", "{SBOpen}-{SP}")
    B.sc(S, "Sdiv", "Sculpting divisor s (fixed point)", "x",
         "IF({LiveOn.c}=1,SUMPRODUCT({CFADSL.r}*{DF.r}*{Flag_Rep.r}*(1-{SA.r}))/({Dsen.c}-SUMPRODUCT({SInt.r}*{DF.r}*{SA.r})),{DSCR_T.c})", fmt="0.000000")
    B.tr(S, "RIALive", "Ramp-up interest account release (live)", "ARD m",
         "IF({Flag_RampUp}=1,MAX(0,{SInt}-{CFADSL}/{RIA_DSCR.c}),0)", total=True)
    B.tr(S, "RIA", "Ramp-up interest account release (applied)", "ARD m", "IF({LiveOn.c}=1,{RIALive},{Locked_RIA})", total=True)
    B.tr(S, "Psched", "Scheduled senior principal (applied)", "ARD m", "IF({LiveOn.c}=1,{SP},{Locked_P})", total=True)
    B.head(S, "Bank mini-perm")
    B.tr(S, "BankOpen", "Opening balance", "ARD m", "{BankClose.p}")
    B.tr(S, "BankInt", "Interest", "ARD m", "{BankOpen}*{BankRate}*{Days}/{DaysYear.c}", total=True)
    B.tr(S, "BankIntDue", "Interest due from operations", "ARD m", "{BankInt}*(1-{Flag_Construction})", total=True)
    B.tr(S, "PDueBank", "Principal due", "ARD m",
         "IF(OR({Flag_Construction}=1,{Flag_Standstill}=1,{Post}=1),0,MIN({Bank_Split.c}*{Psched},{BankOpen}))", total=True)
    B.tr(S, "BankPre", "Closing balance before restructuring", "ARD m", "{BankOpen}+{BankDraw}-{PPaidBank}-{Sweep1Bank}")
    B.tr(S, "BankClose", "Closing balance", "ARD m", "IF({Flag_Restr}=1,0,{BankPre})")
    B.head(S, "Revenue bonds")
    B.tr(S, "BondOpen", "Opening balance", "ARD m", "IF({Period}=1,{BondF.c},{BondClose.p})")
    B.tr(S, "BondInt", "Interest (30/360)", "ARD m", "{BondOpen}*{BondCpn.c}*{D360}/{Days360Y.c}", total=True)
    B.tr(S, "BondIntDue", "Interest due from operations", "ARD m", "{BondInt}*(1-{Flag_Construction})", total=True)
    B.tr(S, "PDueBond", "Principal due", "ARD m",
         "IF(OR({Flag_Construction}=1,{Flag_Standstill}=1,{Post}=1),0,MIN({Bond_Split.c}*{Psched},{BondOpen}))", total=True)
    B.tr(S, "BondPre", "Closing balance before restructuring", "ARD m", "{BondOpen}-{PPaidBond}-{Sweep1Bond}")
    B.tr(S, "BondClose", "Closing balance", "ARD m", "IF({Flag_Restr}=1,0,{BondPre})")
    B.head(S, "Senior arrears (interest unpaid under the standstill)")
    B.tr(S, "ArrBankOpen", "Bank arrears: opening", "ARD m", "{ArrBankClose.p}")
    B.tr(S, "ArrBondOpen", "Bond arrears: opening", "ARD m", "{ArrBondClose.p}")
    B.tr(S, "ArrBankPre", "Bank arrears: closing before restructuring", "ARD m", "({ArrBankOpen}+{BankIntDue}-{IntPaid}*{ShB})*(1-{Flag_Construction})")
    B.tr(S, "ArrBondPre", "Bond arrears: closing before restructuring", "ARD m", "({ArrBondOpen}+{BondIntDue}-{IntPaid}*{ShO})*(1-{Flag_Construction})")
    B.tr(S, "ArrBankClose", "Bank arrears: closing", "ARD m", "IF({Flag_Restr}=1,0,{ArrBankPre})")
    B.tr(S, "ArrBondClose", "Bond arrears: closing", "ARD m", "IF({Flag_Restr}=1,0,{ArrBondPre})")
    B.head(S, "Debt service due (all senior)")
    B.tr(S, "NotesIntDue", "Restructured notes interest due (30/360)", "ARD m", "{NotesOpen}*{RRN}", total=True)
    B.tr(S, "IntDue", "Senior interest due", "ARD m", "{BankIntDue}+{BondIntDue}+{NotesIntDue}", total=True)
    B.tr(S, "FeesDue", "Refinancing / A&E fee and post-completion bridge interest", "ARD m",
         "IF({Hist.c}=1,{AE_Fee.c},{Refi_Fee.c})*{BankOpen}*{Flag_MPM}+{BridgeIntOps}", total=True)
    B.tr(S, "NotesPDue", "Notes principal due", "ARD m", "IF({Post}=1,MIN({NotesPApp},{NotesOpen}),0)", total=True)
    B.tr(S, "PDue", "Senior principal due", "ARD m", "{PDueBank}+{PDueBond}+{NotesPDue}", total=True)
    B.tr(S, "SchedDS", "Scheduled senior cash debt service (covenant basis)", "ARD m",
         "{IntDue}+({Bank_Split.c}+{Bond_Split.c})*{Psched}*(1-{Post})*(1-{Flag_Construction})+{NotesPDue}-{RIA}", total=True)
    B.head(S, "Restructuring at 2023-12-31")
    B.tr(S, "OSB", "Original scheduled senior balance (swap notional basis)", "ARD m", "IF({Flag_FirstPost}=1,{Dsen.c},{OSB.p}-{Psched.p})")
    B.tr(S, "SwapDF", "Swap discount factor at 4.36%", "factor",
         "IF({Flag_SwapMTM}=1,IF({Flag_SwapMTM.p}=1,{SwapDF.p},1)/(1+{SwapMkt.c}/100*{Days}/{DaysYear.c}),0)", fmt="0.000000")
    B.tr(S, "SwapPV", "PV of swap net receipts (market less fixed)", "ARD m",
         "{Flag_SwapMTM}*({SwapMkt.c}-{Swap.c})/100*{Bank_Split.c}*{OSB}*{Days}/{DaysYear.c}*{SwapDF}", total=True)
    B.sc(S, "SwapMTMv", "Swap mark-to-market at 2023-12-31 (owed to the concessionaire)", "ARD m", "SUM({SwapPV.r})")
    B.tr(S, "ClaimsGross", "Senior claims: principal plus arrears", "ARD m",
         "{Flag_Restr}*({BankPre}+{BondPre}+{ArrBankPre}+{ArrBondPre})", total=True)
    B.tr(S, "SwapSet", "Swap termination value set off", "ARD m", "{Flag_Restr}*{SwapMTMv.c}", total=True)
    B.tr(S, "ClaimsNet", "Senior claims net of set-off", "ARD m", "{ClaimsGross}-{SwapSet}", total=True)
    B.tr(S, "NotesIssue", "Restructured Senior Notes issued (76%)", "ARD m", "{NotesShare.c}*{ClaimsNet}", total=True)
    B.tr(S, "ConvEq", "Claims converted to equity (10%)", "ARD m", "{Convert.c}*{ClaimsNet}", total=True)
    B.tr(S, "Cancelled", "Claims cancelled (14%)", "ARD m", "{Cancel.c}*{ClaimsNet}", total=True)
    B.tr(S, "SHLwo", "Shareholder loans written off", "ARD m", "{Flag_Restr}*{SHLPre}", total=True)
    B.sc(S, "NotesFace", "Notes face value", "ARD m", "SUM({NotesIssue.r})")
    B.head(S, "Restructured notes sculpting (live in Scenario 5)")
    B.tr(S, "NDF", "Discount factor at the notes rate", "factor", "IF({Flag_NotesRep}=1,IF({Flag_NotesRep.p}=1,{NDF.p},1)/(1+{RRN}),0)", fmt="0.000000")
    B.tr(S, "NSBOpen", "Scheduled notes balance: opening", "ARD m", "IF({Flag_NotesFirst}=1,{NotesFace.c},{NSBClose.p})")
    B.tr(S, "NSInt", "Scheduled notes interest", "ARD m", "{NSBOpen}*{RRN}", total=True)
    B.tr(S, "NSA", "Flag_InterestOnly (notes)", "flag", "IF(AND({Flag_NotesRep}=1,{NSInt}>{CFADSL}/{SNdiv.c}),1,0)")
    B.tr(S, "NSDS", "Scheduled notes debt service", "ARD m", "IF({Flag_NotesRep}=1,MAX({NSInt},{CFADSL}/{SNdiv.c}),0)", total=True)
    B.tr(S, "NSP", "Scheduled notes principal (live)", "ARD m", "IF({Flag_NotesRep}=1,{NSDS}-{NSInt},0)", total=True)
    B.tr(S, "NSBClose", "Scheduled notes balance: closing", "ARD m", "{NSBOpen}-{NSP}")
    B.sc(S, "SNdiv", "Notes sculpting divisor (fixed point)", "x",
         "IF(AND({LiveROn.c}=1,{NotesFace.c}>0),SUMPRODUCT({CFADSL.r}*{NDF.r}*{Flag_NotesRep.r}*(1-{NSA.r}))/({NotesFace.c}-SUMPRODUCT({NSInt.r}*{NDF.r}*{NSA.r})),{NotesDSCR.c})", fmt="0.000000")
    B.tr(S, "NotesPApp", "Scheduled notes principal (applied)", "ARD m", "IF({LiveROn.c}=1,{NSP},{Locked_NotesP})", total=True)
    B.head(S, "Restructured Senior Notes")
    B.tr(S, "NotesOpen", "Opening balance", "ARD m", "{NotesClose.p}")
    B.tr(S, "NotesClose", "Closing balance", "ARD m", "IF({Flag_Restr}=1,{NotesIssue},{NotesOpen}-{NotesPPaid}-{Sweep2})")
    B.head(S, "NILO loan")
    B.tr(S, "NiloOpen", "Opening balance", "ARD m", "{NiloClose.p}")
    B.tr(S, "NiloAccr", "Interest accrued (cash and capitalized)", "ARD m", "{NiloOpen}*{NRate}*{Days}/{DaysYear.c}", total=True)
    B.tr(S, "NiloCashDue", "Cash interest due", "ARD m", "{NiloAccr}*{NCash}*(1-{Flag_Construction})", total=True)
    B.tr(S, "NiloPDue", "Principal due", "ARD m", "MIN({nPApp},{NiloOpen}+{NiloAccr}-{NiloCashDue})*(1-{Flag_Construction})", total=True)
    B.tr(S, "NiloDue", "Total due incl. shortfall brought forward", "ARD m", "{NiloCashDue}+{NiloPDue}+{NiloShort.p}", total=True)
    B.tr(S, "NiloShort", "Shortfall carried forward", "ARD m", "{NiloDue}-{NiloPaid}")
    B.tr(S, "NiloClose", "Closing balance", "ARD m",
         "IF({Flag_Construction}=1,{NiloOpen}+{NiloDraw}+{NiloAccr},{NiloOpen}+{NiloAccr}-{NiloPaid})")
    B.head(S, "NILO sculpting on CFADS after senior debt service (live in Scenarios 2 and 5)")
    B.tr(S, "Resid", "CFADS after scheduled senior or notes debt service", "ARD m",
         "{CFADSL}-IF({Ext.c}=1,{NSDS},{SDS}-{RIA})", total=True)
    B.tr(S, "KDF", "Discount factor at the NILO rate", "factor",
         "IF({Flag_NiloRep}=1,IF({Flag_NiloRep.p}=1,{KDF.p},1)/(1+{NRate}*{Days}/{DaysYear.c}),0)", fmt="0.000000")
    B.tr(S, "KBOpen", "Scheduled NILO balance: opening", "ARD m", "IF({Flag_NiloStart}=1,{NiloClose.p},{KBClose.p})")
    B.tr(S, "KInt", "Scheduled interest accrued", "ARD m", "{KBOpen}*{NRate}*{Days}/{DaysYear.c}", total=True)
    B.tr(S, "KCash", "Scheduled cash interest", "ARD m", "{KInt}*{NCash}", total=True)
    B.tr(S, "KA", "Flag_InterestOnly (NILO)", "flag", "IF(AND({Flag_NiloRep}=1,{KCash}>{Resid}/{Kdiv.c}),1,0)")
    B.tr(S, "KDS", "Scheduled NILO debt service", "ARD m", "IF({Flag_NiloRep}=1,MAX({KCash},{Resid}/{Kdiv.c}),0)", total=True)
    B.tr(S, "KP", "Scheduled NILO principal (live)", "ARD m", "IF({Flag_NiloRep}=1,{KDS}-{KCash},0)", total=True)
    B.tr(S, "KBClose", "Scheduled NILO balance: closing", "ARD m", "{KBOpen}+{KInt}-{KCash}-{KP}")
    B.sc(S, "KB0", "Scheduled balance at the first NILO repayment period", "ARD m", "SUMPRODUCT({KBOpen.r},{Flag_NiloFirst.r})")
    B.sc(S, "KOn", "NILO live sculpting active", "flag", "IF({Ext.c}=1,{LiveROn.c},{LiveOn.c})")
    B.sc(S, "Kdiv", "NILO sculpting divisor (fixed point)", "x",
         "IF({KOn.c}=1,SUMPRODUCT({Resid.r}*{KDF.r}*{Flag_NiloRep.r}*(1-{KA.r}))/({KB0.c}-SUMPRODUCT({KCash.r}*{KDF.r}*{KA.r})),{DSCR_T.c})", fmt="0.000000")
    B.tr(S, "nPApp", "Scheduled NILO principal (applied)", "ARD m",
         "IF({KOn.c}=1,{KP},IF({Ext.c}=1,{Locked_nPr},{Locked_nP}))", total=True)

    # ===================================================================== RESERVES
    S = "Reserves"
    B.head(S, "Debt service reserve account (six months of scheduled senior cash debt service)")
    B.tr(S, "NextDS", "Next period scheduled cash debt service", "ARD m",
         "IF({Months.n}=0,0,({BankOpen}+{BankDraw}-{PPaidBank})*{BankRate.n}*{Days.n}/{DaysYear.c}"
         "+({BondOpen}-{PPaidBond})*{BondCpn.c}*{D360.n}/{Days360Y.c}+({NotesOpen}-{NotesPPaid})*{RRN.n}"
         "+IF({Flag_Standstill.n}=1,0,IF({Post.n}=1,MIN({NotesPApp.n},{NotesOpen}-{NotesPPaid}),"
         "MIN({Bank_Split.c}*{Psched.n},{BankOpen}+{BankDraw}-{PPaidBank})+MIN({Bond_Split.c}*{Psched.n},{BondOpen}-{PPaidBond})))-{RIA.n})")
    B.tr(S, "DSRATarget", "DSRA target", "ARD m",
         "IF({Flag_Construction}=1,IF({End}={Con_End.c},{NextDS}*{DSRA_Months.c}/{Months.n},0),IF({Flag_Restr}=1,0,IF({Months.n}=0,0,{NextDS}*{DSRA_Months.c}/{Months.n})))")
    B.tr(S, "DSRAOpen", "DSRA: opening", "ARD m", "{DSRAClose.p}")
    B.tr(S, "DSRAAvail", "DSRA after draws", "ARD m", "{DSRAOpen}-{DSRADraw}")
    B.tr(S, "DSRAGap", "Shortfall against target", "ARD m", "{DSRATarget}-{DSRAAvail}")
    B.tr(S, "DSRATop", "Top-up from cash flow", "ARD m", "MIN(MAX(0,{A2}),MAX(0,{DSRAGap}))*(1-{Flag_Construction})", total=True)
    B.tr(S, "DSRARel", "Release of excess", "ARD m", "MAX(0,-{DSRAGap})*(1-{Flag_Construction})", total=True)
    B.tr(S, "DSRAClose", "DSRA: closing", "ARD m",
         "IF({Flag_Construction}=1,{DSRAOpen}+{DSRAInit},{DSRAAvail}+{DSRATop}-{DSRARel})")
    B.head(S, "Ramp-up interest account")
    B.tr(S, "RIAOpen", "Opening", "ARD m", "{RIAClose.p}")
    B.tr(S, "RIAClose", "Closing", "ARD m", "{RIAOpen}+{RIAInit}-{RIA}")
    B.head(S, "Lifecycle reserve (MMRA)")
    B.tr(S, "MMRAOpen", "Opening", "ARD m", "{MMRAClose.p}")
    B.tr(S, "MMRAClose", "Closing", "ARD m", "{MMRAOpen}+{MMRAc}-{LC}")
    B.head(S, "Handback reserve")
    B.tr(S, "HBOpen", "Opening", "ARD m", "{HBClose.p}")
    B.tr(S, "HBSpend", "Handback works at expiry", "ARD m", "({HBOpen}+{HBc})*{Flag_Final}", total=True)
    B.tr(S, "HBClose", "Closing", "ARD m", "{HBOpen}+{HBc}-{HBSpend}")
    B.head(S, "Retention account")
    B.tr(S, "RetOpen", "Opening", "ARD m", "{RetClose.p}")
    B.tr(S, "RetPaid", "Released to the D&C contractor", "ARD m", "{RetOpen}*{Flag_RetRel}", total=True)
    B.tr(S, "RetClose", "Closing", "ARD m", "{RetOpen}+{RetDep}-{RetPaid}")
    B.head(S, "Interchange upgrade account (state new money)")
    B.tr(S, "UpgOpen", "Opening", "ARD m", "{UpgClose.p}")
    B.tr(S, "UpgClose", "Closing", "ARD m", "{UpgOpen}-{UpgCapex}+{Flag_Restr}*{Upgrade.c}")

    # ===================================================================== WATERFALL
    S = "Waterfall"
    B.head(S, "Cash available")
    B.tr(S, "CashOpen", "Cash account (locked-up cash) brought forward", "ARD m", "{CashClose.p}")
    B.tr(S, "CFADSW", "CFADS", "ARD m", "{CFADS}", total=True)
    B.tr(S, "RIAW", "Ramp-up interest account release", "ARD m", "{RIA}", total=True)
    B.tr(S, "Support", "Sponsor support loans", "ARD m", "{Hist.c}*IFERROR(INDEX({SUP_Val.r},MATCH({Code},{SUP_Code.r},0)),0)", total=True)
    B.tr(S, "A1", "Cash available for senior debt service", "ARD m", "({CashOpen}+{CFADSW}+{RIAW}+{Support})*(1-{Flag_Construction})")
    B.head(S, "1. Senior fees, interest (including arrears) and principal")
    B.tr(S, "Arrears", "Senior arrears brought forward", "ARD m", "{ArrBankOpen}+{ArrBondOpen}")
    B.tr(S, "SenDue", "Senior amounts due", "ARD m", "({FeesDue}+{IntDue}+{Arrears}+{PDue})*(1-{Flag_Construction})", total=True)
    B.tr(S, "SenPaid", "Senior amounts paid (cash plus DSRA)", "ARD m", "MIN({SenDue},MAX(0,{A1})+{DSRAOpen})*(1-{Flag_Construction})", total=True)
    B.tr(S, "FeesPaid", "Fees paid", "ARD m", "MIN({FeesDue},{SenPaid})*(1-{Flag_Construction})", total=True)
    B.tr(S, "IntPaid", "Interest paid", "ARD m", "MIN({IntDue}+{Arrears},{SenPaid}-{FeesPaid})", total=True)
    B.tr(S, "PPaid", "Principal paid", "ARD m", "{SenPaid}-{FeesPaid}-{IntPaid}", total=True)
    B.tr(S, "ShB", "Bank share of interest paid", "factor", "IF({IntDue}+{Arrears}>0,({BankIntDue}+{ArrBankOpen})/({IntDue}+{Arrears}),0)", fmt="0.0000")
    B.tr(S, "ShO", "Bond share of interest paid", "factor", "IF({IntDue}+{Arrears}>0,({BondIntDue}+{ArrBondOpen})/({IntDue}+{Arrears}),0)", fmt="0.0000")
    B.tr(S, "PPaidBank", "Bank principal paid", "ARD m", "IF({PDue}>0,{PPaid}*{PDueBank}/{PDue},0)", total=True)
    B.tr(S, "PPaidBond", "Bond principal paid", "ARD m", "IF({PDue}>0,{PPaid}*{PDueBond}/{PDue},0)", total=True)
    B.tr(S, "NotesPPaid", "Notes principal paid", "ARD m", "IF({PDue}>0,{PPaid}*{NotesPDue}/{PDue},0)", total=True)
    B.tr(S, "DSRADraw", "DSRA drawn", "ARD m", "MAX(0,{SenPaid}-MAX(0,{A1}))", total=True)
    B.tr(S, "A2", "Cash after senior debt service", "ARD m", "{A1}-{SenPaid}+{DSRADraw}")
    B.head(S, "2. DSRA top-up or release")
    B.tr(S, "A3", "Cash after DSRA", "ARD m", "{A2}-{DSRATop}+{DSRARel}")
    B.head(S, "3. Standstill cash sweep (100% to senior principal)")
    B.tr(S, "BankAfter", "Bank balance after scheduled principal", "ARD m", "{BankOpen}-{PPaidBank}")
    B.tr(S, "BondAfter", "Bond balance after scheduled principal", "ARD m", "{BondOpen}-{PPaidBond}")
    B.tr(S, "Sweep1", "Standstill sweep", "ARD m", "{Flag_Standstill}*MIN(MAX(0,{A3}),{BankAfter}+{BondAfter})", total=True)
    B.tr(S, "Sweep1Bank", "Standstill sweep to bank", "ARD m", "IF({Sweep1}>0,{Sweep1}*{BankAfter}/({BankAfter}+{BondAfter}),0)", total=True)
    B.tr(S, "Sweep1Bond", "Standstill sweep to bonds", "ARD m", "{Sweep1}-{Sweep1Bank}", total=True)
    B.tr(S, "A4", "Cash after standstill sweep", "ARD m", "{A3}-{Sweep1}")
    B.head(S, "4. NILO debt service")
    B.tr(S, "NiloPaid", "NILO paid", "ARD m", "MIN({NiloDue},MAX(0,{A4}))*(1-{Flag_Construction})", total=True)
    B.tr(S, "A5", "Cash after NILO", "ARD m", "{A4}-{NiloPaid}")
    B.head(S, "5. Notes cash sweep (50% to 2030)")
    B.tr(S, "Sweep2", "Notes cash sweep", "ARD m", "{Flag_NotesSweep}*MIN({NotesSweep.c}*MAX(0,{A5}),{NotesOpen}-{NotesPPaid})", total=True)
    B.tr(S, "A6", "Cash available for distribution", "ARD m", "{A5}-{Sweep2}")
    B.head(S, "6. Distribution tests")
    B.tr(S, "DSCRp", "Senior DSCR, period", "x", "IF({SchedDS}>{Tol.c},{CFADSW}/{SchedDS},0)")
    B.tr(S, "HNum", "Historic CFADS (12 months; 6 months at the first test)", "ARD m",
         "IF(OR({Flag_NotesFirst}=1,{Flag_TestDate.p}=0),{CFADSW},{CFADSW}+{CFADSW.p})")
    B.tr(S, "HDen", "Historic scheduled debt service", "ARD m",
         "IF(OR({Flag_NotesFirst}=1,{Flag_TestDate.p}=0),{SchedDS},{SchedDS}+{SchedDS.p})")
    B.tr(S, "DSCRH", "Senior DSCR, historic (covenant)", "x", "IF(AND({Flag_TestDate}=1,{HDen}>{Tol.c}),{HNum}/{HDen},0)")
    B.tr(S, "LockFail", "Lock-up test failed", "flag", "IF(AND({Flag_TestDate}=1,{HDen}>{Tol.c},{DSCRH}<{LockUp.c}),1,0)")
    B.tr(S, "DefFail", "Default test failed", "flag", "IF(AND({Flag_EoDTest}=1,{HDen}>{Tol.c},{DSCRH}<{DefDSCR.c}),1,0)")
    B.tr(S, "EoD", "Event of default continuing", "flag",
         "IF(AND({Flag_Construction}=0,{Post}=0,OR({EoD.p}=1,{DefFail}=1)),1,0)")
    B.tr(S, "Lock", "Distributions blocked", "flag",
         "IF(OR({Flag_Construction}=1,{Flag_Final}=1),0,IF(OR({Flag_PreFirst}=1,{LockFail}=1,{DSRAClose}<{DSRATarget}-{Tol.c},{EoD}=1,"
         "{Flag_Standstill}=1,{NiloShort}>{Tol.c},{ArrBankPre}+{ArrBondPre}>{Tol.c}),1,0))")
    B.tr(S, "Distr", "Distributions to equity", "ARD m", "IF({Lock}=1,0,MAX(0,{A6}))*(1-{Flag_Construction})", total=True)
    B.tr(S, "StateCash", "State reserve top-up received", "ARD m", "{StateRes.c}*{Flag_Restr}", total=True)
    B.tr(S, "CashClose", "Cash account carried forward", "ARD m", "({A6}-{Distr}+{StateCash})*(1-{Flag_Construction})")
    B.head(S, "Shareholder loans and dividends")
    B.tr(S, "SHLOpen", "Shareholder loans: opening (incl. capitalized interest)", "ARD m", "{SHLClose.p}")
    B.tr(S, "SHLInt", "Shareholder loan interest accrued", "ARD m", "{SHLOpen}*{SHL_Rate.c}*{Days}/{DaysYear.c}", total=True)
    B.tr(S, "SHLAvail", "Shareholder loans outstanding before payment", "ARD m", "{SHLOpen}+{SHLInt}+{Support}+{SHLDraw}")
    B.tr(S, "SHLPaid", "Paid on shareholder loans", "ARD m", "MIN({Distr},{SHLAvail})", total=True)
    B.tr(S, "Div", "Dividends", "ARD m", "{Distr}-{SHLPaid}", total=True)
    B.tr(S, "SHLPre", "Shareholder loans: closing before restructuring", "ARD m", "{SHLAvail}-{SHLPaid}")
    B.tr(S, "SHLClose", "Shareholder loans: closing", "ARD m", "IF({Flag_Restr}=1,0,{SHLPre})")
    B.head(S, "Equity cash flows")
    B.tr(S, "EqCF", "Original sponsors' cash flow", "ARD m",
         "IF({Flag_Construction}=1,-{EqDraw},IF({Post}=1,0,{Distr}-{Support}))", total=True)
    B.tr(S, "NewEqCF", "New equity cash flow (creditors 85%, state 15%)", "ARD m",
         "IF({Post}=1,{Distr},0)-{StateMoney.c}*{Flag_Restr}", total=True)

    # ===================================================================== TAX
    S = "Tax"
    B.head(S, "Tax cost base of the concession asset")
    B.tr(S, "WDVOpen", "Opening tax written-down value", "ARD m", "{WDVClose.p}")
    B.tr(S, "WDVAdd", "Additions (capitalized cost less contribution plus upgrade)", "ARD m", "{CapCost}-{Contribution}+{UpgCapex}", total=True)
    B.tr(S, "Amort", "Tax amortization (straight line over remaining term)", "ARD m", "IF({RemDays}>0,{WDVOpen}*{OpsDays}/{RemDays},0)", total=True)
    B.tr(S, "WPool", "Value before forgiveness", "ARD m", "{WDVOpen}+{WDVAdd}-{Amort}")
    B.tr(S, "ForgWDV", "Debt forgiveness applied to cost base", "ARD m", "MIN({Forg}-{ForgLoss},{WPool})", total=True)
    B.tr(S, "WDVClose", "Closing tax written-down value", "ARD m", "{WPool}-{ForgWDV}")
    B.head(S, "Taxable income")
    B.tr(S, "NiloIntExp", "NILO interest accrued (deductible)", "ARD m", "{NiloAccr}*(1-{Flag_Construction})", total=True)
    B.tr(S, "SHLExp", "Shareholder loan interest accrued (deductible)", "ARD m", "{SHLInt}*(1-{Flag_Construction})", total=True)
    B.tr(S, "TI", "Taxable income before losses", "ARD m",
         "({EBITDA}-{LC}-{HBSpend}-{Amort}-{IntDue}-{NiloIntExp}-{SHLExp}-{FeesDue})*(1-{Flag_Construction})", total=True)
    B.tr(S, "LossOpen", "Tax losses brought forward", "ARD m", "{LossClose.p}")
    B.tr(S, "LossUsed", "Losses used", "ARD m", "MIN({LossOpen},MAX(0,{TI}))", total=True)
    B.tr(S, "LossAdd", "Losses added", "ARD m", "MAX(0,-{TI})", total=True)
    B.tr(S, "Forg", "Forgiven commercial debt (cancelled claims plus shareholder loans)", "ARD m", "{Flag_Restr}*({Cancelled}+{SHLwo})", total=True)
    B.tr(S, "ForgLoss", "Forgiveness applied to losses", "ARD m", "MIN({Forg},{LossOpen}-{LossUsed}+{LossAdd})", total=True)
    B.tr(S, "LossClose", "Tax losses carried forward", "ARD m", "{LossOpen}-{LossUsed}+{LossAdd}-{ForgLoss}")
    B.tr(S, "Taxable", "Taxable income after losses", "ARD m", "MAX(0,{TI})-{LossUsed}", total=True)
    B.tr(S, "Tax", "Tax paid", "ARD m", "{TaxRate.c}*{Taxable}", total=True)

    # ===================================================================== FINANCIALS
    S = "Financials"
    B.head(S, "Income statement")
    B.tr(S, "BAmort", "Book amortization of the concession asset", "ARD m", "IF({RemDays}>0,{NBV.p}*{OpsDays}/{RemDays},0)", total=True)
    B.tr(S, "NBV", "Concession intangible: net book value", "ARD m", "{NBV.p}+{WDVAdd}-{BAmort}")
    B.tr(S, "Gain", "Restructuring gain (cancelled claims and swap termination)", "ARD m", "{Cancelled}+{SwapSet}", total=True)
    B.tr(S, "NI", "Net income", "ARD m",
         "({EBITDA}-{LC}-{HBSpend}-{BAmort}-{IntDue}-{NiloIntExp}-{SHLExp}-{FeesDue}-{Tax})*(1-{Flag_Construction})+{Gain}", total=True)
    B.head(S, "Balance sheet")
    B.tr(S, "TA", "Total assets", "ARD m",
         "{NBV}+{Rec}+{EscrowClose}+{DSRAClose}+{RIAClose}+{MMRAClose}+{HBClose}+{RetClose}+{UpgClose}+{CashClose}")
    B.tr(S, "TL", "Total liabilities", "ARD m",
         "{BankClose}+{BondClose}+{ArrBankClose}+{ArrBondClose}+{NotesClose}+{NiloClose}+{BridgeClose}+{Pay}+{SHLClose}+{RetClose}")
    B.tr(S, "ShCap", "Share capital", "ARD m", "{ShCap.p}+{ShareCapDraw}+{ConvEq}+{StateMoney.c}*{Flag_Restr}")
    B.tr(S, "SHLRes", "Shareholder loans forgiven (equity reserve)", "ARD m", "{SHLRes.p}+{SHLwo}")
    B.tr(S, "RE", "Retained earnings", "ARD m", "{RE.p}+{NI}-{Div}")
    B.tr(S, "TE", "Total equity", "ARD m", "{ShCap}+{SHLRes}+{RE}")
    B.tr(S, "BSCheck", "Balance check (assets less liabilities less equity)", "ARD m", "ROUND({TA}-{TL}-{TE},6)")

    # ===================================================================== RATIOS
    S = "Ratios"
    B.head(S, "Senior ratios (original debt, repayment periods)")
    B.tr(S, "DSCRRep", "DSCR, repayment periods", "x", "IF(AND({Flag_Rep}=1,{SchedDS}>{Tol.c}),{CFADSW}/{SchedDS},\"\")")
    B.sc(S, "MinDSCR", "Minimum DSCR, repayment periods", "x", "MIN({DSCRRep.r})")
    B.sc(S, "AvgDSCR", "Average DSCR, repayment periods", "x", "AVERAGE({DSCRRep.r})")
    B.tr(S, "DSCRRamp", "DSCR, ramp-up periods", "x", "IF({Flag_RampUp}=1,{CFADSW}/{SchedDS},\"\")")
    B.sc(S, "MinDSCRRamp", "Minimum DSCR, ramp-up", "x", "MIN({DSCRRamp.r})")
    B.tr(S, "DSCRComb", "Senior plus NILO DSCR", "x",
         "IF(AND({Flag_Rep}=1,{Flag_NiloRep}=1),{CFADSW}/({SchedDS}+{NiloCashDue}+{NiloPDue}),\"\")")
    B.sc(S, "MinComb", "Minimum senior plus NILO DSCR", "x", "MIN({DSCRComb.r})")
    B.tr(S, "DFAll", "Discount factor from 2021-01-01 (all periods)", "factor",
         "IF({End}>={First_Rep.c},IF({DFAll.p}>0,{DFAll.p},1)/(1+{RR}),0)", fmt="0.000000")
    B.tr(S, "LLCR", "LLCR", "x",
         "IF(AND({Flag_Rep}=1,{BankOpen}+{BondOpen}>{Tol.c}),SUMPRODUCT({CFADSW.rt},{DFAll.rt},{Flag_Rep.rt})/({DFAll}*(1+{RR}))/({BankOpen}+{BondOpen}),\"\")")
    B.sc(S, "LLCR1", "LLCR at the first repayment period", "x", "SUMPRODUCT({LLCR.r},{Flag_FirstRep.r})")
    B.sc(S, "MinLLCR", "Minimum LLCR", "x", "MIN({LLCR.r})")
    B.sc(S, "PLCR1", "PLCR at the first repayment period", "x",
         "SUMPRODUCT({CFADSW.r},{DFAll.r},{Flag_Ops.r})/SUMPRODUCT({BankOpen.r}+{BondOpen.r},{Flag_FirstRep.r})")
    B.tr(S, "DSCRNotes", "Notes DSCR", "x", "IF(AND({Flag_NotesRep}=1,{SchedDS}>{Tol.c}),{CFADSW}/{SchedDS},\"\")")
    B.sc(S, "MinNotes", "Minimum notes DSCR", "x", "IF(COUNT({DSCRNotes.r})>0,MIN({DSCRNotes.r}),0)")
    B.sc(S, "AvgNotes", "Average notes DSCR", "x", "IF(COUNT({DSCRNotes.r})>0,AVERAGE({DSCRNotes.r}),0)")
    B.sc(S, "LockPeriods", "Lock-up periods after the first repayment date", "#",
         "SUMPRODUCT({Lock.r},{Flag_Ops.r},1-{Flag_PreFirst.r})", fmt="0")
    B.sc(S, "DSRADraws", "DSRA drawn in total", "ARD m", "SUM({DSRADraw.r})")
    B.sc(S, "GearSen", "Senior debt, % of funding net of contribution", "%", "{Dsen.c}/{Fnet.c}")
    B.sc(S, "GearNilo", "NILO, % of funding net of contribution", "%", "{Nsub.c}/{Fnet.c}")
    B.sc(S, "GearEq", "Equity, % of funding net of contribution", "%", "{Eq.c}/{Fnet.c}")
    B.sc(S, "NiloElig", "NILO, % of eligible costs", "%", "{Nsub.c}/{TotalUses.c}")

    # ===================================================================== RETURNS
    S = "Returns"
    B.head(S, "Original sponsors' equity")
    B.tr(S, "EqCFR", "Equity cash flow", "ARD m", "{EqCF}", total=True)
    B.sc(S, "EqIRR", "Equity IRR (XIRR on period-end dates)", "%", "IFERROR(XIRR({EqCFR.r},{End.r}),\"n/a\")")
    B.sc(S, "EqNPV", "Equity NPV at 11.4% at financial close", "ARD m",
         "SUMPRODUCT({EqCFR.r}/(1+{Target_IRR.c})^(({End.r}-{FC_Date.c})/{DaysYear.c}))")
    B.sc(S, "EqInv", "Equity invested (incl. sponsor support)", "ARD m", "SUM({EqDraw.r})+SUM({Support.r})")
    B.head(S, "Project cash flow (post-tax, before financing)")
    B.tr(S, "ProjCF", "Project cash flow", "ARD m", "-{Capex}+{Contribution}+{CFADSW}*(1-{Flag_Construction})-{UpgCapex}", total=True)
    B.sc(S, "ProjIRR", "Project IRR, post-tax", "%", "IFERROR(XIRR({ProjCF.r},{End.r}),\"n/a\")")
    B.tr(S, "ProjCFPre", "Project cash flow before tax", "ARD m", "{ProjCF}+{Tax}", total=True)
    B.sc(S, "ProjIRRPre", "Project IRR, pre-tax", "%", "IFERROR(XIRR({ProjCFPre.r},{End.r}),\"n/a\")")
    B.head(S, "Termination values at 2022-06-30")
    B.sc(S, "NPVDistr", "NPV at 11.4% of distributions after 2022-06-30 (authority-default equity compensation when run on the bid base)", "ARD m",
         "SUMPRODUCT(({End.r}>{FV_Date.c})*{Distr.r}/(1+{Target_IRR.c})^(({End.r}-{FV_Date.c})/{DaysYear.c}))")
    B.tr(S, "FVCF", "Pre-tax unlevered cash flow (retender valuation)", "ARD m", "{EBITDA}-{dWC}-{LC}-{HBSpend}+{RestrCost}", total=True)
    B.sc(S, "FV2022", "Estimated fair value at 2022-06-30 (run in Scenario 6)", "ARD m",
         "SUMPRODUCT(({End.r}>{FV_Date.c})*{FVCF.r}/(1+{FV_Rate.c})^(({End.r}-{FV_Date.c})/{DaysYear.c}))")
    B.sc(S, "CDComp", "Concessionaire-default compensation (fair value less retendering costs)", "ARD m", "{FV2022.c}-{Retender_Cost.c}")
    B.sc(S, "SenClaim22", "Senior claims at 2022-06-30 (principal plus arrears)", "ARD m",
         "SUMPRODUCT(({End.r}={FV_Date.c})*({BankClose.r}+{BondClose.r}+{ArrBankClose.r}+{ArrBondClose.r}))")
    B.sc(S, "Nilo22", "NILO outstanding at 2022-06-30", "ARD m", "SUMPRODUCT(({End.r}={FV_Date.c})*{NiloClose.r})")
    B.sc(S, "EqNet22", "Equity contributed less distributions to 2022-06-30", "ARD m",
         "SUMPRODUCT(({End.r}<={FV_Date.c})*({EqDraw.r}+{Support.r}-{Distr.r}))")
    B.sc(S, "FMComp", "Prolonged force majeure compensation", "ARD m", "{SenClaim22.c}+{Nilo22.c}+{EqNet22.c}")
    B.head(S, "Restructuring valuation (run in Scenario 5)")
    B.sc(S, "EqV", "New equity value at the plan rate (14.0%)", "ARD m",
         "SUMPRODUCT(({End.r}>{Restr_Date.c})*{Distr.r}/(1+{EqV_Rate.c})^(({End.r}-{Restr_Date.c})/{DaysYear.c}))")
    B.tr(S, "NotesRepaid", "Flag_NotesRepaid", "flag", "IF(AND({Ext.c}=1,{End}>{Restr_Date.c},{NotesOpen}<={Tol.c}),1,0)")
    B.sc(S, "WarrV", "Warrants: 3% of distributions after the notes are repaid", "ARD m",
         "{Warrant.c}*SUMPRODUCT({NotesRepaid.r}*{Distr.r}/(1+{EqV_Rate.c})^(({End.r}-{Restr_Date.c})/{DaysYear.c}))")
    B.sc(S, "EqVCred", "Equity value to senior creditors (85%)", "ARD m", "{CredEq.c}*({EqV.c}-{WarrV.c})")
    B.sc(S, "EqVState", "Equity value to the state (15%)", "ARD m", "{StateEq.c}*({EqV.c}-{WarrV.c})")
    B.sc(S, "StateGrant", "State money in excess of plan value (implied capital grant)", "ARD m", "{StateMoney.c}-{EqVState.c}")
    B.tr(S, "NotesCF", "Payments to noteholders", "ARD m", "IF({Post}=1,{IntPaid}+{NotesPPaid}+{Sweep2},0)", total=True)
    B.sc(S, "NotesMV", "Notes market value at the 7.50% yield", "ARD m",
         "SUMPRODUCT(({End.r}>{Restr_Date.c})*{NotesCF.r}/(1+{Notes_Yield.c})^(({End.r}-{Restr_Date.c})/{DaysYear.c}))")
    B.sc(S, "ClaimsNetV", "Senior claims net of set-off", "ARD m", "SUM({ClaimsNet.r})")
    B.sc(S, "SenRec", "Senior recovery (notes at market plus equity at plan value)", "%", "IF({ClaimsNetV.c}>0,({NotesMV.c}+{EqVCred.c})/{ClaimsNetV.c},0)")
    B.sc(S, "SenRecPar", "Senior recovery with notes at par", "%", "IF({ClaimsNetV.c}>0,({NotesFace.c}+{EqVCred.c})/{ClaimsNetV.c},0)")
    B.sc(S, "NiloClaim", "NILO claim at 2023-12-31", "ARD m", "SUMPRODUCT({Flag_Restr.r},{NiloClose.r})")
    B.sc(S, "NiloPV", "PV of NILO receipts at 3.06%", "ARD m",
         "SUMPRODUCT(({End.r}>{Restr_Date.c})*{NiloPaid.r}/(1+{Nilo_Rate.c})^(({End.r}-{Restr_Date.c})/{DaysYear.c}))")
    B.sc(S, "NiloRec", "NILO recovery in PV terms", "%", "IF({NiloClaim.c}>0,{NiloPV.c}/{NiloClaim.c},0)")
    B.head(S, "Public sector comparator and value for money (PV at 2012-12-31)")
    B.sc(S, "PSCTot", "Risk-adjusted PSC", "ARD m",
         "{PSC_Capex.c}+{PSC_OM.c}+{PSC_Toll.c}+{PSC_ConRisk.c}+{PSC_TrafRisk.c}+{PSC_OpRisk.c}+{PSC_CN.c}")
    B.sc(S, "PSCRaw", "Raw PSC (before risk and competitive neutrality)", "ARD m", "{PSC_Capex.c}+{PSC_OM.c}+{PSC_Toll.c}")
    B.sc(S, "PSCDF", "Discount factor, contribution date to PV date", "factor",
         "(1+{PSC_Rate.c})^(-({Contrib_Date.c}-{PSC_Date.c})/{DaysYear.c})", fmt="0.000000")
    B.sc(S, "PPPRef", "PPP reference project cost", "ARD m", "{PSC_Ref.c}*{PSCDF.c}+{PSC_Retained.c}+{PSC_CM.c}")
    B.sc(S, "PPPBid", "PPP cost at the winning bid", "ARD m", "{Contrib_Bid.c}*{PSCDF.c}+{PSC_Retained.c}+{PSC_CM.c}")
    B.sc(S, "VfMRef", "Value for money, reference", "ARD m", "{PSCTot.c}-{PPPRef.c}")
    B.sc(S, "VfMRefPct", "Value for money, reference, % of risk-adjusted PSC", "%", "{VfMRef.c}/{PSCTot.c}")
    B.sc(S, "VfMBid", "Value for money, winning bid", "ARD m", "{PSCTot.c}-{PPPBid.c}")
    B.sc(S, "VfMBidPct", "Value for money, winning bid, % of risk-adjusted PSC", "%", "{VfMBid.c}/{PSCTot.c}")

    # ===================================================================== CHECKS
    S = "Checks"
    B.head(S, "Integrity checks (0 = pass)")
    B.sc(S, "CK_SU", "Sources equal uses in every construction period", "ARD m", "ROUND(SUMPRODUCT(ABS({SUCheck.r})),6)")
    B.sc(S, "CK_Fund", "Total sources at close equal total uses", "ARD m",
         "ROUND({Eq.c}+{Dsen.c}+{Nsub.c}+{Contrib.c}-{TotalUses.c},6)")
    B.sc(S, "CK_Eq", "Equity drawn equals equity commitment", "ARD m", "ROUND(SUM({EqDraw.r})-{Eq.c},6)")
    B.sc(S, "CK_Escrow", "Bond escrow fully used at completion", "ARD m", "ROUND(SUMPRODUCT({EscrowClose.r},{Flag_FirstPost.r}),6)")
    B.sc(S, "CK_Bridge", "Bridge repaid at opening", "ARD m", "ROUND(SUMPRODUCT(ABS({BridgeClose.r})*{Flag_Ops.r}),6)")
    B.sc(S, "CK_BS", "Balance sheet balances", "ARD m", "ROUND(SUMPRODUCT(ABS({BSCheck.r})),6)")
    B.sc(S, "CK_Sculpt", "Live senior schedule repays by 2048 (Scenario 2)", "ARD m",
         "IF({LiveOn.c}=1,ROUND(SUMPRODUCT({SBClose.r},{Flag_Rep.r}*({End.r}={Final_Rep.c})),6),0)")
    B.sc(S, "CK_Locked", "Live sizing equals locked financing (Scenario 2, bid contribution)", "ARD m",
         "IF(AND({LiveOn.c}=1,{ContribOpt.c}=1),ROUND(ABS({D_Live.c}-{Locked_D.c})+SUMPRODUCT(ABS({SP.r}-{Locked_P.r}))+SUMPRODUCT(ABS({RIALive.r}-{Locked_RIA.r}))+SUMPRODUCT(ABS({KP.r}-{Locked_nP.r})),3),0)")
    B.sc(S, "CK_Notes", "Live notes schedule equals locked (Scenario 5)", "ARD m",
         "IF({LiveROn.c}=1,ROUND(SUMPRODUCT(ABS({NSP.r}-{Locked_NotesP.r}))+SUMPRODUCT(ABS({KP.r}-{Locked_nPr.r})),3),0)")
    B.sc(S, "CK_SenRepaid", "Senior debt repaid by concession expiry", "ARD m",
         "ROUND(SUMPRODUCT(({BankClose.r}+{BondClose.r}+{NotesClose.r})*{Flag_Final.r}),6)")
    B.sc(S, "CK_NiloRepaid", "NILO repaid by concession expiry", "ARD m", "ROUND(SUMPRODUCT({NiloClose.r}*{Flag_Final.r}),6)")
    B.sc(S, "CK_Res", "Reserves never negative", "ARD m",
         "ROUND(-MIN(0,MIN({DSRAClose.r}),MIN({MMRAClose.r}),MIN({HBClose.r}),MIN({RIAClose.r}),MIN({UpgClose.r})),6)")
    B.sc(S, "CK_RIA", "Ramp-up account fully used", "ARD m", "ROUND(SUMPRODUCT({RIAClose.r},{Flag_FirstRep.r}),6)")
    B.sc(S, "CK_Cash", "Cash account never negative", "ARD m", "ROUND(-MIN(0,MIN({CashClose.r})),6)")
    B.sc(S, "CK_Nilo33", "NILO within 33% of eligible costs", "ARD m", "ROUND(MAX(0,{Nsub.c}-{Nilo_Max.c}*{TotalUses.c}),6)")
    B.sc(S, "CK_Eq22", "Equity at least 22% of funding net of contribution", "ARD m", "ROUND(MAX(0,{EqShare.c}*{Fnet.c}-{Eq.c}),6)")
    B.sc(S, "CK_Restr", "Restructuring: notes plus converted plus cancelled equal net claims", "ARD m",
         "ROUND(SUM({NotesIssue.r})+SUM({ConvEq.r})+SUM({Cancelled.r})-SUM({ClaimsNet.r}),6)")
    B.sc(S, "CK_All", "All checks", "ARD m",
         "ABS({CK_SU.c})+ABS({CK_Fund.c})+ABS({CK_Eq.c})+ABS({CK_Escrow.c})+ABS({CK_Bridge.c})+ABS({CK_BS.c})+ABS({CK_Sculpt.c})+ABS({CK_Locked.c})+ABS({CK_Notes.c})+ABS({CK_SenRepaid.c})+ABS({CK_NiloRepaid.c})+ABS({CK_Res.c})+ABS({CK_RIA.c})+ABS({CK_Cash.c})+ABS({CK_Nilo33.c})+ABS({CK_Eq22.c})+ABS({CK_Restr.c})")

    # ===================================================================== OUTPUTS
    S = "Outputs"
    B.head(S, "Dashboard (selected scenario)")
    for nm, lab, u, f in [
        ("O_Scen", "Scenario", "#", "{Scenario.c}"),
        ("O_TotUses", "Total uses", "ARD m", "{TotalUses.c}"),
        ("O_Fnet", "Funding requirement net of contribution", "ARD m", "{Fnet.c}"),
        ("O_D", "Senior debt (bank plus bonds)", "ARD m", "{Dsen.c}"),
        ("O_Bank", "Bank mini-perm", "ARD m", "{BankC.c}"),
        ("O_Bond", "Revenue bonds", "ARD m", "{BondF.c}"),
        ("O_N", "NILO loan", "ARD m", "{Nsub.c}"),
        ("O_E", "Equity", "ARD m", "{Eq.c}"),
        ("O_Bind", "Binding sizing constraint", "#", "{Binding.c}"),
        ("O_s", "Senior sculpting divisor", "x", "{Sdiv.c}"),
        ("O_IDC", "Senior IDC net of escrow income plus bridge interest", "ARD m",
         "SUM({BankIntC.r})+SUM({BondIntC.r})-SUM({EscrowInt.r})+SUM({BridgeIntC.r})"),
        ("O_Fees", "Upfront and commitment fees", "ARD m", "SUM({Fees.r})+SUM({Commit.r})"),
        ("O_DSRA", "DSRA initial funding", "ARD m", "SUM({DSRAInit.r})"),
        ("O_RIA", "Ramp-up interest account", "ARD m", "SUM({RIAInit.r})"),
        ("O_MinDSCR", "Minimum senior DSCR (repayment periods)", "x", "{MinDSCR.c}"),
        ("O_AvgDSCR", "Average senior DSCR", "x", "{AvgDSCR.c}"),
        ("O_MinComb", "Minimum senior plus NILO DSCR", "x", "{MinComb.c}"),
        ("O_LLCR", "LLCR at first repayment", "x", "{LLCR1.c}"),
        ("O_PLCR", "PLCR at first repayment", "x", "{PLCR1.c}"),
        ("O_EqIRR", "Equity IRR", "%", "{EqIRR.c}"),
        ("O_EqNPV", "Equity NPV at 11.4%", "ARD m", "{EqNPV.c}"),
        ("O_ProjIRR", "Project IRR post-tax", "%", "{ProjIRR.c}"),
        ("O_VfM", "Value for money, winning bid", "ARD m", "{VfMBid.c}"),
        ("O_Checks", "All checks (0 = pass)", "ARD m", "{CK_All.c}"),
    ]:
        B.sc(S, nm, lab, u, f)
    B.write(path, scenario, contribution_option)
    return B


if __name__ == "__main__":
    b = build(os.path.join(HERE, "Case_T_Model.xlsx"))
    print("rows:", len(b.rows))
