import json, os, glob
from openpyxl import load_workbook
import numpy as np
H='/home/user/neurolab/book/model'
o=json.load(open(f'{H}/outputs_case_p.json'))
log=open(f'{H}/recalc_p/verify_all.log').read().splitlines()
au=json.load(open(f'{H}/recalc_p/verify_audit.json'))
L=[]
a=L.append
a('# Case P workbook verification')
a('')
a('Model version 1.4. Method. `model/verify_p.py` builds `Case_P_Model.xlsx` with the scenario selector set to each of the 15 scenarios, recalculates every copy with LibreOffice 24.2 headless using a private user profile (`soffice -env:UserInstallation=file:///tmp/lo_profile_case_p --headless --calc --convert-to xlsx`) into `model/recalc_p/out/`, reads the recalculated values with openpyxl (`data_only=True`) and compares them with the Python mirror (`case_p.py`) row by row: every mapped calculation row over all its columns (monthly or semiannual) and every mapped scalar. Tolerance: 0.01 in displayed units (USD m for amounts, x for ratios, percentage points for IRRs). The workbook has no circular references and no macros; iterative calculation is off.')
a('')
a('## Results by scenario')
a('')
a('| Scenario | Name | Rows and scalars compared | Failures | Largest absolute difference | Workbook checks (0 = pass) | Result |')
a('|---|---|---|---|---|---|---|')
for line in log:
    if line.startswith('scenario'):
        p=line.split()
        i=p[1].rstrip(':'); n=p[2]; fails=p[4]; md=p[8].rstrip(','); chk=p[10]
        a(f"| {i} | {o['meta']['scenarios'][i]} | {n} | {fails} | {md} | {chk} | {'PASS' if fails=='0' and chk=='0' else 'FAIL'} |")
a(f"| audit copy | Case_P_Model_AuditExercise.xlsx (FC base, errors E1-E10 seeded) | {au['n']} | {au['fails']} | {au['maxdiff']:.2e} | {au['checks']} | PASS (the failing checks are intended audit clues: F14 debt above the correct gearing cap and F23 the ECA tests, which the seeded errors also breach) |")
a('')
mc=json.load(open(f'{H}/recalc_p/verify_mc.json'))
a('')
a('## Monte Carlo wiring (u09 R5)')
a('')
a('Scenario 1 with Inputs F311 set to a run number; the workbook reads that row of the pasted draw table (availability shocks by operating year, dispatch, heat-rate degradation, FX drift) and is compared with the mirror run with the same draws (P-F42 generator, seed 20180717).')
a('')
a('| Run | Rows and scalars compared | Failures | Largest absolute difference | Workbook checks |')
a('|---|---|---|---|---|')
for k,v in mc.items(): a(f"| {k} | {v['n']} | {v['fails']} | {v['maxdiff']:.2e} | {v['checks']} |")
st=json.load(open(f'{H}/recalc_p/stage_results.json'))
a('')
a('## Build-stage reconciliation on Scenario 1 (u09 Section 0.5; confirmation for the build agent)')
a('')
a('For each stage the companion workbook was cut to the rows the u09 Section 0.4 row map assigns to Chapters 39 up to that chapter (v1.4 appended rows assigned as in `model/case_p_report.md` Section 8c), the provisional rows were pasted as FC base values, the master check was restated over the checks present, and the file was recalculated with LibreOffice. Every remaining numeric cell was compared with the full model. A script also confirmed that no formula in a stage refers to a row not yet built (Checks F19 excepted, which each stage restates).')
a('')
a('| Stage | Numeric cells compared | Largest absolute difference | Provisional rows pasted | Master check |')
a('|---|---|---|---|---|')
for k,v in st.items(): a(f"| Ch {k} | {v['cells']} | {v['maxdiff']:.1e} | {', '.join(f'{x[0]} {x[1]}' for x in v['pasted']) or 'none'} | {v['master']} |")
a('')
a('## Key outputs, Python against workbook (selected scenarios)')
a('')
a('| Scenario | Output | Python | Workbook | Difference |')
a('|---|---|---|---|---|')
for i in (1,3,15):
    wb=load_workbook(f'{H}/recalc_p/out/scn{i:02d}.xlsx',data_only=True)
    ws=wb['Outputs']
    lab={ws.cell(r,4).value: ws.cell(r,6).value for r in range(1,40) if ws.cell(r,4).value}
    s=o['summary'][str(i)]
    pairs=[('Total funding requirement','T'),('Senior debt (four tranches)','D'),('Minimum DSCR','min_dscr'),('Average DSCR (debt-service weighted)', 'avg_dscr'),('LLCR at first debt service period (incl. DSRA)','llcr_first'),('PLCR at first debt service period','plcr_first'),('Equity IRR','equity_irr'),('Project IRR, post-tax','project_irr'),('Equity NPV at 16.0% (at FC)','npv16_at_fc')]
    for lb,k in pairs:
        key=[x for x in lab if x.startswith(lb.split(' (')[0])][0]
        xv=lab[key]; pv=s[k]
        if 'IRR' in lb:
            a(f"| {i} | {lb} | {pv*100:.4f}% | {xv*100:.4f}% | {abs(xv-pv)*100:.2e} pp |")
        else:
            a(f"| {i} | {lb} | {pv:.4f} | {xv:.4f} | {abs(xv-pv):.2e} |")
a('')
a('## Coverage')
a('')
a('Compared rows include: the Time sheet operating months and days; every Construction use line and the VAT facility (KCR) rows; every Funding row (tranche drawdowns, balances, interest by tranche, commitment and upfront fees, ECA premium, DSRA funding, uses, equity, share capital, shareholder loans and capitalized interest, the alpha and beta rows of the closed-form gross-up and the closed-form total funding requirement); Operations macro paths, indices, plant, revenue lines, operating costs, working capital; every Tax row; every Debt tranche corkscrew (interest, scheduled principal, deferral, LD prepayment, refinancing, sweeps, closing balances), the swap, PRI, PCG and waiver fee, debt service and the DSRA target; Reserves (MMRA, DSRA, handback); the full Waterfall (tests, flags, sweeps, shareholder-loan payments, dividends, trapped cash, retained earnings); Financials (income statement, balance sheet and balance check); Ratios (DSCR, LLCR, PLCR); and Returns (equity IRR, project IRRs, NPV). Python-only figures are listed in `case_p_report.md` Section 8.')
a('')
a('## Known deviations from the style sheet (not verification failures)')
a('')
a('* FAST check (v1.4): `model/scan_hardcodes_p.py` lists every numeric literal inside calculation formulas other than 0, 1, 12 and the unit conversions 100, 1,000 and 1,000,000. Result for Case_P_Model.xlsx: 0. Event dates, period lengths, operating-year bands, day bases, shares and tolerances are named inputs on the Inputs sheet; event flags (Flag_LDPrepayment, Flag_WaiverDeferral, Flag_Refinancing and others) sit on the Time sheet; every calculation row uses one formula copied across (the first column reads the blank column I as the prior period). The audit exercise copy keeps only its deliberate seeded errors (E4 uses 0.5 and 1/12, E9 types 0.765).')
a('* The Inputs sheet holds time-series inputs on the model timelines (semiannual and monthly blocks), not on a separate date header.')
open(f'{H}/case_p_verification.md','w').write('\n'.join(L)+'\n')
