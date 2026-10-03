#!/usr/bin/env python3
"""Build ex43_17_solution.xlsx: the solution workbook of Exercise 43.17 (Cactus Ridge Solar and Storage,
Illustrative), live formulas, macro-free, no circular reference, no iterative calculation.

Circularity (debt -> IDC, fees, DSRA -> total funding; debt -> interest -> tax -> CFADS -> sculpted debt)
is resolved as in the Case P companion (u09 Section 0.6): the converged senior debt, equity commitment
and principal profile from ex43_17_mirror.py are contract inputs; the workbook recomputes the funding
and the sculpted profile live, and the Checks sheet proves they are the fixed point.
"""
import os, sys
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.formatting.rule import CellIsRule
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils import get_column_letter as L

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import ex43_17_mirror as mir      # noqa: E402

BLUE = Font(color='0000FF'); GREEN = Font(color='008000'); BOLD = Font(bold=True)
YEL = PatternFill('solid', fgColor='FFF2CC'); RED = PatternFill('solid', fgColor='FF9999'); HEAD = PatternFill('solid', fgColor='DDEBF7')
NM, NS = 12, 40
J = 10
cm = lambda i: L(J + i)                     # column of month/period i (0-based)
LASTM, LASTS = cm(NM - 1), cm(NS - 1)       # U, AW
NR = NM + NS                                # Returns strip
LASTR = cm(NR - 1)


def build(path):
    out, sol, cases = mir.answers()
    wb = Workbook(); wb.remove(wb.active)
    sheets = {}
    def sheet(name, title, tl=None):
        ws = wb.create_sheet(name); sheets[name] = ws
        ws['A1'] = 'Cactus Ridge Solar and Storage (fictional; Illustrative): Exercise 43.17 solution'; ws['A1'].font = Font(bold=True, size=13)
        ws['A2'] = title
        ws['A3'] = 'Column D label, E unit, F constant or scalar, G row total; first period in column J. Inputs: blue on yellow. Links from other sheets: green. Checks: 0 when passing.'
        for col, w in zip('ABCDEFGHI', [2, 2, 2, 52, 12, 14, 12, 2, 2]):
            ws.column_dimensions[col].width = w
        for c, t in zip('DEFG', ['Label', 'Unit', 'Constant', 'Total']):
            ws[f'{c}4'] = t; ws[f'{c}4'].font = BOLD
        if tl:
            labels = {'M': [date(2027, m, 1).strftime('%Y-%m') for m in range(1, 13)],
                      'S': [f'{2028 + k // 2}H{1 + k % 2}' for k in range(NS)],
                      'R': [date(2027, m, 1).strftime('%Y-%m') for m in range(1, 13)] + [f'{2028 + k // 2}H{1 + k % 2}' for k in range(NS)]}[tl]
            for i, t in enumerate(labels):
                c = ws.cell(5, J + i, t); c.font = BOLD; c.fill = HEAD
                ws.column_dimensions[cm(i)].width = 10
            ws.freeze_panes = 'J6'
        return ws

    def lab(ws, r, label, unit, sec=False):
        if sec:
            ws.cell(r, 2, label).font = BOLD
        else:
            ws.cell(r, 4, label); ws.cell(r, 5, unit)

    def inp(ws, ref, v, fmt=None):
        c = ws[ref]; c.value = v; c.font = BLUE; c.fill = YEL
        if fmt: c.number_format = fmt

    def row(ws, r, n, fn, fmt='#,##0.000;(#,##0.000);"-"', total=False):
        for i in range(n):
            f = fn(i); c = ws.cell(r, J + i, f); c.number_format = fmt
            if f.startswith('=') and f.count('!') == 1 and not any(o in f[1:] for o in '+-*/(,'):
                c.font = GREEN
        if total:
            ws.cell(r, 7, f'=SUM(J{r}:{cm(n - 1)}{r})').number_format = fmt

    def sc(ws, r, f, fmt='#,##0.000'):
        c = ws.cell(r, 6, f); c.number_format = fmt
        if f.startswith('=') and f.count('!') == 1 and not any(o in f[1:] for o in '+-*/(,'):
            c.font = GREEN

    # ------------------------------------------------------------------ Cover
    cv = sheet('Cover', 'Cover: purpose, layout, method and answers')
    lines = [
        'Solution workbook for Exercise 43.17 (Part VII capstone, guided by Exhibit 43.8). Macro-free; no circular reference; iterative calculation off.',
        'Sheets: Cover, Inputs, Construction (monthly, January to December 2027), Operations, Tax, Debt, Waterfall (half-years 2028H1 to 2047H2), Returns, Checks.',
        'Case selector: Inputs F8 (named Scenario): 1 base (P50), 2 sizing (one-year P90), 3 downside (P90 and opex +10%). The debt is locked at the sizing profile in every case.',
        'Circularity: the committed debt (Inputs F33), equity (F34) and principal profile (row 57) are the converged sizing-case values (ex43_17_mirror.py, 12 passes);',
        'Construction recomputes total funding live and Debt rows 15 to 19 recompute the sculpted profile live; Checks F9 to F11 show they are the fixed point.',
        'Python mirror: ex43_17_mirror.py (same answers; interpretations listed there and in model/exercises/README.md).',
    ]
    for j, t in enumerate(lines):
        cv.cell(6 + j, 4, t)
    cv.cell(13, 4, 'Master check (0 = all pass)').font = BOLD

    # ------------------------------------------------------------------ Inputs
    ws = sheet('Inputs', 'Inputs: case control, costs, financing, operations, tax, timing, contract terms (Exhibit 43.8)', 'S')
    for i in range(NS):
        ws.cell(5, J + i).value = None; ws.cell(5, J + i).fill = PatternFill()
    ws['D5'] = 'Case table: columns J to L = cases 1 to 3; monthly rows: J to U = Jan to Dec 2027; half-year rows: J to AW = 2028H1 to 2047H2'; ws['D5'].font = BOLD
    lab(ws, 7, 'Case control', '', True)
    lab(ws, 8, 'Scenario (1 base P50, 2 sizing one-year P90, 3 downside P90 and opex +10%)', 'index'); inp(ws, 'F8', 1, '0')
    lab(ws, 9, 'Case table (columns J to L = cases 1 to 3) and live values (column F)', '', True)
    lab(ws, 10, 'Energy factor applied to P50 [table]', 'factor')
    lab(ws, 11, 'Operating cost factor [table]', 'factor')
    for j, (a, b) in enumerate([(1.0, 1.0), (0.917, 1.0), (0.917, 1.10)]):
        inp(ws, f'{cm(j)}10', a, '0.000'); inp(ws, f'{cm(j)}11', b, '0.000')
    lab(ws, 12, 'Energy factor (live)', 'factor'); sc(ws, 12, '=INDEX($J$10:$L$10,1,Scenario)', '0.000')
    lab(ws, 13, 'Operating cost factor (live)', 'factor'); sc(ws, 13, '=INDEX($J$11:$L$11,1,Scenario)', '0.000')
    lab(ws, 14, 'Construction costs (USD m)', '', True)
    items = [(15, 'PV EPC price', 'USD m', 168.4), (16, 'Battery supply price', 'USD m', 61.7),
             (17, 'Battery share paid at the first payment', '%', 20.0), (18, 'Battery share paid at the second payment', '%', 80.0),
             (19, 'Battery first payment month (January 2027 = 1)', 'month', 1), (20, 'Battery second payment month (October 2027)', 'month', 10),
             (21, "Development and owner's costs (paid in month 1)", 'USD m', 12.9), (22, 'Contingency (pro rata with PV EPC payments, months 2 to 12)', 'USD m', 9.6)]
    for r, l_, u, v in items:
        lab(ws, r, l_, u); inp(ws, f'F{r}', v, '#,##0.000')
    lab(ws, 23, 'Monthly PV EPC payment profile', '', True)
    lab(ws, 24, 'PV EPC payment profile (columns J to U = January to December 2027)', '%')
    for i, p in enumerate(mir.PV_PCT):
        inp(ws, f'{cm(i)}24', p, '0.0')
    ws['G24'] = f'=SUM(J24:{LASTM}24)'
    lab(ws, 25, 'Financing', '', True)
    fin = [(26, 'Senior loan all-in fixed rate (swapped)', '% pa', 7.25), (27, 'Months per year', 'months', 12),
           (28, 'Upfront fee (of the commitment, month 1)', '%', 1.75), (29, 'Commitment fee on the undrawn commitment', '% pa', 0.70),
           (30, 'Gearing cap (share of total funding requirement)', '%', 75.0), (31, 'Sizing DSCR (sizing case)', 'x', 1.30),
           (32, 'Distribution test: historic 12-month DSCR at least', 'x', 1.15),
           (33, 'Senior debt commitment (contract; converged sizing)', 'USD m', sol['D']),
           (34, 'Equity commitment (contract; converged sizing)', 'USD m', sol['T'] - sol['D']),
           (35, 'Number of repayment half-years (2028H1 to 2042H2)', 'count', 30)]
    for r, l_, u, v in fin:
        lab(ws, r, l_, u); inp(ws, f'F{r}', v, '#,##0.000000' if r in (33, 34) else '#,##0.000')
    lab(ws, 36, 'Operations', '', True)
    ops = [(37, 'Solar P50 energy, operating year 1', 'GWh', 512.8), (38, 'Share of annual energy in H1 (Jan-Jun)', '%', 52.0),
           (39, 'Degradation (compounding, from operating year 2)', '% pa', 0.40), (40, 'PPA price in the price base year', 'USD/MWh', 49.80),
           (41, 'Escalation of PPA price and battery payment (each January 1)', '% pa', 1.5), (42, 'Battery contracted power', 'MW', 60.0),
           (43, 'Battery capacity payment in the price base year', 'USD/kW-month', 11.40), (44, 'Fixed operating cost in the price base year', 'USD m pa', 4.36),
           (45, 'Fixed operating cost escalation (each January 1)', '% pa', 2.5), (46, 'Variable operating cost', 'USD/MWh', 1.12),
           (47, 'Price base year', 'year', 2028)]
    for r, l_, u, v in ops:
        lab(ws, r, l_, u); inp(ws, f'F{r}', v, '#,##0.000' if u != 'year' else '0')
    lab(ws, 48, 'Tax', '', True)
    for r, l_, u, v in [(49, 'Corporate income tax rate', '%', 30.0), (50, 'Last tax-holiday year (2028 to 2030; holiday losses are lost)', 'year', 2030),
                        (51, 'Tax depreciation life (straight line, half-years)', 'half-years', 40)]:
        lab(ws, r, l_, u); inp(ws, f'F{r}', v, '#,##0.000' if u == '%' else '0')
    lab(ws, 52, 'Timing', '', True)
    lab(ws, 53, 'Notice to proceed (month 1 start)', 'date'); inp(ws, 'F53', date(2027, 1, 1), 'yyyy-mm-dd')
    lab(ws, 54, 'Commercial operation date (first half-year start)', 'date'); inp(ws, 'F54', date(2028, 1, 1), 'yyyy-mm-dd')
    lab(ws, 55, 'Months per operating period', 'months'); inp(ws, 'F55', 6, '0')
    lab(ws, 56, 'Contract repayment profile', '', True)
    lab(ws, 57, 'Scheduled principal (contract; converged sizing; columns J to AW)', 'USD m')
    for i in range(NS):
        inp(ws, f'{cm(i)}57', sol['prin'][i] if i < mir.N_DEBT else 0.0, '#,##0.000')
    ws['G57'] = f'=SUM(J57:{LASTS}57)'
    lab(ws, 58, 'Numerical zero (debt service test)', 'USD m'); inp(ws, 'F58', 1e-9, '0.0E+00')
    lab(ws, 59, 'Check tolerance', 'USD m'); inp(ws, 'F59', 1e-6, '0.0E+00')
    wb.defined_names['Scenario'] = DefinedName('Scenario', attr_text='Inputs!$F$8')

    # ------------------------------------------------------------------ Construction (monthly)
    ws = sheet('Construction', 'Construction: monthly uses and pro rata funding (January to December 2027)', 'M')
    lab(ws, 7, 'Monthly timeline', '', True)
    lab(ws, 8, 'Month number', 'index'); row(ws, 8, NM, lambda i: f'={cm(i - 1) if i else "I"}8+1', '0')
    lab(ws, 9, 'Month start', 'date'); row(ws, 9, NM, lambda i: f'=IF({cm(i)}8=1,Inputs!$F$53,EDATE({cm(i - 1) if i else "I"}9,1))', 'yyyy-mm-dd')
    lab(ws, 10, 'Month end (cash flow date)', 'date'); row(ws, 10, NM, lambda i: f'=EOMONTH({cm(i)}9,0)', 'yyyy-mm-dd')
    lab(ws, 11, 'Uses (USD m)', '', True)
    lab(ws, 12, 'PV EPC payments', 'USD m'); row(ws, 12, NM, lambda i: f'=Inputs!$F$15*Inputs!{cm(i)}24/100', total=True)
    lab(ws, 13, 'Battery supply payments', 'USD m')
    row(ws, 13, NM, lambda i: f'=Inputs!$F$16*(IF({cm(i)}8=Inputs!$F$19,Inputs!$F$17,0)+IF({cm(i)}8=Inputs!$F$20,Inputs!$F$18,0))/100', total=True)
    lab(ws, 14, "Development and owner's costs", 'USD m'); row(ws, 14, NM, lambda i: f'=IF({cm(i)}8=1,Inputs!$F$21,0)', total=True)
    lab(ws, 15, 'Contingency (pro rata with PV EPC payments, months 2 to 12)', 'USD m')
    row(ws, 15, NM, lambda i: f'=IF({cm(i)}8>1,Inputs!$F$22*Inputs!{cm(i)}24/SUMIFS(Inputs!$J$24:${LASTM}$24,$J$8:${LASTM}$8,">1"),0)', total=True)
    lab(ws, 16, 'Hard costs', 'USD m'); row(ws, 16, NM, lambda i: f'=SUM({cm(i)}12:{cm(i)}15)', total=True)
    lab(ws, 17, 'Senior loan: opening balance', 'USD m'); row(ws, 17, NM, lambda i: f'={cm(i - 1) if i else "I"}24')
    lab(ws, 18, 'Interest during construction (rate / 12 on the opening balance)', 'USD m'); row(ws, 18, NM, lambda i: f'={cm(i)}17*Inputs!$F$26/100/Inputs!$F$27', total=True)
    lab(ws, 19, 'Commitment fee (on the undrawn commitment at month start)', 'USD m'); row(ws, 19, NM, lambda i: f'=(Inputs!$F$33-{cm(i)}17)*Inputs!$F$29/100/Inputs!$F$27', total=True)
    lab(ws, 20, 'Upfront fee (month 1)', 'USD m'); row(ws, 20, NM, lambda i: f'=IF({cm(i)}8=1,Inputs!$F$28/100*Inputs!$F$33,0)', total=True)
    lab(ws, 21, 'DSRA initial funding (first half-year debt service, month before COD)', 'USD m')
    row(ws, 21, NM, lambda i: f'=IF({cm(i)}10=EOMONTH(Inputs!$F$54,-1),Debt!$F$13,0)', total=True)
    lab(ws, 22, 'Total uses', 'USD m'); row(ws, 22, NM, lambda i: f'={cm(i)}16+{cm(i)}18+{cm(i)}19+{cm(i)}20+{cm(i)}21', total=True)
    lab(ws, 23, 'Senior loan drawdown (debt share g of uses)', 'USD m'); row(ws, 23, NM, lambda i: f'=$F$26*{cm(i)}22', total=True)
    lab(ws, 24, 'Senior loan: closing balance', 'USD m'); row(ws, 24, NM, lambda i: f'={cm(i)}17+{cm(i)}23')
    lab(ws, 25, 'Equity contribution', 'USD m'); row(ws, 25, NM, lambda i: f'={cm(i)}22-{cm(i)}23', total=True)
    lab(ws, 26, 'Debt share of funding g = D / (D + E)', 'fraction'); sc(ws, 26, '=Inputs!$F$33/(Inputs!$F$33+Inputs!$F$34)', '0.000000')
    lab(ws, 27, 'Total funding requirement T (sum of uses)', 'USD m'); sc(ws, 27, f'=SUM(J22:{LASTM}22)')
    lab(ws, 28, 'Capitalized cost for tax (total uses less DSRA)', 'USD m'); sc(ws, 28, f'=F27-SUM(J21:{LASTM}21)')
    lab(ws, 29, 'Gearing (senior debt / total funding)', '%'); sc(ws, 29, f'=SUM(J23:{LASTM}23)/F27*100', '0.00')
    lab(ws, 30, 'Senior debt at the gearing cap', 'USD m'); sc(ws, 30, '=Inputs!$F$30/100*F27')

    # ------------------------------------------------------------------ Operations
    ws = sheet('Operations', 'Operations: timeline, energy, revenue, operating costs (half-years 2028H1 to 2047H2)', 'S')
    p = lambda i: cm(i - 1) if i else 'I'
    lab(ws, 7, 'Semiannual timeline', '', True)
    lab(ws, 8, 'Period number', 'index'); row(ws, 8, NS, lambda i: f'={p(i)}8+1', '0')
    lab(ws, 9, 'Period start', 'date'); row(ws, 9, NS, lambda i: f'=IF({cm(i)}8=1,Inputs!$F$54,EDATE({p(i)}9,Inputs!$F$55))', 'yyyy-mm-dd')
    lab(ws, 10, 'Period end (cash flow date)', 'date'); row(ws, 10, NS, lambda i: f'=EOMONTH({cm(i)}9,Inputs!$F$55-1)', 'yyyy-mm-dd')
    lab(ws, 11, 'Calendar year', 'year'); row(ws, 11, NS, lambda i: f'=YEAR({cm(i)}9)', '0')
    lab(ws, 12, 'Operating year', 'OY'); row(ws, 12, NS, lambda i: f'={cm(i)}11-YEAR(Inputs!$F$54)+1', '0')
    lab(ws, 13, 'Flag_FirstHalf (Jan-Jun)', 'flag'); row(ws, 13, NS, lambda i: f'=IF(MONTH({cm(i)}9)=1,1,0)', '0')
    lab(ws, 14, 'Flag_TaxHoliday', 'flag'); row(ws, 14, NS, lambda i: f'=IF({cm(i)}11<=Inputs!$F$50,1,0)', '0')
    lab(ws, 15, 'Flag_Repayment (2028H1 to 2042H2)', 'flag'); row(ws, 15, NS, lambda i: f'=IF({cm(i)}8<=Inputs!$F$35,1,0)', '0')
    lab(ws, 16, 'Revenue (USD m)', '', True)
    lab(ws, 17, 'Solar energy (case factor applied)', 'GWh')
    row(ws, 17, NS, lambda i: f'=Inputs!$F$37*(1-Inputs!$F$39/100)^({cm(i)}12-1)*IF({cm(i)}13=1,Inputs!$F$38/100,1-Inputs!$F$38/100)*Inputs!$F$12', total=True)
    lab(ws, 18, 'PPA price', 'USD/MWh'); row(ws, 18, NS, lambda i: f'=Inputs!$F$40*(1+Inputs!$F$41/100)^({cm(i)}11-Inputs!$F$47)')
    lab(ws, 19, 'Energy revenue', 'USD m'); row(ws, 19, NS, lambda i: f'={cm(i)}17*{cm(i)}18/1000', total=True)
    lab(ws, 20, 'Battery capacity payment', 'USD m')
    row(ws, 20, NS, lambda i: f'=Inputs!$F$42*1000*Inputs!$F$43*(1+Inputs!$F$41/100)^({cm(i)}11-Inputs!$F$47)*Inputs!$F$55/1000000', total=True)
    lab(ws, 21, 'Total revenue', 'USD m'); row(ws, 21, NS, lambda i: f'={cm(i)}19+{cm(i)}20', total=True)
    lab(ws, 22, 'Operating costs (USD m)', '', True)
    lab(ws, 23, 'Fixed operating cost (equal halves)', 'USD m')
    row(ws, 23, NS, lambda i: f'=Inputs!$F$44*(1+Inputs!$F$45/100)^({cm(i)}11-Inputs!$F$47)*Inputs!$F$55/Inputs!$F$27*Inputs!$F$13', total=True)
    lab(ws, 24, 'Variable operating cost', 'USD m'); row(ws, 24, NS, lambda i: f'={cm(i)}17*Inputs!$F$46/1000*Inputs!$F$13', total=True)
    lab(ws, 25, 'Total operating costs', 'USD m'); row(ws, 25, NS, lambda i: f'={cm(i)}23+{cm(i)}24', total=True)
    lab(ws, 26, 'EBITDA', 'USD m'); row(ws, 26, NS, lambda i: f'={cm(i)}21-{cm(i)}25', total=True)

    # ------------------------------------------------------------------ Tax
    ws = sheet('Tax', 'Tax: holiday 2028 to 2030 (losses lost), straight-line depreciation, losses carried forward', 'S')
    lab(ws, 7, 'Depreciation (straight line over 40 half-years)', 'USD m'); row(ws, 7, NS, lambda i: f'=Construction!$F$28/Inputs!$F$51*IF(Operations!{cm(i)}8<=Inputs!$F$51,1,0)', total=True)
    lab(ws, 8, 'Senior interest (deductible)', 'USD m'); row(ws, 8, NS, lambda i: f'=Debt!{cm(i)}9', total=True)
    lab(ws, 9, 'Taxable income before losses', 'USD m'); row(ws, 9, NS, lambda i: f'=Operations!{cm(i)}26-{cm(i)}8-{cm(i)}7', total=True)
    lab(ws, 10, 'Tax losses used', 'USD m'); row(ws, 10, NS, lambda i: f'=IF(Operations!{cm(i)}14=1,0,MIN({p(i)}12,MAX(0,{cm(i)}9)))', total=True)
    lab(ws, 11, 'Tax losses arising (holiday losses lost)', 'USD m'); row(ws, 11, NS, lambda i: f'=IF(Operations!{cm(i)}14=1,0,MAX(0,-{cm(i)}9))', total=True)
    lab(ws, 12, 'Tax losses carried forward, closing', 'USD m'); row(ws, 12, NS, lambda i: f'={p(i)}12-{cm(i)}10+{cm(i)}11')
    lab(ws, 13, 'Taxable income', 'USD m'); row(ws, 13, NS, lambda i: f'=IF(Operations!{cm(i)}14=1,0,MAX(0,{cm(i)}9-{cm(i)}10))', total=True)
    lab(ws, 14, 'Tax paid', 'USD m'); row(ws, 14, NS, lambda i: f'={cm(i)}13*Inputs!$F$49/100', total=True)

    # ------------------------------------------------------------------ Debt
    ws = sheet('Debt', 'Debt: senior loan on the contract profile; live sculpting check (sizing case)', 'S')
    lab(ws, 7, 'Senior loan (USD m)', '', True)
    lab(ws, 8, 'Opening balance', 'USD m'); row(ws, 8, NS, lambda i: f'=IF(Operations!{cm(i)}8=1,Inputs!$F$33,{p(i)}12)')
    lab(ws, 9, 'Interest (rate x months / 12 on the opening balance)', 'USD m'); row(ws, 9, NS, lambda i: f'={cm(i)}8*Inputs!$F$26/100*Inputs!$F$55/Inputs!$F$27', total=True)
    lab(ws, 10, 'Principal (contract profile)', 'USD m'); row(ws, 10, NS, lambda i: f'=Inputs!{cm(i)}57', total=True)
    lab(ws, 11, 'Debt service', 'USD m'); row(ws, 11, NS, lambda i: f'={cm(i)}9+{cm(i)}10', total=True)
    lab(ws, 12, 'Closing balance', 'USD m'); row(ws, 12, NS, lambda i: f'={cm(i)}8-{cm(i)}10')
    lab(ws, 13, 'First half-year debt service (DSRA initial funding)', 'USD m'); sc(ws, 13, '=J11')
    lab(ws, 14, 'Live sculpting check (meaningful in the sizing case, Scenario 2)', '', True)
    lab(ws, 15, 'Sculpted debt service = CFADS / sizing DSCR', 'USD m'); row(ws, 15, NS, lambda i: f'=Operations!{cm(i)}15*Waterfall!{cm(i)}7/Inputs!$F$31', total=True)
    lab(ws, 16, 'PV of sculpted debt service at the loan rate (backward recursion)', 'USD m')
    row(ws, 16, NS, lambda i: f'=IF(Operations!{cm(i)}15=1,({cm(i)}15+{cm(i + 1)}16)/(1+Inputs!$F$26/100*Inputs!$F$55/Inputs!$F$27),0)')
    lab(ws, 17, 'Debt capacity at the sizing DSCR', 'USD m'); sc(ws, 17, '=J16')
    lab(ws, 18, 'Sculpted principal', 'USD m'); row(ws, 18, NS, lambda i: f'={cm(i)}15-{cm(i)}8*Inputs!$F$26/100*Inputs!$F$55/Inputs!$F$27', total=True)
    lab(ws, 19, 'Sculpted principal less contract principal', 'USD m'); row(ws, 19, NS, lambda i: f'={cm(i)}18-{cm(i)}10')

    # ------------------------------------------------------------------ Waterfall
    ws = sheet('Waterfall', 'Waterfall: CFADS, debt service, DSRA, distribution test, distributions', 'S')
    lab(ws, 7, 'CFADS (EBITDA less tax)', 'USD m'); row(ws, 7, NS, lambda i: f'=Operations!{cm(i)}26-Tax!{cm(i)}14', total=True)
    lab(ws, 8, 'Senior debt service', 'USD m'); row(ws, 8, NS, lambda i: f'=Debt!{cm(i)}11', total=True)
    lab(ws, 9, 'DSCR, period', 'x'); row(ws, 9, NS, lambda i: f'=IF({cm(i)}8>Inputs!$F$58,{cm(i)}7/{cm(i)}8,0)', '0.0000')
    lab(ws, 10, 'DSCR, historic 12 months (this and the previous half-year)', 'x'); row(ws, 10, NS, lambda i: f'=IF({cm(i)}8>Inputs!$F$58,({cm(i)}7+{p(i)}7)/({cm(i)}8+{p(i)}8),0)', '0.0000')
    lab(ws, 11, 'DSRA opening', 'USD m'); row(ws, 11, NS, lambda i: f'=IF(Operations!{cm(i)}8=1,SUM(Construction!$J$21:${LASTM}$21),{p(i)}14)')
    lab(ws, 12, 'DSRA target (next half-year debt service; nil after final maturity)', 'USD m'); row(ws, 12, NS, lambda i: f'={cm(i + 1)}8')
    lab(ws, 13, 'DSRA release (+) or top-up (-)', 'USD m'); row(ws, 13, NS, lambda i: f'={cm(i)}11-{cm(i)}12', total=True)
    lab(ws, 14, 'DSRA closing', 'USD m'); row(ws, 14, NS, lambda i: f'={cm(i)}12')
    lab(ws, 15, 'Cash available for distribution (incl. locked-up cash brought forward)', 'USD m'); row(ws, 15, NS, lambda i: f'={cm(i)}7-{cm(i)}8+{cm(i)}13+{p(i)}18')
    lab(ws, 16, 'Distribution test passed (historic DSCR at least 1.15x, or no debt service)', 'flag')
    row(ws, 16, NS, lambda i: f'=IF(OR({cm(i)}8<=Inputs!$F$58,{cm(i)}10>=Inputs!$F$32),1,0)', '0')
    lab(ws, 17, 'Distributions to equity', 'USD m'); row(ws, 17, NS, lambda i: f'={cm(i)}15*{cm(i)}16', total=True)
    lab(ws, 18, 'Locked-up cash, closing', 'USD m'); row(ws, 18, NS, lambda i: f'={cm(i)}15-{cm(i)}17')
    lab(ws, 19, 'Lock-up (1 = distribution blocked)', 'flag'); row(ws, 19, NS, lambda i: f'=1-{cm(i)}16', '0', total=True)

    # ------------------------------------------------------------------ Returns
    ws = sheet('Returns', 'Returns: dated equity cash flows (12 month ends, then 40 half-year ends) and summary outputs', 'R')
    lab(ws, 7, 'Cash flow date', 'date')
    row(ws, 7, NR, lambda i: f'=Construction!{cm(i)}10' if i < NM else f'=Operations!{cm(i - NM)}10', 'yyyy-mm-dd')
    lab(ws, 8, 'Equity cash flow (contributions negative)', 'USD m')
    row(ws, 8, NR, lambda i: f'=-Construction!{cm(i)}25' if i < NM else f'=Waterfall!{cm(i - NM)}17', total=True)
    lab(ws, 9, 'Summary outputs (selected case)', '', True)
    outs = [(10, 'Equity IRR (XIRR)', '%', f'=XIRR(J8:{LASTR}8,J7:{LASTR}7)*100', '0.00'),
            (11, 'Minimum DSCR (repayment half-years)', 'x', f'=_xlfn.MINIFS(Waterfall!J9:{LASTS}9,Waterfall!J8:{LASTS}8,">"&Inputs!$F$58)', '0.000'),
            (12, 'Average DSCR (CFADS / debt service over the repayment half-years)', 'x', f'=SUMIFS(Waterfall!J7:{LASTS}7,Waterfall!J8:{LASTS}8,">"&Inputs!$F$58)/SUM(Waterfall!J8:{LASTS}8)', '0.000'),
            (13, 'Number of lock-ups', 'count', f'=SUM(Waterfall!J19:{LASTS}19)', '0'),
            (14, 'Hard costs', 'USD m', '=Construction!$G$16', '#,##0.000'),
            (15, 'Interest during construction', 'USD m', '=Construction!$G$18', '#,##0.000'),
            (16, 'Commitment fees', 'USD m', '=Construction!$G$19', '#,##0.000'),
            (17, 'Upfront fee', 'USD m', '=Construction!$G$20', '#,##0.000'),
            (18, 'DSRA initial funding', 'USD m', '=Construction!$G$21', '#,##0.000'),
            (19, 'Total funding requirement', 'USD m', '=Construction!$F$27', '#,##0.000'),
            (20, 'Senior debt', 'USD m', '=Construction!$G$23', '#,##0.000'),
            (21, 'Senior debt at the gearing cap', 'USD m', '=Construction!$F$30', '#,##0.000'),
            (22, 'Gearing', '%', '=Construction!$F$29', '0.0'),
            (23, 'Equity', 'USD m', '=Construction!$G$25', '#,##0.000'),
            (24, 'First half-year debt service', 'USD m', '=Debt!$F$13', '#,##0.000')]
    for r, l_, u, f, fmt in outs:
        lab(ws, r, l_, u); sc(ws, r, f, fmt)

    # ------------------------------------------------------------------ Checks
    ws = sheet('Checks', 'Checks: every check shows 0 when passing')
    tol = 'Inputs!$F$59'
    chk = [(7, 'Sources less uses (construction)', f'=IF(ABS(Construction!$G$23+Construction!$G$25-Construction!$F$27)<={tol},0,Construction!$G$23+Construction!$G$25-Construction!$F$27)'),
           (8, 'PV EPC profile sums to 100%', f'=IF(ABS(Inputs!$G$24-100)<={tol},0,Inputs!$G$24-100)'),
           (9, 'Committed debt drawn in full (debt drawn less commitment)', f'=IF(ABS(Construction!$G$23-Inputs!$F$33)<={tol},0,Construction!$G$23-Inputs!$F$33)'),
           (10, 'Sizing case: live debt capacity at 1.30x equals the commitment', f'=IF(Scenario=2,IF(ABS(Debt!$F$17-Inputs!$F$33)<={tol},0,Debt!$F$17-Inputs!$F$33),0)'),
           (11, 'Sizing case: live sculpted principal equals the contract profile (max abs)', f'=IF(Scenario=2,IF(MAX(MAX(Debt!J19:{LASTS}19),-MIN(Debt!J19:{LASTS}19))<={tol},0,MAX(MAX(Debt!J19:{LASTS}19),-MIN(Debt!J19:{LASTS}19))),0)'),
           (12, 'Debt not above the gearing cap', '=IF(Inputs!$F$33<=Construction!$F$30,0,1)'),
           (13, 'Debt repaid by final maturity (closing balance in the last repayment half-year)', f'=IF(ABS(INDEX(Debt!J12:{LASTS}12,1,Inputs!$F$35))<={tol},0,INDEX(Debt!J12:{LASTS}12,1,Inputs!$F$35))'),
           (14, 'Contract profile sums to the commitment', f'=IF(ABS(Inputs!$G$57-Inputs!$F$33)<={tol},0,Inputs!$G$57-Inputs!$F$33)'),
           (15, 'No negative cash after debt service and DSRA (min, if below 0)', f'=IF(MIN(Waterfall!J15:{LASTS}15)>=-{tol},0,MIN(Waterfall!J15:{LASTS}15))'),
           (16, 'DSRA released at final maturity (closing DSRA in the last repayment half-year)', f'=IF(ABS(INDEX(Waterfall!J14:{LASTS}14,1,Inputs!$F$35))<={tol},0,1)'),
           (17, 'Energy shares H1 + H2 = 100% (operating year 1)', f'=IF(ABS(Operations!J17+Operations!K17-Inputs!$F$37*Inputs!$F$12)<={tol},0,1)'),
           (18, 'Equity flows: contributions equal the equity commitment', f'=IF(ABS(-SUMIF(Returns!J8:{LASTR}8,"<0")-Inputs!$F$34)<={tol},0,1)')]
    for r, l_, f in chk:
        ws.cell(r, 4, l_); ws.cell(r, 5, 'check'); c = ws.cell(r, 6, f); c.number_format = '0.000000'
    ws.cell(20, 4, 'Sum of checks (0 = all pass)').font = BOLD; ws.cell(20, 5, 'check')
    ws.cell(20, 6, '=' + '+'.join(f'ABS(F{r})' for r, _, _ in chk)).number_format = '0.000000'
    ws.conditional_formatting.add('F7:F20', CellIsRule(operator='notEqual', formula=['0'], fill=RED))
    cv['F13'] = '=Checks!$F$20'; cv['F13'].font = GREEN
    cv.conditional_formatting.add('F13', CellIsRule(operator='notEqual', formula=['0'], fill=RED))
    # answers block on the Cover (live links)
    cv.cell(15, 4, 'Answers for the selected case (live; Returns sheet)').font = BOLD
    for j, (r, l_, u, f, fmt) in enumerate(outs):
        cv.cell(16 + j, 4, l_); cv.cell(16 + j, 5, u)
        c = cv.cell(16 + j, 6, f'=Returns!$F${r}'); c.number_format = fmt; c.font = GREEN
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
    return out


if __name__ == '__main__':
    p = os.path.join(HERE, 'ex43_17_solution.xlsx')
    build(p)
    print('written', p)
