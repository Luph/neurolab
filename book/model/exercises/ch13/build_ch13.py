#!/usr/bin/env python3
"""Build the three Chapter 13 exercise workbooks named in bible/briefs/u03.md (macro-free .xlsx):

  Ch13_Practice_Solution.xlsx   Exercise 13.12 and Section 13.10 (walkthrough sec:13.10):
                                Quebracho Alto practice workbook, sheets Inputs, Time, Calc, Checks,
                                reproducing Examples 13.2, 13.3, 13.6, 13.7, 13.8 and 13.9.
  Ch13_Inherited_Sheet.xlsx     Exercise 13.13: the seeded sheet exactly as specified (five errors kept).
  Ch13_LlanoPardo_Solution.xlsx Exercise 13.15: Llano Pardo projection (u01 Section 1.A.3 to 1.A.7),
                                k = 1.38074 stored as a goal-seek result with a check.

Rates in these Chapter 13 files are Excel percentages (0.0875 shown as 8.75%), as the chapter's own
formulas use them (=-F5+PV(F9,F10,...)); the companion model's percentage-number convention starts in
Chapter 39.  Inputs blue on yellow; links from other sheets green; checks 0 when passing.
"""
import os, math
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.formatting.rule import CellIsRule
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.formula import DataTableFormula
from openpyxl.utils import get_column_letter as L

HERE = os.path.dirname(os.path.abspath(__file__))
BLUE = Font(color='0000FF'); GREEN = Font(color='008000'); BOLD = Font(bold=True)
YEL = PatternFill('solid', fgColor='FFF2CC'); RED = PatternFill('solid', fgColor='FF9999'); HEAD = PatternFill('solid', fgColor='DDEBF7')
J = 10
cm = lambda i: L(J + i)
pv_ = lambda i: cm(i - 1) if i else 'I'


def header(ws, title, case='Quebracho Alto (fictional; Illustrative)'):
    ws['A1'] = f'{case}: Chapter 13 practice files'; ws['A1'].font = Font(bold=True, size=13)
    ws['A2'] = title
    ws['A3'] = 'Column D label, E unit, F value or scalar, G row total; first period in column J. Inputs: blue on yellow. Links: green. Checks: 0 when passing.'
    for col, w in zip('ABCDEFGHI', [2, 2, 2, 50, 12, 14, 12, 4, 4]):
        ws.column_dimensions[col].width = w


def inp(ws, ref, v, fmt=None):
    c = ws[ref]; c.value = v; c.font = BLUE; c.fill = YEL
    if fmt: c.number_format = fmt
    return c


def lab(ws, r, label, unit=None, sec=False):
    if sec:
        ws.cell(r, 2, label).font = BOLD
    else:
        ws.cell(r, 4, label)
        if unit: ws.cell(r, 5, unit)


def frm(ws, ref, f, fmt=None):
    c = ws[ref]; c.value = f
    if fmt: c.number_format = fmt
    if isinstance(f, str) and f.startswith('=') and f.count('!') == 1 and not any(o in f[1:] for o in '+-*/(,'):
        c.font = GREEN
    return c


def row(ws, r, n, fn, fmt='#,##0.000', total=False):
    for i in range(n):
        frm(ws, f'{cm(i)}{r}', fn(i), fmt)
    if total:
        ws.cell(r, 7, f'=SUM(J{r}:{cm(n - 1)}{r})').number_format = fmt


# =============================================================================================
# Practice workbook (Exercise 13.12)
# =============================================================================================
NMON = 30        # monthly columns J..AM: March 2027 to August 2029 (26 construction months + 4)
NSEM = 41        # semiannual columns J..AX: 2029H1 to 2049H1
NYR = 20         # operating years J..AC
NLD = 8          # delay months of Example 13.3


def practice(path):
    wb = Workbook(); wb.remove(wb.active)
    I = wb.create_sheet('Inputs'); T = wb.create_sheet('Time'); C = wb.create_sheet('Calc'); K = wb.create_sheet('Checks')
    header(I, 'Inputs: every number the workbook uses, with units (Examples 13.2 to 13.9)')
    header(T, 'Time: monthly construction timeline (rows 4 to 12) and semiannual operations timeline (rows 18 to 24)')
    header(C, 'Calc: LD cap (Ex 13.3), revenue and NPV rows (Ex 13.5, 13.6), data table, goal seek (Ex 13.7), IDC closed form and iteration (Ex 13.8), sources and uses (Ex 13.9)')
    header(K, 'Checks: each check 0 when passing; master check (Example 13.9)')
    # ---------------- Inputs
    lab(I, 5, 'Check and macro settings', sec=True)
    lab(I, 6, 'Check tolerance', 'USD m'); inp(I, 'F6', 0.001, '0.000')
    lab(I, 7, 'Scenario control (Example 13.4)', sec=True)
    lab(I, 8, 'Scenario (1 Base, 2 Low, 3 High; cell named Scenario)', 'index'); inp(I, 'F8', 1, '0')
    lab(I, 9, 'Escalation factors by operating year (columns F to Y = years 1 to 20; Example 13.5)', sec=True)
    lab(I, 10, 'Escalation factor (tariff escalation 2.4% a year)', None)
    for k in range(NYR):
        inp(I, f'{L(6 + k)}10', 1.024 ** k, '0.000000')
    lab(I, 11, 'Timing (Example 13.2)', sec=True)
    lab(I, 12, 'Notice to proceed', 'date'); inp(I, 'F12', date(2027, 3, 1), 'yyyy-mm-dd')
    lab(I, 13, 'Construction months', 'months'); inp(I, 'F13', 26, '0')
    lab(I, 14, 'Commercial operation date', 'date'); inp(I, 'F14', date(2029, 5, 1), 'yyyy-mm-dd')
    lab(I, 15, 'Operating years', 'years'); inp(I, 'F15', NYR, '0')
    lab(I, 16, 'Months per operating period', 'months'); inp(I, 'F16', 6, '0')
    lab(I, 17, 'Operations and revenue (Examples 13.5 to 13.7)', sec=True)
    lab(I, 18, 'Annual energy (no degradation in these examples)', 'GWh'); inp(I, 'F18', 551.4, '#,##0.0')
    lab(I, 19, 'Base tariff (year 1)', 'USD/MWh'); inp(I, 'F19', 63.40, '0.00')
    lab(I, 20, 'Tariff escalation (Example 13.5; factors in row 10)', '% pa'); inp(I, 'F20', 0.024, '0.0%')
    lab(I, 21, 'Operating cost', 'USD m pa'); inp(I, 'F21', 7.94, '0.00')
    lab(I, 22, 'Discount rate (base)', '% pa'); inp(I, 'F22', 0.0875, '0.00%')
    lab(I, 23, 'Delay LDs (Example 13.3)', sec=True)
    lab(I, 24, 'Delay LD rate', 'USD/day'); inp(I, 'F24', 39500, '#,##0')
    lab(I, 25, 'Delay LD cap', 'USD'); inp(I, 'F25', 9620000, '#,##0')
    lab(I, 26, 'Days in each delay month (columns J to Q = delay months 1 to 8)', 'days')
    for i, d in enumerate([31, 30, 31, 31, 30, 31, 30, 31]):
        inp(I, f'{cm(i)}26', d, '0')
    lab(I, 27, 'Funding (Example 13.8)', sec=True)
    lab(I, 28, 'Gearing (debt share of total funding)', '%'); inp(I, 'F28', 0.72, '0.00%')
    lab(I, 29, 'IDC as a share of total debt (whole construction period; simplification)', '%'); inp(I, 'F29', 0.0874, '0.00%')
    lab(I, 30, 'Costs before financing: scenario table H30:J30 (Base, Low, High); live value in F30', 'USD m')
    for col, v in zip('HIJ', [248.6, 231.7, 271.9]):
        inp(I, f'{col}30', v, '0.0')
    frm(I, 'F30', '=INDEX(H30:J30,Scenario)', '0.0')
    I['H29'] = 'Base'; I['I29'] = 'Low'; I['J29'] = 'High'
    lab(I, 31, 'Macro tolerance (optional Converge macro, ssec:13.8.5)', 'USD m'); inp(I, 'F31', 0.001, '0.000')
    lab(I, 32, 'Macro pass limit', 'passes'); inp(I, 'F32', 50, '0')
    lab(I, 33, 'Goal seek (Example 13.7)', sec=True)
    be = (248.6 / ((1 - 1.0875 ** -20) / 0.0875) + 7.94) / 551.4 * 1000
    lab(I, 34, 'Breakeven tariff at the base rate, pasted from Goal Seek (NPV = 0 by changing the tariff)', 'USD/MWh'); inp(I, 'F34', be, '0.0000')
    lab(I, 35, 'Breakeven test rate, low (Example 13.7)', '% pa'); inp(I, 'F35', 0.0775, '0.00%')
    lab(I, 36, 'Breakeven test rate, high (Example 13.7)', '% pa'); inp(I, 'F36', 0.0975, '0.00%')
    wb.defined_names['Scenario'] = DefinedName('Scenario', attr_text='Inputs!$F$8')

    # ---------------- Time
    lab(T, 4, 'Month counter', 'index'); row(T, 4, NMON, lambda i: f'={pv_(i)}4+1', '0')
    for i in range(NMON):
        c = T.cell(5, J + i, date(2027 + (2 + i) // 12, (2 + i) % 12 + 1, 1).strftime('%Y-%m')); c.font = BOLD; c.fill = HEAD
    lab(T, 7, 'Days in month', 'days'); row(T, 7, NMON, lambda i: f'={cm(i)}10-{cm(i)}9+1', '0')
    lab(T, 9, 'Month start', 'date'); row(T, 9, NMON, lambda i: f'=IF({cm(i)}4=1,$F$14,{pv_(i)}10+1)', 'yyyy-mm-dd')
    lab(T, 10, 'Month end', 'date'); row(T, 10, NMON, lambda i: f'=IF({cm(i)}4=1,EOMONTH($F$14,0),EOMONTH({pv_(i)}10,1))', 'yyyy-mm-dd')
    lab(T, 12, 'Flag_Construction', 'flag'); row(T, 12, NMON, lambda i: f'=({cm(i)}10<=$F$15)*({cm(i)}9>=$F$14)', '0', total=True)
    lab(T, 13, 'Key dates (column F)', sec=True)
    lab(T, 14, 'Start of construction (notice to proceed)', 'date'); frm(T, 'F14', '=Inputs!$F$12', 'yyyy-mm-dd')
    lab(T, 15, 'End of construction (last day of the last construction month)', 'date'); frm(T, 'F15', '=EOMONTH(EDATE(F14,Inputs!$F$13-1),0)', 'yyyy-mm-dd')
    lab(T, 16, 'Commercial operation date', 'date'); frm(T, 'F16', '=Inputs!$F$14', 'yyyy-mm-dd')
    lab(T, 17, 'End of operations', 'date'); frm(T, 'F17', '=EDATE(F16,12*Inputs!$F$15)-1', 'yyyy-mm-dd')
    lab(T, 18, 'Half-year counter (operations timeline)', 'index'); row(T, 18, NSEM, lambda i: f'={pv_(i)}18+1', '0')
    lab(T, 19, 'Half-year start', 'date')
    row(T, 19, NSEM, lambda i: f'=IF({cm(i)}18=1,DATE(YEAR($F$16),1+Inputs!$F$16*INT((MONTH($F$16)-1)/Inputs!$F$16),1),{pv_(i)}20+1)', 'yyyy-mm-dd')
    lab(T, 20, 'Half-year end', 'date'); row(T, 20, NSEM, lambda i: f'=EOMONTH({cm(i)}19,Inputs!$F$16-1)', 'yyyy-mm-dd')
    lab(T, 21, 'Days in half-year', 'days'); row(T, 21, NSEM, lambda i: f'={cm(i)}20-{cm(i)}19+1', '0')
    lab(T, 22, 'Flag_Operations', 'flag'); row(T, 22, NSEM, lambda i: f'=IF(AND({cm(i)}20>=$F$16,{cm(i)}19<=$F$17),1,0)', '0')
    lab(T, 23, 'Flag_FirstOperationsPeriod', 'flag'); row(T, 23, NSEM, lambda i: f'=IF(AND({cm(i)}22=1,{pv_(i)}22=0),1,0)', '0')
    lab(T, 24, 'Operations fraction of the period (actual days)', 'fraction')
    row(T, 24, NSEM, lambda i: f'=MAX(0,MIN({cm(i)}20,$F$17)-MAX({cm(i)}19,$F$16)+1)/({cm(i)}20-{cm(i)}19+1)', '0.0000')
    lab(T, 25, 'Operations fraction in the first operating period', 'fraction'); frm(T, 'F25', f'=SUMPRODUCT(J23:{cm(NSEM - 1)}23,J24:{cm(NSEM - 1)}24)', '0.0000')
    T.freeze_panes = 'J6'

    # ---------------- Calc
    lab(C, 4, 'Operating year counter (revenue and NPV rows, columns J to AC)', 'index'); row(C, 4, NYR, lambda i: f'={pv_(i)}4+1', '0')
    lab(C, 6, 'Example 13.3: delay LDs with a cumulative cap (columns J to Q = delay months 1 to 8)', sec=True)
    lab(C, 7, 'Days in the delay month', 'days'); row(C, 7, NLD, lambda i: f'=Inputs!{cm(i)}26', '0')
    lab(C, 20, 'Delay LD rate', 'USD/day'); frm(C, 'F20', '=Inputs!$F$24', '#,##0')
    lab(C, 21, 'Delay LD cap', 'USD'); frm(C, 'F21', '=Inputs!$F$25', '#,##0')
    lab(C, 22, 'Delay LD accrual (capped)', 'USD'); row(C, 22, NLD, lambda i: f'=MAX(0,MIN($F$20*{cm(i)}$7,$F$21-{pv_(i)}23))', '#,##0', total=True)
    lab(C, 23, 'Delay LDs, cumulative', 'USD'); row(C, 23, NLD, lambda i: f'={pv_(i)}23+{cm(i)}22', '#,##0')
    lab(C, 24, 'Days to reach the cap', 'days'); frm(C, 'F24', '=F21/F20', '0.0')
    lab(C, 25, 'Examples 13.5 and 13.6: 20-year revenue and NPV rows (columns J to AC = operating years 1 to 20)', sec=True)
    lab(C, 26, 'Escalation factor (INDEX on the Inputs factor row; Example 13.5 replacement for OFFSET)', 'factor')
    row(C, 26, NYR, lambda i: f'=INDEX(Inputs!$F$10:$Y$10,{cm(i)}$4)', '0.000000')
    lab(C, 27, 'Revenue with escalation (Example 13.5)', 'USD m'); row(C, 27, NYR, lambda i: f'=Inputs!$F$18*Inputs!$F$19*{cm(i)}26/1000', total=True)
    lab(C, 28, 'Discount factor at the base rate (end of year)', 'factor'); row(C, 28, NYR, lambda i: f'=1/(1+Inputs!$F$22)^{cm(i)}$4', '0.000000')
    lab(C, 29, 'PV of revenue with escalation at the base rate', 'USD m'); frm(C, 'F29', f'=SUMPRODUCT(J27:{cm(NYR - 1)}27,J28:{cm(NYR - 1)}28)')
    lab(C, 30, 'Data-table input cell: flat tariff (row input; must sit on this sheet)', 'USD/MWh'); inp(C, 'F30', 62.00, '0.00')
    lab(C, 31, 'Data-table input cell: discount rate (column input; must sit on this sheet)', '% pa'); inp(C, 'F31', 0.0875, '0.00%')
    lab(C, 32, 'Revenue at the flat tariff (Example 13.6: no escalation)', 'USD m'); row(C, 32, NYR, lambda i: f'=Inputs!$F$18*$F$30/1000', total=True)
    lab(C, 33, 'Operating cost', 'USD m'); row(C, 33, NYR, lambda i: f'=Inputs!$F$21', total=True)
    lab(C, 34, 'Net cash flow', 'USD m'); row(C, 34, NYR, lambda i: f'={cm(i)}32-{cm(i)}33', total=True)
    lab(C, 35, 'Discount factor at the data-table rate', 'factor'); row(C, 35, NYR, lambda i: f'=1/(1+$F$31)^{cm(i)}$4', '0.000000')
    lab(C, 36, 'Present value of net cash flow', 'USD m'); row(C, 36, NYR, lambda i: f'={cm(i)}34*{cm(i)}35', total=True)
    lab(C, 37, 'NPV = -capex + sum of present values (feeds the data table)', 'USD m'); frm(C, 'F37', '=-Inputs!$F$30+G36')
    lab(C, 38, 'NPV, annuity form (Example 13.6 Excel form)', 'USD m'); frm(C, 'F38', '=-Inputs!$F$30+PV(F31,Inputs!$F$15,-(Inputs!$F$18*F30/1000-Inputs!$F$21))')
    lab(C, 39, 'Example 13.8: IDC circularity (C, g, r_c from Inputs)', sec=True)
    lab(C, 40, 'Senior debt, closed form D = gC / (1 - g r_c)', 'USD m')
    frm(C, 'F40', '=Inputs!$F$28*Inputs!$F$30/(1-Inputs!$F$28*Inputs!$F$29)', '0.0000')
    lab(C, 41, 'Pasted debt (written by the optional Converge macro; 0 = not run)', 'USD m'); inp(C, 'F41', 0.0, '0.0000')
    lab(C, 42, 'Debt computed from the pasted debt: g x (C + r_c x pasted)', 'USD m'); frm(C, 'F42', '=Inputs!$F$28*(Inputs!$F$30+Inputs!$F$29*F41)', '0.0000')
    lab(C, 43, 'Convergence residual (computed less pasted)', 'USD m'); frm(C, 'F43', '=F42-F41', '0.0000')
    lab(C, 44, 'Iteration D0 = g x C', 'USD m'); frm(C, 'F44', '=Inputs!$F$28*Inputs!$F$30', '0.0000')
    for k in range(1, 6):
        lab(C, 44 + k, f'Iteration D{k} = g x (C + r_c x D{k - 1})', 'USD m'); frm(C, f'F{44 + k}', f'=Inputs!$F$28*(Inputs!$F$30+Inputs!$F$29*F{43 + k})', '0.0000')
    lab(C, 50, 'NPV at the pasted goal-seek tariff and the base rate (Example 13.7; must be 0)', 'USD m')
    frm(C, 'F50', '=-Inputs!$F$30+PV(Inputs!$F$22,Inputs!$F$15,-(Inputs!$F$18*Inputs!$F$34/1000-Inputs!$F$21))', '0.000000')
    lab(C, 51, 'Breakeven tariff, closed form, at the low test rate', 'USD/MWh')
    frm(C, 'F51', '=(Inputs!$F$30/PV(Inputs!$F$35,Inputs!$F$15,-1)+Inputs!$F$21)/Inputs!$F$18*1000', '0.00')
    lab(C, 52, 'Breakeven tariff, closed form, at the base rate', 'USD/MWh')
    frm(C, 'F52', '=(Inputs!$F$30/PV(Inputs!$F$22,Inputs!$F$15,-1)+Inputs!$F$21)/Inputs!$F$18*1000', '0.00')
    lab(C, 53, 'Breakeven tariff, closed form, at the high test rate', 'USD/MWh')
    frm(C, 'F53', '=(Inputs!$F$30/PV(Inputs!$F$36,Inputs!$F$15,-1)+Inputs!$F$21)/Inputs!$F$18*1000', '0.00')
    lab(C, 54, 'Annuity factor at the base rate', 'factor'); frm(C, 'F54', '=PV(Inputs!$F$22,Inputs!$F$15,-1)', '0.00000')
    lab(C, 55, 'Interest during construction = r_c x D', 'USD m'); frm(C, 'F55', '=Inputs!$F$29*F40', '0.0000')
    lab(C, 56, 'Total funding = C + IDC', 'USD m'); frm(C, 'F56', '=Inputs!$F$30+F55', '0.0000')
    lab(C, 57, 'Equity = total funding - debt', 'USD m'); frm(C, 'F57', '=F56-F40', '0.0000')
    lab(C, 59, 'Example 13.9: sources and uses', sec=True)
    lab(C, 60, 'Sources: debt + equity', 'USD m'); frm(C, 'F60', '=F40+F57', '0.0000')
    lab(C, 61, 'Uses: costs before financing + IDC', 'USD m'); frm(C, 'F61', '=Inputs!$F$30+F55', '0.0000')
    lab(C, 63, 'Example 13.6: two-way data table of NPV (USD m). Corner E64 = output; tariffs across (row input F30); rates down (column input F31)', sec=True)
    frm(C, 'E64', '=F37', '0.00')
    for j, t in enumerate([58.0, 62.0, 66.0]):
        inp(C, f'{L(6 + j)}64', t, '0.00')
    for i, r_ in enumerate([0.0775, 0.0875, 0.0975]):
        inp(C, f'E{65 + i}', r_, '0.00%')
    C['F65'] = DataTableFormula(ref='F65:H67', dt2D=True, r1='F30', r2='F31')
    for rr in range(65, 68):
        for cc in 'FGH':
            C[f'{cc}{rr}'].number_format = '#,##0.00;(#,##0.00)'
    C['D64'] = 'Data table (Data, What-If Analysis, Data Table; computed by Excel)'
    lab(C, 69, 'Closed-form grid: the same NPVs computed directly (live check on the data table)', sec=True)
    C['D70'] = 'Rate \\ tariff'
    for j in range(3):
        frm(C, f'{L(6 + j)}70', f'={L(6 + j)}64', '0.00')
    for i in range(3):
        frm(C, f'E{71 + i}', f'=E{65 + i}', '0.00%')
        for j in range(3):
            col = L(6 + j)
            frm(C, f'{col}{71 + i}', f'=-Inputs!$F$30+PV($E{71 + i},Inputs!$F$15,-(Inputs!$F$18*{col}$70/1000-Inputs!$F$21))', '#,##0.00;(#,##0.00)')
    C.freeze_panes = 'J6'

    # ---------------- Checks
    tol = 'Inputs!$F$6'
    chk = [
        (7, 'Sources less uses (Example 13.9)', f'=--(ABS(Calc!F60-Calc!F61)>{tol})'),
        (8, 'Construction flags sum to the construction months', f'=--(Time!G12<>Inputs!$F$13)'),
        (9, 'Operations fraction between 0 and 1 in every period', f'=--OR(MIN(Time!J24:{cm(NSEM - 1)}24)<0,MAX(Time!J24:{cm(NSEM - 1)}24)>1)'),
        (10, 'Revenue positive in every operating period (catches the Example 13.5 OFFSET failure)', f'=--(MIN(Calc!J27:{cm(NYR - 1)}27)<=0)'),
        (11, 'Goal-seek tariff still gives NPV = 0 (Example 13.7)', f'=--(ABS(Calc!F50)>{tol})'),
        (12, 'Closed-form debt solves D = g(C + r_c D)', f'=--(ABS(Calc!F40-Inputs!$F$28*(Inputs!$F$30+Inputs!$F$29*Calc!F40))>{tol})'),
        (13, 'Iteration D5 within tolerance of the closed form', f'=--(ABS(Calc!F49-Calc!F40)>{tol})'),
        (14, 'Delay LDs never exceed the cap', '=--(Calc!G22>Calc!F21)'),
        (15, 'NPV row equals the annuity form', f'=--(ABS(Calc!F37-Calc!F38)>{tol})'),
        (16, 'Data table equals the closed-form grid (Excel computes the table on open; blank table = not yet computed)',
         f'=IF(COUNT(Calc!F65:H67)=0,0,--(MAX(ABS(Calc!F65-Calc!F71),ABS(Calc!G65-Calc!G71),ABS(Calc!H65-Calc!H71),ABS(Calc!F66-Calc!F72),ABS(Calc!G66-Calc!G72),ABS(Calc!H66-Calc!H72),ABS(Calc!F67-Calc!F73),ABS(Calc!G67-Calc!G73),ABS(Calc!H67-Calc!H73))>{tol}))'),
    ]
    for r, l_, f in chk:
        K.cell(r, 4, l_); K.cell(r, 5, 'check'); K.cell(r, 6, f).number_format = '0'
    K.cell(18, 4, 'Master check (sum of the checks above; 0 = all pass)').font = BOLD; K.cell(18, 5, 'check')
    K.cell(18, 6, '=SUM(F7:F16)').number_format = '0'
    K.conditional_formatting.add('F7:F18', CellIsRule(operator='notEqual', formula=['0'], fill=RED))
    K['D20'] = ('Stress test (Section 13.10 step 5): insert a column on Inputs and a row on Calc, change Scenario, and confirm F18 stays 0; '
                'then replace Calc row 26 with =OFFSET(Inputs!$E$10,0,J$4) and insert a column at Inputs F: F10 fires (1).')
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)


# =============================================================================================
# Inherited sheet (Exercise 13.13): exactly as specified, errors kept
# =============================================================================================
def inherited(path):
    wb = Workbook(); ws = wb.active; ws.title = 'Sheet1'
    ws['A1'] = 'Quebracho Alto: first three operating years (inherited sheet, Exercise 13.13)'; ws['A1'].font = BOLD
    for col, w in zip('ABCDEFGHI', [2, 2, 2, 30, 10, 12, 4, 4, 4]):
        ws.column_dimensions[col].width = w
    for col in 'JKL':
        ws.column_dimensions[col].width = 12
    ws['D4'] = 'Operating year'
    for i in range(3):
        ws.cell(4, J + i, i + 1)
    rows = [(8, 'Tariff', 'USD/MWh', 63.40, '0.00'), (9, 'Escalation', '% pa', 0.024, '0.0%'), (10, 'Energy', 'GWh', 551.4, '0.00'),
            (11, 'Opex', 'USD m', 7.94, '0.00'), (12, 'Opex escalation', '% pa', 0.02, '0.0%'), (13, 'Degradation', '% pa', 0.005, '0.0%')]
    for r, l_, u, v, f in rows:
        ws.cell(r, 4, l_); ws.cell(r, 5, u); c = ws.cell(r, 6, v); c.number_format = f
    calc = [(14, 'Tariff', 'USD/MWh', lambda c, p: f'=$F$8*(1+{L(6 + p)}9)^({c}$4-1)'),
            (15, 'Energy', 'GWh', lambda c, p: f'=$F$10*0.995^({c}$4-1)'),
            (16, 'Revenue', 'USD m', lambda c, p: f'={c}14*{c}15'),
            (17, 'Opex', 'USD m', lambda c, p: f'=OFFSET($E$11,0,1)*(1+$F$12)^({c}$4-1)'),
            (18, 'Margin', 'USD m', lambda c, p: f'=IFERROR({c}16-{c}17,0)')]
    for r, l_, u, fn in calc:
        ws.cell(r, 4, l_); ws.cell(r, 5, u)
        for i in range(3):
            c = ws.cell(r, J + i, fn(L(J + i), i)); c.number_format = '#,##0.00'
    ws['L15'] = 'tbc'
    wb.save(path)


# =============================================================================================
# Llano Pardo projection (Exercise 13.15; u01 Section 1.A.3 to 1.A.7)
# =============================================================================================
LP = dict(energy=134724.0, deg=0.005, tariff=58.20, opex=1085.0, esc=0.02, dep=3136.0, tax=0.25, debt=48809.0, rate=0.054,
          y0=2018, ny=20, nrep=18)


def llano_python(k):
    E = []; bal = LP['debt']; out = []
    for t in range(LP['ny']):
        en = LP['energy'] * (1 - LP['deg']) ** t
        rev = en * LP['tariff'] / 1000
        opx = LP['opex'] * (1 + LP['esc']) ** t
        intr = bal * LP['rate'] if t < LP['nrep'] else 0.0
        tax = max(0.0, LP['tax'] * (rev - opx - LP['dep'] - intr))
        cf = rev - opx - tax
        ds = cf / k if t < LP['nrep'] else 0.0
        pr = ds - intr
        bal -= pr
        out.append(dict(year=LP['y0'] + t, energy=en, rev=rev, opex=opx, tax=tax, cfads=cf, int=intr, prin=pr, ds=ds, bal=bal))
    return out


def llano_k():
    lo, hi = 1.0, 2.0
    for _ in range(200):
        mid = (lo + hi) / 2
        if llano_python(mid)[LP['nrep'] - 1]['bal'] > 0:   # debt not repaid: debt service too low -> lower k
            hi = mid
        else:
            lo = mid
    return (lo + hi) / 2


def llano(path):
    k = llano_k()
    wb = Workbook(); wb.remove(wb.active)
    I = wb.create_sheet('Inputs'); C = wb.create_sheet('Calc'); K = wb.create_sheet('Checks')
    case = 'Llano Pardo Solar (fictional; Illustrative)'
    header(I, 'Inputs: Chapter 1 Exhibits 1.4 and 1.6 (u01 Section 1.A.3 to 1.A.6); USD thousand', case)
    header(C, 'Calc: annual projection 2018 to 2037 (columns J to AC); sculpted debt service = CFADS / k', case)
    header(K, 'Checks: 0 when passing', case)
    items = [(6, 'Check tolerance', 'USD k', 0.5, '0.0'),
             (7, 'P50 energy, 2018 (operating year 1)', 'MWh', LP['energy'], '#,##0'),
             (8, 'Degradation (compounding)', '% pa', LP['deg'], '0.0%'),
             (9, 'PPA tariff (flat nominal)', 'USD/MWh', LP['tariff'], '0.00'),
             (10, 'Operating costs, 2018', 'USD k', LP['opex'], '#,##0'),
             (11, 'Operating cost escalation', '% pa', LP['esc'], '0.0%'),
             (12, 'Tax depreciation (62,721 / 20)', 'USD k pa', LP['dep'], '#,##0'),
             (13, 'Corporate income tax', '%', LP['tax'], '0%'),
             (14, 'Senior debt at COD', 'USD k', LP['debt'], '#,##0'),
             (15, 'Interest rate (fixed all-in, on the opening balance)', '% pa', LP['rate'], '0.00%'),
             (16, 'First operating year', 'year', LP['y0'], '0'),
             (17, 'Last repayment year', 'year', 2035, '0'),
             (18, 'k: CFADS / debt service, pasted from Goal Seek (closing debt 2035 = 0 by changing k)', 'x', k, '0.00000')]
    for r, l_, u, v, f in items:
        lab(I, r, l_, u); inp(I, f'F{r}', v, f)
    n = LP['ny']
    for i in range(n):
        c = C.cell(5, J + i, LP['y0'] + i); c.font = BOLD; c.fill = HEAD
    lab(C, 6, 'Year counter', 'index'); row(C, 6, n, lambda i: f'={pv_(i)}6+1', '0')
    lab(C, 7, 'Year', 'year'); row(C, 7, n, lambda i: f'=Inputs!$F$16+{cm(i)}6-1', '0')
    lab(C, 8, 'Operating year', 'OY'); row(C, 8, n, lambda i: f'={cm(i)}7-Inputs!$F$16+1', '0')
    lab(C, 9, 'Flag_Repayment (2018 to 2035)', 'flag'); row(C, 9, n, lambda i: f'=IF({cm(i)}7<=Inputs!$F$17,1,0)', '0')
    lab(C, 10, 'Energy', 'MWh'); row(C, 10, n, lambda i: f'=Inputs!$F$7*(1-Inputs!$F$8)^({cm(i)}8-1)', '#,##0', total=True)
    lab(C, 11, 'Revenue', 'USD k'); row(C, 11, n, lambda i: f'={cm(i)}10*Inputs!$F$9/1000', '#,##0', total=True)
    lab(C, 12, 'Operating costs', 'USD k'); row(C, 12, n, lambda i: f'=Inputs!$F$10*(1+Inputs!$F$11)^({cm(i)}8-1)', '#,##0', total=True)
    lab(C, 13, 'Debt, opening', 'USD k'); row(C, 13, n, lambda i: f'=IF({cm(i)}8=1,Inputs!$F$14,{pv_(i)}21)', '#,##0')
    lab(C, 14, 'Interest', 'USD k'); row(C, 14, n, lambda i: f'={cm(i)}13*Inputs!$F$15', '#,##0', total=True)
    lab(C, 15, 'Taxable income', 'USD k'); row(C, 15, n, lambda i: f'={cm(i)}11-{cm(i)}12-Inputs!$F$12-{cm(i)}14', '#,##0', total=True)
    lab(C, 16, 'Tax (none negative in the base case; checked)', 'USD k'); row(C, 16, n, lambda i: f'=MAX(0,{cm(i)}15)*Inputs!$F$13', '#,##0', total=True)
    lab(C, 17, 'CFADS', 'USD k'); row(C, 17, n, lambda i: f'={cm(i)}11-{cm(i)}12-{cm(i)}16', '#,##0', total=True)
    lab(C, 18, 'Debt service = CFADS / k in repayment years', 'USD k'); row(C, 18, n, lambda i: f'={cm(i)}9*{cm(i)}17/Inputs!$F$18', '#,##0', total=True)
    lab(C, 19, 'Principal', 'USD k'); row(C, 19, n, lambda i: f'={cm(i)}18-{cm(i)}14', '#,##0', total=True)
    lab(C, 20, 'DSCR', 'x'); row(C, 20, n, lambda i: f'=IF({cm(i)}18>0,{cm(i)}17/{cm(i)}18,0)', '0.00')
    lab(C, 21, 'Debt, closing', 'USD k'); row(C, 21, n, lambda i: f'={cm(i)}13-{cm(i)}19', '#,##0')
    lab(C, 22, 'Closing debt in the last repayment year', 'USD k'); frm(C, 'F22', f'=INDEX(J21:{cm(n - 1)}21,1,MATCH(Inputs!$F$17,J7:{cm(n - 1)}7,0))', '#,##0.000')
    C.freeze_panes = 'J6'
    tol = 'Inputs!$F$6'
    chk = [(7, 'Goal-seek k repays the debt by the last repayment year', f'=--(ABS(Calc!F22)>{tol})'),
           (8, 'Principal sums to the debt at COD', f'=--(ABS(Calc!G19-Inputs!$F$14)>{tol})'),
           (9, 'No negative taxable income (no losses to carry forward)', f'=--(MIN(Calc!J15:{cm(n - 1)}15)<0)'),
           (10, 'No debt after the last repayment year', f'=--(SUMPRODUCT((Calc!J7:{cm(n - 1)}7>Inputs!$F$17)*ABS(Calc!J21:{cm(n - 1)}21))>{tol})')]
    for r, l_, f in chk:
        K.cell(r, 4, l_); K.cell(r, 5, 'check'); K.cell(r, 6, f).number_format = '0'
    K.cell(12, 4, 'Master check (0 = all pass)').font = BOLD; K.cell(12, 6, '=SUM(F7:F10)')
    K.conditional_formatting.add('F7:F12', CellIsRule(operator='notEqual', formula=['0'], fill=RED))
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
    return k


if __name__ == '__main__':
    practice(os.path.join(HERE, 'Ch13_Practice_Solution.xlsx'))
    inherited(os.path.join(HERE, 'Ch13_Inherited_Sheet.xlsx'))
    k = llano(os.path.join(HERE, 'Ch13_LlanoPardo_Solution.xlsx'))
    print('written; Llano Pardo k =', k)
