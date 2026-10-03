#!/usr/bin/env python3
"""Verify every workbook in model/exercises/ against Python computations of the briefs' stated answers.

Recalculation: LibreOffice headless with the private profile -env:UserInstallation=file:///tmp/lo_profile_buildalong,
into this script's own output directory.  Writes <WORK>/verify_exercises.json (summarized in README.md).
Run after model/build/verify_stages.py (it reuses that run's recalculated Ch43 files for the 43.14 comparison).
"""
import os, sys, json, zipfile, datetime
import numpy as np
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
MODEL = os.path.dirname(HERE)
sys.path.insert(0, MODEL); sys.path.insert(0, os.path.join(MODEL, 'build'))
for d in ('ex43_17', 'ch13', 'ex85_12', 'ex86_13'):
    sys.path.insert(0, os.path.join(HERE, d))
import build_stages as bs            # noqa: E402
import verify_stages as vs           # noqa: E402
import build_excel_p as bx           # noqa: E402
import case_p as cp                  # noqa: E402
import ex43_17_mirror as mir         # noqa: E402
import build_ch13 as b13             # noqa: E402
import build_ex85_12 as b85          # noqa: E402
import build_ex86_13 as b86          # noqa: E402

WORK = os.path.join(bs.WORK, 'exercises'); IN = os.path.join(WORK, 'in'); OUT = os.path.join(WORK, 'out')
os.makedirs(IN, exist_ok=True); os.makedirs(OUT, exist_ok=True)
R = {}


def macro_free(path):
    z = zipfile.ZipFile(path); names = z.namelist(); wbx = z.read('xl/workbook.xml').decode()
    return path.endswith('.xlsx') and not any('vbaProject' in n for n in names) and 'iterate="1"' not in wbx


def copy_with(src, name, edits):
    wb = load_workbook(src)
    for (sh, ref), v in edits.items():
        wb[sh][ref].value = v
    p = os.path.join(IN, name); wb.save(p); return p


def ok(x, target, dp):
    return round(x, dp) == round(target, dp)


def main():
    todo = []
    # ---- Exercise 43.14
    start = os.path.join(HERE, 'ex43_14', 'Ch43_start_no_mc.xlsx')
    for s in range(1, 16):
        todo.append(copy_with(start, f'start_s{s:02d}.xlsx', {('Inputs', 'F8'): s}))
    ch43 = os.path.join(MODEL, 'build', 'Ch43_outputs.xlsx')
    for k in (1, 2, 1000):
        todo.append(copy_with(ch43, f'ch43_run{k:04d}.xlsx', {('Inputs', 'F311'): k}))
    # ---- Exercise 43.17
    sol = os.path.join(HERE, 'ex43_17', 'ex43_17_solution.xlsx')
    for s in (1, 2, 3):
        todo.append(copy_with(sol, f'ex43_17_s{s}.xlsx', {('Inputs', 'F8'): s}))
    # ---- Chapter 13, 85.12, 86.13
    simple = {'Ch13_Practice_Solution': os.path.join(HERE, 'ch13', 'Ch13_Practice_Solution.xlsx'),
              'Ch13_Inherited_Sheet': os.path.join(HERE, 'ch13', 'Ch13_Inherited_Sheet.xlsx'),
              'Ch13_LlanoPardo_Solution': os.path.join(HERE, 'ch13', 'Ch13_LlanoPardo_Solution.xlsx'),
              'ex85_12_solution': os.path.join(HERE, 'ex85_12', 'ex85_12_solution.xlsx'),
              'ex86_13_solution': os.path.join(HERE, 'ex86_13', 'ex86_13_solution.xlsx')}
    for n, p in simple.items():
        todo.append(copy_with(p, n + '.xlsx', {}))
    print('recalculating', len(todo)); sys.stdout.flush()
    bs.soffice_recalc(todo, OUT)
    ld = lambda n: load_workbook(os.path.join(OUT, n), data_only=True)

    # ======== 43.14 start file
    st = vs.structure(43, start)
    runs = {}
    ch43_out = os.path.join(bs.WORK, 'recalc_out')
    removed = lambda sh, r, c: sh == 'Inputs' and (311 <= r <= 1312 or r == 1321)
    for s in range(1, 16):
        a = ld(f'start_s{s:02d}.xlsx'); b = load_workbook(os.path.join(ch43_out, f'Ch43_s{s:02d}.xlsx'), data_only=True)
        n = 0; worst = 0.0; extra = []
        for sh in a.sheetnames:
            for row in a[sh].iter_rows():
                for c in row:
                    x = vs.num(c.value)
                    if x is None or (sh == 'Cover' and c.row >= bs.STAGE_NOTE_ROW):
                        continue
                    y = vs.num(b[sh][c.coordinate].value)
                    if y is None:
                        extra.append(f'{sh}!{c.coordinate}'); continue
                    n += 1; worst = max(worst, abs(x - y))
        ck = {r: a['Checks'].cell(r, 6).value for r in range(7, 26)}
        runs[s] = dict(cells=n, maxdiff=worst, master=a['Checks']['F19'].value, checks_zero=all((v or 0) == 0 for v in ck.values()), extra=extra[:3])
    a = ld('start_s01.xlsx'); ret = a['Returns']; rat = a['Ratios']; t1 = int(a['Time']['F13'].value)
    cum = [(ret.cell(7, 10 + i).value, ret.cell(11, 10 + i).value) for i in range(bx.SHEETS['Returns'].n)]
    payback = next(d for d, v in cum if v is not None and v >= 0)
    t4312 = dict(equity_irr=ret['F12'].value * 100, project_irr=ret['F13'].value * 100, project_irr_pretax=ret['F14'].value * 100,
                 npv16=ret['F15'].value, payback=payback.date().isoformat(), llcr=rat['F18'].value,
                 llcr_ex_dsra=rat.cell(14, 9 + t1).value, plcr=rat['F19'].value)
    t4312_ok = dict(equity_irr=ok(t4312['equity_irr'], 13.2, 1), project_irr=ok(t4312['project_irr'], 11.0, 1),
                    project_irr_pretax=ok(t4312['project_irr_pretax'], 11.6, 1), npv16=ok(t4312['npv16'], -45.3, 1),
                    payback=t4312['payback'] == '2031-06-30', llcr=ok(t4312['llcr'], 1.42, 2), 
                    plcr=ok(t4312['plcr'], 1.85, 2))   # v1.5 ledger P-F16, P-F41 (no ledger line for LLCR excluding the DSRA)
    # draw table
    DR = cp.mc_draws(); dw = load_workbook(os.path.join(HERE, 'ex43_14', 'ch43_draw_table.xlsx')).active
    dmax = max(abs(dw.cell(313 + r, 10 + j).value - DR[r, j]) for r in range(DR.shape[0]) for j in range(DR.shape[1]))
    c43 = load_workbook(ch43)['Inputs']
    dmax2 = max(abs(dw.cell(313 + r, 10 + j).value - c43.cell(313 + r, 10 + j).value) for r in range(DR.shape[0]) for j in range(DR.shape[1]))
    has_results = any(dw.cell(r, c).value is not None for r in range(311, 1313) for c in range(40, 45))
    # answer: runs 1, 2, 1000 in Ch43_outputs against its pasted per-run results; P-F42 statistics
    res = {}
    for k in (1, 2, 1000):
        w = ld(f'ch43_run{k:04d}.xlsx'); inp_ = w['Inputs']; wf = w['Waterfall']; dsr = w['Debt']
        ds = [vs.num(dsr.cell(114, 10 + i).value) or 0 for i in range(bx.NS)]
        hist = [vs.num(wf.cell(18, 10 + i).value) for i in range(bx.NS)]
        mh = min(h for h, d in zip(hist, ds) if d > 1e-9)
        live = [w['Ratios']['F16'].value, mh, w['Returns']['F12'].value]
        pasted = [inp_.cell(312 + k, 40 + j).value for j in range(3)]
        res[k] = dict(live=live, pasted=pasted, maxdiff=max(abs(a_ - b_) for a_, b_ in zip(live, pasted)),
                      checks=w['Checks']['F19'].value, active=inp_['F312'].value)
    pr = np.array([[c43.cell(313 + r, 40 + j).value for j in range(5)] for r in range(1000)], dtype=float)
    p_dscr = np.percentile(pr[:, 0], [10, 50, 90]); p_irr = np.percentile(pr[:, 2], [10, 50, 90]) * 100
    bins = [1.0, 1.1, 1.2, 1.25, 1.3, 1.35, 1.4, 1.5]
    hist_counts = [int(((pr[:, 0] >= bins[i]) & (pr[:, 0] < bins[i + 1])).sum()) for i in range(len(bins) - 1)] + [int((pr[:, 0] >= bins[-1]).sum())]
    pf42 = dict(min_dscr_p=[round(x, 4) for x in p_dscr], irr_p=[round(x, 3) for x in p_irr], lockup_share=float(pr[:, 3].mean()),
                default_share=float(pr[:, 4].mean()), histogram=hist_counts,
                ok=[round(x, 2) for x in p_dscr] == [1.32, 1.34, 1.35] and [round(x, 1) for x in p_irr] == [12.8, 13.1, 13.6] and pr[:, 3].mean() == 0)
    R['ex43_14'] = dict(structure=dict(macro_free=st['macro_free'], iterative=st['iterative'], forward_refs=st['forward_refs']),
                        scenarios=runs, ex43_12_targets=t4312, ex43_12_ok=t4312_ok, draw_table=dict(maxdiff_vs_mirror=dmax, maxdiff_vs_ch43=dmax2, has_results=has_results,
                                                                                                        macro_free=macro_free(os.path.join(HERE, 'ex43_14', 'ch43_draw_table.xlsx'))),
                        answer_runs=res, pf42=pf42)
    # ======== 43.17
    out, sol_, cases = mir.answers()
    e17 = {}
    for s, case in zip((1, 2, 3), ('base', 'sizing', 'downside')):
        w = ld(f'ex43_17_s{s}.xlsx'); r_ = w['Returns']; wf = w['Waterfall']
        cmp_rows = {}
        for key, rr in (('cfads', 7), ('ds', 8), ('dist', 17)):
            xs = [wf.cell(rr, 10 + i).value for i in range(40)]
            cmp_rows[key] = max(abs(a_ - b_) for a_, b_ in zip(xs, cases[case][key]))
        tx = [w['Tax'].cell(14, 10 + i).value for i in range(40)]
        cmp_rows['tax'] = max(abs(a_ - b_) for a_, b_ in zip(tx, cases[case]['tax']))
        e17[case] = dict(irr=r_['F10'].value, min=r_['F11'].value, avg=r_['F12'].value, lockups=r_['F13'].value,
                         master=w['Checks']['F20'].value, rows_maxdiff=cmp_rows,
                         scalars=dict(hard=r_['F14'].value, idc=r_['F15'].value, cfee=r_['F16'].value, upf=r_['F17'].value, dsra=r_['F18'].value,
                                      T=r_['F19'].value, D=r_['F20'].value, gearcap=r_['F21'].value, gearing=r_['F22'].value, equity=r_['F23'].value, ds1=r_['F24'].value))
    b = mir.BRIEF; sc = e17['base']['scalars']
    e17_ok = {k: ok(sc[k], b[k], len(str(b[k]).split('.')[1])) for k in ('hard', 'idc', 'cfee', 'upf', 'dsra', 'T', 'D', 'gearcap', 'gearing', 'equity', 'ds1')}
    e17_ok.update(irr_base=ok(e17['base']['irr'], b['irr_base'], 2), irr_sizing=ok(e17['sizing']['irr'], b['irr_sizing'], 2), irr_down=ok(e17['downside']['irr'], b['irr_down'], 2),
                  dscr_sizing=ok(e17['sizing']['min'], 1.30, 2) and ok(e17['sizing']['avg'], 1.30, 2),
                  dscr_base=ok(e17['base']['min'], 1.372, 3) and ok(e17['base']['avg'], 1.383, 3),
                  dscr_down=ok(e17['downside']['min'], 1.275, 3) and ok(e17['downside']['avg'], 1.279, 3),
                  no_lockups=all(e17[c]['lockups'] == 0 for c in e17), checks=all(e17[c]['master'] == 0 for c in e17))
    R['ex43_17'] = dict(cases=e17, ok=e17_ok, mirror_vs_brief=[list(x) for x in mir.compare(out)], macro_free=macro_free(sol))
    # ======== Chapter 13 practice
    w = ld('Ch13_Practice_Solution.xlsx'); T = w['Time']; C = w['Calc']; K = w['Checks']
    grid = [[C.cell(r, c).value for c in range(6, 9)] for r in range(71, 74)]
    exp_grid = [[-8.10, 13.96, 36.02], [-25.17, -4.67, 15.82], [-40.38, -21.28, -2.18]]
    pr13 = dict(flags=T['G12'].value, ops_fraction=T['F25'].value, ld_total=C['G22'].value, ld_months=[C.cell(22, 10 + i).value for i in range(8)],
                rev_y1=C['J27'].value, rev_y2=C['K27'].value, rev_total=C['G27'].value, rev_pv=C['F29'].value, npv_62_875=C['F37'].value,
                grid=grid, breakevens=[C['F51'].value, C['F52'].value, C['F53'].value], annuity=C['F54'].value, debt=C['F40'].value, idc=C['F55'].value,
                total=C['F56'].value, equity=C['F57'].value, iterations=[C.cell(r, 6).value for r in range(44, 50)], gs_npv=C['F50'].value,
                checks=[K.cell(r, 6).value for r in range(7, 17)], master=K['F18'].value, datatable_cells=[C.cell(65, c).value for c in range(6, 9)])
    pr13_ok = dict(flags=pr13['flags'] == 26, ops_fraction=ok(pr13['ops_fraction'], 0.3370, 4), ld_total=ok(pr13['ld_total'], 9620000, 0),
                   ld_months=[round(x) for x in pr13['ld_months']] == [1224500, 1185000, 1224500, 1224500, 1185000, 1224500, 1185000, 1167000],
                   rev=ok(pr13['rev_y1'], 34.96, 2) and ok(pr13['rev_y2'], 35.80, 2) and ok(pr13['rev_total'], 884.08, 2) and ok(pr13['rev_pv'], 385.26, 2),
                   grid=all(ok(grid[i][j], exp_grid[i][j], 2) for i in range(3) for j in range(3)),
                   breakevens=[round(x, 2) for x in pr13['breakevens']] == [59.47, 62.91, 66.46], annuity=ok(pr13['annuity'], 9.29353, 5),
                   debt=ok(pr13['debt'], 191.01, 2) and ok(pr13['idc'], 16.69, 2) and ok(pr13['equity'], 74.28, 2) and ok(pr13['total'], 265.29, 2),
                   iterations=[round(x, 4) for x in pr13['iterations'][1:]] == [190.2556, 190.9644, 191.0090, 191.0118, 191.0120],
                   checks=pr13['master'] == 0)
    R['ch13_practice'] = dict(values=pr13, ok=pr13_ok, macro_free=macro_free(simple['Ch13_Practice_Solution']))
    # ======== Chapter 13 inherited
    w = ld('Ch13_Inherited_Sheet.xlsx').active
    shown = {r: [w.cell(r, 10 + i).value for i in range(3)] for r in range(14, 19)}
    exp = {14: [63.40, 63.40, 63.40], 15: [551.40, 548.64, 'tbc'], 16: [34958.76, 34783.97, '#VALUE!'], 17: [7.94, 8.10, 8.26], 18: [34950.82, 34775.87, 0.00]}
    inh_ok = all((round(a_, 2) == b_) if isinstance(b_, float) else (a_ == b_) for r in exp for a_, b_ in zip(shown[r], exp[r]))
    tar = [63.40 * 1.024 ** k for k in range(3)]; en = [551.4 * 0.995 ** k for k in range(3)]
    rev = [t * e / 1000 for t, e in zip(tar, en)]; opx = [7.94 * 1.02 ** k for k in range(3)]; mar = [r_ - o for r_, o in zip(rev, opx)]
    corr = dict(tariff=[round(x, 2) for x in tar], energy=[round(x, 2) for x in en], revenue=[round(x, 2) for x in rev], opex=[round(x, 2) for x in opx], margin=[round(x, 2) for x in mar])
    corr_ok = corr == dict(tariff=[63.40, 64.92, 66.48], energy=[551.40, 548.64, 545.90], revenue=[34.96, 35.62, 36.29], opex=[7.94, 8.10, 8.26], margin=[27.02, 27.52, 28.03])
    R['ch13_inherited'] = dict(shown=shown, shown_ok=inh_ok, corrected_python=corr, corrected_ok=corr_ok, macro_free=macro_free(simple['Ch13_Inherited_Sheet']))
    # ======== Llano Pardo
    w = ld('Ch13_LlanoPardo_Solution.xlsx'); C = w['Calc']
    cf = [C.cell(17, 10 + i).value for i in range(20)]
    exh = [6510, 6436, 6361, 6286, 6209, 6131, 6052, 5972, 5890, 5808, 5723, 5638, 5551, 5462, 5372, 5281, 5187, 5092, 4995, 4945]
    lp = dict(k=w['Inputs']['F18'].value, cfads=[round(x) for x in cf], debt_2018=C['J21'].value, debt_2025=C['Q21'].value, debt_2035=C['F22'].value,
              dscr=[round(C.cell(20, 10 + i).value, 4) for i in range(18)], master=w['Checks']['F12'].value)
    lp_ok = dict(k=ok(lp['k'], 1.38074, 5), cfads=lp['cfads'] == exh, debt=round(lp['debt_2018']) == 46730 and round(lp['debt_2025']) == 30382 and abs(lp['debt_2035']) < 0.5,
                 dscr=all(round(x, 2) == 1.38 for x in lp['dscr']), checks=lp['master'] == 0)
    R['ch13_llano'] = dict(values=lp, ok=lp_ok, python_k=b13.llano_k(), macro_free=macro_free(simple['Ch13_LlanoPardo_Solution']))
    # ======== 85.12
    w = ld('ex85_12_solution.xlsx').active; py = b85.python_answers()
    v85 = dict(capacity=w['C16'].value, decline_default=w['G10'].value, decline_lockup=w['G15'].value, gearing=w['C34'].value, dscr_p90=w['C32'].value,
               aadt_default=w['G11'].value, aadt_lockup=w['G16'].value, verdicts=[w['C38'].value, w['G22'].value], master=w['C48'].value)
    R['ex85_12'] = dict(values=v85, python=py, ok=dict(capacity=ok(v85['capacity'], 68.53, 2), declines=ok(v85['decline_default'] * 100, 25.3, 1) and ok(v85['decline_lockup'] * 100, 14.0, 1),
                                                      aadt=round(v85['aadt_default']) == 23467 and round(v85['aadt_lockup']) == 27031,
                                                      python_equal=abs(v85['capacity'] - py['capacity']) < 1e-9, checks=v85['master'] == 0),
                        macro_free=macro_free(simple['ex85_12_solution']))
    # ======== 86.13
    w = ld('ex86_13_solution.xlsx').active
    order = [w.cell(r, 11).value for r in range(6, 10)]; chg = [w.cell(r, 13).value for r in range(6, 10)]
    py = b86.python_answers()
    R['ex86_13'] = dict(order=order, changes=chg, sortby_cells=[w.cell(6, c).value for c in range(6, 9)], master=w['C16'].value,
                        ok=dict(order=order == [x[0] for x in py], changes=all(abs(a_ - x[2]) < 1e-12 for a_, x in zip(chg, py)), checks=w['C16'].value == 0),
                        has_chart=len(load_workbook(simple['ex86_13_solution']).active._charts) == 1, macro_free=macro_free(simple['ex86_13_solution']))
    # ======== audit reader copy (modeler's file, R8)
    rd = os.path.join(HERE, 'Case_P_Model_AuditExercise_reader.xlsx'); full = os.path.join(MODEL, 'Case_P_Model_AuditExercise.xlsx')
    a, b = load_workbook(rd), load_workbook(full)
    same = all(a[sh][c.coordinate].value == c.value for sh in a.sheetnames for row in b[sh].iter_rows() for c in row if c.value is not None)
    R['audit_reader'] = dict(sheets=a.sheetnames, no_key='AuditKey' not in a.sheetnames, key_in_full='AuditKey' in b.sheetnames,
                             equal_to_full_minus_key=same and a.sheetnames == [s for s in b.sheetnames if s != 'AuditKey'], macro_free=macro_free(rd))
    json.dump(R, open(os.path.join(WORK, 'verify_exercises.json'), 'w'), indent=1, default=str)
    for k, v in R.items():
        print(k, {kk: vv for kk, vv in v.items() if kk in ('ok', 'ex43_12_ok', 'pf42', 'shown_ok', 'corrected_ok', 'no_key', 'equal_to_full_minus_key', 'macro_free')})


if __name__ == '__main__':
    main()
