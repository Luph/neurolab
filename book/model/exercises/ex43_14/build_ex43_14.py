#!/usr/bin/env python3
"""Build the Exercise 43.14 starting file Ex43_14_start.xlsx (u09 Chapter 43; D-125; coordinator note:
the starting file must not contain its answer).

Pre-run state: Ch43_outputs.xlsx with the Monte Carlo NOT implemented and NOT run.
  Supplied: the draw table (Inputs J312 heading, rows 313 to 1312, columns J to AL: 26 availability
  shocks by operating year, dispatch, heat-rate degradation, FX drift; seed 20180717, P-F42
  parameters) and its description cell Inputs F1321 (26 shock years).
  Removed (the reader builds or produces them): the run index Inputs F311 and the active flag F312;
  the per-run results pasted in Inputs AN311:AR1312 (the answer); the Monte Carlo branches of Time F16
  and Operations rows 31, 37 and 39, and the run-suspension conditions in Checks F8, F13, F14, F23 and
  Outputs F44 (these cells hold their pre-hook form, as in Ch42_waterfall.xlsx); the Cover's Monte
  Carlo line (row 29).
The answer (P-F42) is in model/build/Ch43_outputs.xlsx (Inputs AN313:AR1312) and the ledger.
"""
import os, sys
from openpyxl import load_workbook
from openpyxl.styles import Font

HERE = os.path.dirname(os.path.abspath(__file__))
BUILD = os.path.join(os.path.dirname(HERE), '..', 'build')
sys.path.insert(0, BUILD)
import build_stages as bs     # noqa: E402

OUT = os.path.join(HERE, 'Ex43_14_start.xlsx')


def build():
    wb = load_workbook(os.path.join(BUILD, 'Ch43_outputs.xlsx'))
    pre = load_workbook(os.path.join(BUILD, 'Ch42_waterfall.xlsx'))
    ws = wb['Inputs']
    # run index and active flag (reader builds them)
    for r in (311, 312):
        for col in 'DEF':
            ws[f'{col}{r}'].value = None
    # pasted per-run results (the answer)
    for (r, c) in [k for k in ws._cells if 311 <= k[0] <= 1312 and k[1] >= 40]:   # AN = column 40
        del ws._cells[(r, c)]
    # pre-hook cells (D-125), copied from the Chapter 42 file
    for sh, r, kind in bs.V13_ROWS + [('Checks', 23, 'F')]:
        w, p = wb[sh], pre[sh]
        w.cell(r, 4).value = p.cell(r, 4).value
        cols = [6] if kind == 'F' else range(10, 10 + 57)
        for c in cols:
            w.cell(r, c).value = p.cell(r, c).value
    # Outputs F44: drop the Monte Carlo suspension wrapper
    o = wb['Outputs']['F44']
    pre_ = '=IF(Inputs!$F$312=1,0,'
    assert o.value.startswith(pre_) and o.value.endswith(')')
    o.value = '=' + o.value[len(pre_):-1]
    # master check unchanged in form (it lists Checks F7 to F25)
    cv = wb['Cover']
    cv['D29'].value = ('Exercise 43.14 starting file: the Monte Carlo draw table is supplied on Inputs (rows 312 to 1312, columns J to AL; '
                       'seed 20180717, P-F42 parameters); the run index, the hooks into Time F16 and Operations rows 31, 37 and 39, '
                       'and the results are yours to build. The answer is P-F42 (Ch43_outputs.xlsx holds the solution).')
    note = ['Exercise 43.14 starting file (pre-run state). Built from model/build/Ch43_outputs.xlsx with the Monte Carlo removed.',
            'Supplied: the draw table (Inputs J312:AL1312) and Inputs F1321 (26 operating years with an availability shock).',
            'Removed: Inputs F311 (run index) and F312 (active flag); the pasted per-run results (Inputs AN:AR); the Monte Carlo branches.',
            'Every other cell equals Ch43_outputs.xlsx; all checks are 0 in every scenario.',
            'Verification: model/exercises/README.md.']
    for j, t in enumerate(note):
        c = cv.cell(bs.STAGE_NOTE_ROW + j, 4, t)
        c.font = Font(bold=(j == 0))
    wb.calculation.fullCalcOnLoad = True
    wb.save(OUT)
    return OUT


if __name__ == '__main__':
    print('written', build())
