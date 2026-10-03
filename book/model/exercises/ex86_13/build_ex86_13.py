#!/usr/bin/env python3
"""Exercise 86.13 (bible/briefs/u17.md): tornado chart from Example 86.3's sensitivity table (Tsiskari
hydro, Illustrative), bars ordered automatically by absolute impact.  Writes ex86_13_solution.xlsx
(macro-free).

Sheet Tornado:
  rows 6 to 9   inputs: sensitivity name and DSCR; base DSCR in C4; change and absolute change live;
  F6:H9         the ordering the exercise asks for: =SORTBY(B6:D9,E6:E9,-1) entered over F6:H9
                (SORTBY needs Excel for Microsoft 365, 2021 or 2024; t-excel-versions);
  J6:M9         the same ordering for any Excel version (LARGE / MATCH / INDEX), which the chart reads;
  the chart     horizontal bars of the change in DSCR from base, largest impact at the top
                (chart type per ssec:9.8.1, a tornado chart).
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.formatting.rule import CellIsRule
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.chart import BarChart, Reference

HERE = os.path.dirname(os.path.abspath(__file__))
BLUE = Font(color='0000FF'); BOLD = Font(bold=True)
YEL = PatternFill('solid', fgColor='FFF2CC'); RED = PatternFill('solid', fgColor='FF9999')
SENS = [('Generation -10%', 1.158), ('Operating cost +10%', 1.293), ('Availability -2 points (95.0% against 97.0%)', 1.310),
        ('Interest +100 bps on the 20% unhedged share', 1.332)]


def build(path):
    wb = Workbook(); ws = wb.active; ws.title = 'Tornado'
    for col, w in zip('ABCDEFGHIJKLM', [2, 44, 10, 12, 12, 44, 10, 12, 2, 4, 44, 12, 12]):
        ws.column_dimensions[col].width = w
    ws['A1'] = 'Example 86.3: bank-case DSCR sensitivities in order of impact (Illustrative)'; ws['A1'].font = Font(bold=True, size=12)
    ws['A2'] = 'Inputs blue on yellow. Debt service fixed at USD 13.105 million; base bank-case DSCR 1.35x.'
    ws['B4'] = 'Base DSCR'; c = ws['C4']; c.value = 1.35; c.font = BLUE; c.fill = YEL; c.number_format = '0.000'
    for col, h in zip('BCDE', ['Sensitivity', 'DSCR (x)', 'Change (x)', 'Abs change']):
        ws[f'{col}5'] = h; ws[f'{col}5'].font = BOLD
    for i, (n, v) in enumerate(SENS):
        r = 6 + i
        a = ws[f'B{r}']; a.value = n; a.font = BLUE; a.fill = YEL
        b = ws[f'C{r}']; b.value = v; b.font = BLUE; b.fill = YEL; b.number_format = '0.000'
        ws[f'D{r}'] = f'=C{r}-$C$4'; ws[f'D{r}'].number_format = '0.000'
        ws[f'E{r}'] = f'=ABS(D{r})'; ws[f'E{r}'].number_format = '0.000'
    # SORTBY ordering (as the exercise specifies)
    for col, h in zip('FGH', ['Sorted by impact (SORTBY)', 'DSCR (x)', 'Change (x)']):
        ws[f'{col}5'] = h; ws[f'{col}5'].font = BOLD
    ws['F6'] = ArrayFormula('F6:H9', '=_xlfn.SORTBY(B6:D9,E6:E9,-1)')
    for r in range(6, 10):
        for col in 'GH':
            ws[f'{col}{r}'].number_format = '0.000'
    # version-independent ordering (chart source)
    for col, h in zip('JKLM', ['Rank', 'Sensitivity (any Excel version)', 'DSCR (x)', 'Change (x)']):
        ws[f'{col}5'] = h; ws[f'{col}5'].font = BOLD
    for i in range(4):
        r = 6 + i
        ws[f'J{r}'] = f'={"J" + str(r - 1) + "+1" if i else "1"}'
        ws[f'K{r}'] = f'=INDEX($B$6:$B$9,MATCH(LARGE($E$6:$E$9,J{r}),$E$6:$E$9,0))'
        ws[f'L{r}'] = f'=INDEX($C$6:$C$9,MATCH(LARGE($E$6:$E$9,J{r}),$E$6:$E$9,0))'; ws[f'L{r}'].number_format = '0.000'
        ws[f'M{r}'] = f'=INDEX($D$6:$D$9,MATCH(LARGE($E$6:$E$9,J{r}),$E$6:$E$9,0))'; ws[f'M{r}'].number_format = '0.000'
    # checks
    ws['B12'] = 'Checks (0 = pass)'; ws['B12'].font = BOLD
    ws['B13'] = 'No ties in absolute change (ordering is unique)'; ws['C13'] = '=IF(SUMPRODUCT(1/COUNTIF(E6:E9,E6:E9))=COUNT(E6:E9),0,1)'
    ws['B14'] = 'Ordered list is in descending absolute change'; ws['C14'] = '=IF(AND(ABS(M6)>=ABS(M7),ABS(M7)>=ABS(M8),ABS(M8)>=ABS(M9)),0,1)'
    ws['B15'] = 'Ordered list keeps every sensitivity (sum of changes equal)'; ws['C15'] = '=IF(ABS(SUM(M6:M9)-SUM(D6:D9))<1E-9,0,1)'
    ws['B16'] = 'Master check'; ws['B16'].font = BOLD; ws['C16'] = '=SUM(C13:C15)'
    ws.conditional_formatting.add('C13:C16', CellIsRule(operator='notEqual', formula=['0'], fill=RED))
    # tornado chart
    ch = BarChart(); ch.type = 'bar'; ch.style = 10
    ch.title = 'Change in bank-case DSCR from 1.35x, by sensitivity'
    ch.y_axis.title = 'Change in DSCR (x)'
    ch.x_axis.scaling.orientation = 'maxMin'          # largest impact at the top
    data = Reference(ws, min_col=13, min_row=5, max_row=9)
    cats = Reference(ws, min_col=11, min_row=6, max_row=9)
    ch.add_data(data, titles_from_data=True); ch.set_categories(cats)
    ch.legend = None; ch.height = 8; ch.width = 18
    ws.add_chart(ch, 'B19')
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)


def python_answers():
    ch = sorted([(n, v, v - 1.35) for n, v in SENS], key=lambda x: -abs(x[2]))
    return ch


if __name__ == '__main__':
    build(os.path.join(HERE, 'ex86_13_solution.xlsx'))
    for x in python_answers():
        print(x)
