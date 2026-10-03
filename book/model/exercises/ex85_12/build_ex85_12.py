#!/usr/bin/env python3
"""Exercise 85.12 (bible/briefs/u17.md): one-tab screening sheet implementing Equations 85.1 and 85.2
with inputs, checks and a verdict cell, applied to Examples 85.1 (Écija solar, EUR) and 85.2 (Via Tâmega
Norte toll motorway, EUR).  Writes ex85_12_solution.xlsx (macro-free, live formulas).

Cell layout follows the brief's Excel forms:
  Equation 85.1 block, column C: C6 and C7 the two DSCR targets, C8 the rate, C9 the tenor, C10 the gearing
  cap, C11 total cost, C12 and C13 the two CFADS, C15 the uncapped capacity:
      C15 =MIN(C12/C6,C13/C7)*PV(C8,C9,-1)      C16 =MIN(C15,C10*C11)
  Equation 85.2 block, column G (same rows as the brief's column-C form, shifted to sit beside it):
  G3 revenue, G4 debt service, G5 DSCR test, G6 fixed opex, G7 variable cost share, G9 breakeven revenue:
      G9 =(G5*G4+G6)/(1-G7)      G10 =1-G9/G3
  (the lock-up test repeats the formula with the lock-up ratio in G13).
Rates and shares are Excel percentages (Chapter 85 uses Excel's own forms).
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.formatting.rule import CellIsRule

HERE = os.path.dirname(os.path.abspath(__file__))
BLUE = Font(color='0000FF'); BOLD = Font(bold=True)
YEL = PatternFill('solid', fgColor='FFF2CC'); RED = PatternFill('solid', fgColor='FF9999')


def build(path):
    wb = Workbook(); ws = wb.active; ws.title = 'Screen'
    for col, w in zip('ABCDEFGHI', [52, 12, 14, 4, 52, 12, 14, 4, 4]):
        ws.column_dimensions[col].width = w
    ws['A1'] = 'One-hour screen: debt capacity by annuity (Equation 85.1) and breakeven revenue (Equation 85.2). Illustrative.'
    ws['A1'].font = Font(bold=True, size=12)
    ws['A2'] = 'Inputs blue on yellow; checks 0 when passing. Left: Example 85.1 (EUR m). Right: Example 85.2 (EUR m).'
    def inp(ref, v, fmt):
        c = ws[ref]; c.value = v; c.font = BLUE; c.fill = YEL; c.number_format = fmt
    def put(lr, ur, vr, label, unit, val, fmt, is_inp=False):
        ws[lr] = label; ws[ur] = unit
        if is_inp:
            inp(vr, val, fmt)
        else:
            ws[vr] = val; ws[vr].number_format = fmt
    # ---------------- Equation 85.1 (Example 85.1), column C
    ws['A4'] = 'Equation 85.1: screening debt capacity (Example 85.1, Écija solar)'; ws['A4'].font = BOLD
    put('A6', 'B6', 'C6', 'DSCR target on P50', 'x', 1.30, '0.00', True)
    put('A7', 'B7', 'C7', 'DSCR target on one-year P99', 'x', 1.00, '0.00', True)
    put('A8', 'B8', 'C8', 'All-in cost of debt', '% pa', 0.0515, '0.00%', True)
    put('A9', 'B9', 'C9', 'Tenor', 'years', 11, '0', True)
    put('A10', 'B10', 'C10', 'Gearing cap', '%', 0.80, '0%', True)
    put('A11', 'B11', 'C11', 'Total project cost = capacity x cost per MWac', 'EUR m', '=C20*C25', '#,##0.00')
    put('A12', 'B12', 'C12', 'CFADS, P50 = revenue P50 - opex', 'EUR m', '=C27-C24', '#,##0.000')
    put('A13', 'B13', 'C13', 'CFADS, one-year P99 = revenue x P99 factor - opex', 'EUR m', '=C27*C23-C24', '#,##0.000')
    put('A14', 'B14', 'C14', 'Annuity factor PV(C8, C9, -1)', 'factor', '=PV(C8,C9,-1)', '0.0000')
    put('A15', 'B15', 'C15', 'Debt capacity before the gearing cap =MIN(C12/C6,C13/C7)*PV(C8,C9,-1)', 'EUR m', '=MIN(C12/C6,C13/C7)*PV(C8,C9,-1)', '#,##0.00')
    put('A16', 'B16', 'C16', 'Debt capacity =MIN(C15,C10*C11)', 'EUR m', '=MIN(C15,C10*C11)', '#,##0.00')
    ws['A18'] = 'Example 85.1 inputs'; ws['A18'].font = BOLD
    put('A19', 'B19', 'C19', 'Hours per year', 'h', 8760, '#,##0', True)
    put('A20', 'B20', 'C20', 'Capacity', 'MWac', 148.6, '0.0', True)
    put('A21', 'B21', 'C21', 'P50 net capacity factor (AC)', '%', 0.243, '0.0%', True)
    put('A22', 'B22', 'C22', 'One-year P90 as a share of P50', '%', 0.921, '0.0%', True)
    put('A23', 'B23', 'C23', 'One-year P99 as a share of P50', '%', 0.874, '0.0%', True)
    put('A24', 'B24', 'C24', 'Operating cost (flat for the screen)', 'EUR m pa', 2.38, '0.00', True)
    put('A25', 'B25', 'C25', 'Total project cost per MWac', 'EUR m/MWac', 0.612, '0.000', True)
    put('A26', 'B26', 'C26', 'PPA price (12 years, flat nominal)', 'EUR/MWh', 41.70, '0.00', True)
    put('A27', 'B27', 'C27', 'Revenue, P50 = MWac x hours x factor x price', 'EUR m', '=C20*C19*C21*C26/1000000', '#,##0.000')
    put('A28', 'B28', 'C28', 'P50 energy', 'MWh', '=C20*C19*C21', '#,##0.0')
    put('A29', 'B29', 'C29', 'CFADS, one-year P90', 'EUR m', '=C27*C22-C24', '#,##0.000')
    put('A30', 'B30', 'C30', 'Annual debt service at the chosen debt', 'EUR m', '=C16/C14', '#,##0.000')
    put('A31', 'B31', 'C31', 'DSCR at P50', 'x', '=C12/C30', '0.0000')
    put('A32', 'B32', 'C32', 'DSCR at one-year P90', 'x', '=C29/C30', '0.0000')
    put('A33', 'B33', 'C33', 'DSCR at one-year P99', 'x', '=C13/C30', '0.0000')
    put('A34', 'B34', 'C34', 'Gearing = debt / total cost', '%', '=C16/C11', '0.0%')
    put('A35', 'B35', 'C35', 'Capacity allowed by the P99 test alone', 'EUR m', '=C13/C7*C14', '#,##0.00')
    put('A36', 'B36', 'C36', 'Binding constraint', 'text', '=IF(C15>C10*C11,"gearing cap",IF(C12/C6<=C13/C7,"P50 DSCR test","P99 DSCR test"))', '@')
    ws['A38'] = 'Verdict (Equation 85.1)'; ws['A38'].font = BOLD
    ws['C38'] = ('=IF(AND(C31>=C6-$C$40,C33>=C7-$C$40,C34<=C10),"Screen passes: debt EUR "&TEXT(C16,"0.00")&" million, gearing "&TEXT(C34,"0.0%")&", "&C36&" binds",'
                 '"Screen fails: revisit the debt or the structure")')
    # ---------------- Equation 85.2 (Example 85.2), column G
    ws['E4'] = 'Equation 85.2: breakeven revenue for a DSCR test (Example 85.2, toll motorway, OY3 2030)'; ws['E4'].font = BOLD
    put('E3', 'F3', 'G3', 'Revenue in the test year = AADT x days x toll', 'EUR m', '=G18*G19*G20/1000000', '#,##0.000')
    ws['E3'].font = Font(italic=True)
    put('E5', 'F5', 'G5', 'DSCR test (default level)', 'x', 1.00, '0.00', True)
    put('E6', 'F6', 'G6', 'Fixed operating cost', 'EUR m pa', 9.40, '0.00', True)
    put('E7', 'F7', 'G7', 'Variable cost as a share of revenue', '%', 0.06, '0.0%', True)
    put('E8', 'F8', 'G8', 'CFADS = revenue x (1 - v) - fixed cost', 'EUR m', '=G3*(1-G7)-G6', '#,##0.000')
    put('E9', 'F9', 'G9', 'Breakeven revenue for the DSCR test =(G5*G4+G6)/(1-G7)', 'EUR m', '=(G5*G4+G6)/(1-G7)', '#,##0.000')
    put('E10', 'F10', 'G10', 'Revenue (traffic) decline to the DSCR test =1-G9/G3', '%', '=1-G9/G3', '0.0%')
    put('E11', 'F11', 'G11', 'AADT at the DSCR test', 'vehicles', '=G18*(1-G10)', '#,##0')
    put('E12', 'F12', 'G12', 'DSCR in the test year', 'x', '=G8/G4', '0.000')
    put('E13', 'F13', 'G13', 'Lock-up ratio', 'x', 1.20, '0.00', True)
    put('E14', 'F14', 'G14', 'Breakeven revenue for lock-up =(G13*G4+G6)/(1-G7)', 'EUR m', '=(G13*G4+G6)/(1-G7)', '#,##0.000')
    put('E15', 'F15', 'G15', 'Revenue (traffic) decline to lock-up =1-G14/G3', '%', '=1-G14/G3', '0.0%')
    put('E16', 'F16', 'G16', 'AADT at lock-up', 'vehicles', '=G18*(1-G15)', '#,##0')
    ws['E17'] = 'Example 85.2 inputs'; ws['E17'].font = BOLD
    put('E18', 'F18', 'G18', 'Forecast AADT', 'vehicles', 31420, '#,##0', True)
    put('E19', 'F19', 'G19', 'Days per year', 'days', 365, '0', True)
    put('E20', 'F20', 'G20', 'Average toll per vehicle', 'EUR', 4.85, '0.00', True)
    put('E21', 'F21', 'G21', 'Scheduled debt service', 'EUR m', 29.65, '0.00', True)
    ws['E4'].value = ws['E4'].value
    ws['G4'] = '=G21'; ws['G4'].number_format = '0.00'; ws['F4'] = 'EUR m'
    ws['E22'] = 'Verdict (Equation 85.2)'; ws['E22'].font = BOLD
    ws['G22'] = ('="Debt service is met down to a "&TEXT(G10,"0.0%")&" traffic fall; distributions lock up after a "&TEXT(G15,"0.0%")&" fall"')
    # ---------------- Checks
    ws['A40'] = 'Check tolerance'; ws['B40'] = 'x or EUR m'; inp('C40', 0.000001, '0.0E+00')
    ws['A41'] = 'Checks (0 = pass)'; ws['A41'].font = BOLD
    chk = [(42, 'Debt capacity not above the gearing cap', '=IF(C16<=C10*C11+$C$40,0,1)'),
           (43, 'P99 DSCR at the chosen debt not below its target', '=IF(C33>=C7-$C$40,0,1)'),
           (44, 'Annuity check: debt = PV of the annual debt service at the rate', '=IF(ABS(PV(C8,C9,-C30)-C16)<=$C$40,0,1)'),
           (45, 'Breakeven revenue reproduces the DSCR test', '=IF(ABS((G9*(1-G7)-G6)/G4-G5)<=$C$40,0,1)'),
           (46, 'Breakeven revenue reproduces the lock-up ratio', '=IF(ABS((G14*(1-G7)-G6)/G4-G13)<=$C$40,0,1)'),
           (47, 'Lock-up breakeven lies above the DSCR-test breakeven', '=IF(G14>G9,0,1)')]
    for r, l_, f in chk:
        ws[f'A{r}'] = l_; ws[f'B{r}'] = 'check'; ws[f'C{r}'] = f
    ws['A48'] = 'Master check'; ws['A48'].font = BOLD; ws['C48'] = '=SUM(C42:C47)'
    ws.conditional_formatting.add('C42:C48', CellIsRule(operator='notEqual', formula=['0'], fill=RED))
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)


EXPECT = dict(capacity=68.53, decline_default=0.253, decline_lockup=0.140)


def python_answers():
    en = 148.6 * 8760 * 0.243; rev = en * 41.70 / 1e6
    c50 = rev - 2.38; c99 = rev * 0.874 - 2.38
    af = (1 - 1.0515 ** -11) / 0.0515
    cap = min(min(c50 / 1.30, c99 / 1.00) * af, 0.80 * 148.6 * 0.612)
    r = 31420 * 365 * 4.85 / 1e6
    be1 = (1.00 * 29.65 + 9.40) / 0.94; be2 = (1.20 * 29.65 + 9.40) / 0.94
    return dict(capacity=cap, decline_default=1 - be1 / r, decline_lockup=1 - be2 / r, revenue=r, af=af,
                dscr_p90=(rev * 0.921 - 2.38) / (cap / af), gearing=cap / (148.6 * 0.612))


if __name__ == '__main__':
    build(os.path.join(HERE, 'ex85_12_solution.xlsx'))
    print(python_answers())
