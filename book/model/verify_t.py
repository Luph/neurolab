"""Verify Case_T_Model.xlsx against the Python mirror (case_t.py) in every scenario.

For each scenario the workbook is recalculated in headless LibreOffice (lo_recalc.py, repeated
hard recalculation with iterative calculation on), read back with openpyxl (data_only=True) and
compared, period by period and scalar by scalar, with the Python mirror. Writes
model/case_t_verification.md. Tolerance: 0.01 in displayed units (ARD m to one decimal, ratios
to two decimals, percentages to two decimals): absolute difference <= 0.01 for ARD m series and
<= 0.0001 for ratios, rates, factors and IRRs.
"""
import datetime as dt
import math
import os
import sys

import numpy as np
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import case_t as M  # noqa: E402
import build_excel_t as BX  # noqa: E402
from lo_recalc import recalc  # noqa: E402

SERIES = [  # (python key, excel row, unit class)
    ("traffic", "Traffic", "q"), ("toll", "Toll", "r"), ("wmult", "WMult", "r"), ("cpif", "CPIf", "r"),
    ("opsdays", "OpsDays", "m"), ("gross", "Gross", "m"), ("netrev", "NetRev", "m"), ("ld", "LD", "m"),
    ("opex", "Opex", "m"), ("revshare", "RevSh", "m"), ("restrcost", "RestrCost", "m"), ("ebitda", "EBITDA", "m"),
    ("lc", "LC", "m"), ("mmra_c", "MMRAc", "m"), ("hb_c", "HBc", "m"), ("dwc", "dWC", "m"), ("tax", "Tax", "m"),
    ("cfads", "CFADS", "m"), ("capex", "Capex", "m"), ("uses", "Uses", "m"), ("fees", "Fees", "m"),
    ("eqdraw", "EqDraw", "m"), ("bridgedraw", "BridgeDraw", "m"), ("sendraw", "SenDraw", "m"),
    ("nilodraw", "NiloDraw", "m"), ("bank_int", "BankInt", "m"), ("commit_fee", "Commit", "m"),
    ("bond_int", "BondInt", "m"), ("escrow_int", "EscrowInt", "m"), ("bridge_int", "BridgeInt", "m"),
    ("bank_close", "BankClose", "m"), ("bond_close", "BondClose", "m"), ("notes_close", "NotesClose", "m"),
    ("nilo_accr", "NiloAccr", "m"), ("nilo_cashdue", "NiloCashDue", "m"), ("nilo_Pdue", "NiloPDue", "m"),
    ("nilo_paid", "NiloPaid", "m"), ("nilo_close", "NiloClose", "m"), ("shl_int", "SHLInt", "m"),
    ("shl_close", "SHLClose", "m"), ("int_due", "IntDue", "m"), ("fees_due", "FeesDue", "m"), ("P_due", "PDue", "m"),
    ("sen_paid", "SenPaid", "m"), ("int_paid", "IntPaid", "m"), ("P_paid", "PPaid", "m"),
    ("arr_bank_close", "ArrBankClose", "m"), ("arr_bond_close", "ArrBondClose", "m"), ("ds_sched", "SchedDS", "m"),
    ("dscr", "DSCRp", "r"), ("dscr_hist", "DSCRH", "r"), ("dsra_target", "DSRATarget", "m"),
    ("dsra_draw", "DSRADraw", "m"), ("dsra_close", "DSRAClose", "m"), ("ria_rel", "RIA", "m"),
    ("ria_close", "RIAClose", "m"), ("mmra_close", "MMRAClose", "m"), ("hb_spend", "HBSpend", "m"),
    ("hb_close", "HBClose", "m"), ("support", "Support", "m"), ("A1", "A1", "m"), ("sweep1", "Sweep1", "m"),
    ("sweep2", "Sweep2", "m"), ("lockup", "Lock", "r"), ("eod", "EoD", "r"), ("distr", "Distr", "m"),
    ("cash_close", "CashClose", "m"), ("shl_paid", "SHLPaid", "m"), ("div", "Div", "m"), ("eq_cf", "EqCF", "m"),
    ("neweq_cf", "NewEqCF", "m"), ("amort", "Amort", "m"), ("ti", "TI", "m"), ("loss_close", "LossClose", "m"),
    ("wdv_close", "WDVClose", "m"), ("net_income", "NI", "m"), ("nbv", "NBV", "m"), ("total_assets", "TA", "m"),
    ("total_liab", "TL", "m"), ("total_equity", "TE", "m"), ("bs_check", "BSCheck", "m"),
    ("project_cf", "ProjCF", "m"), ("fv_cf", "FVCF", "m"), ("upgrade", "UpgCapex", "m"), ("llcr", "LLCR", "r"),
]


def tol_of(u):
    return 0.01 if u in ("m", "q") else 0.0001


def scalar_pairs(X, scn, d):
    st = X["st"]
    p = [("Senior debt", st["D"], "Dsen", "m"), ("NILO", st["N"], "Nsub", "m"), ("Equity", st["E"], "Eq", "m"),
         ("Total uses", X["total_uses"], "TotalUses", "m"), ("Funding net of contribution", X["Fnet"], "Fnet", "m"),
         ("Equity IRR", X["equity_irr"], "EqIRR", "r"), ("Equity NPV at 11.4%", X["equity_npv_target"], "EqNPV", "m"),
         ("Project IRR post-tax", X["project_irr"], "ProjIRR", "r"),
         ("Project IRR pre-tax", X["project_irr_pretax"], "ProjIRRPre", "r"),
         ("Min DSCR (repayment)", X["min_dscr_rep"], "MinDSCR", "r"), ("Avg DSCR", X["avg_dscr_rep"], "AvgDSCR", "r"),
         ("Min ramp-up DSCR", X["min_dscr_rampup"], "MinDSCRRamp", "r"),
         ("Min senior+NILO DSCR", X["min_comb_dscr"], "MinComb", "r"),
         ("LLCR first repayment", X["llcr_first"], "LLCR1", "r"), ("PLCR first repayment", X["plcr_first"], "PLCR1", "r"),
         ("Min LLCR", X["min_llcr"], "MinLLCR", "r"), ("Lock-up periods", X["lockup_periods"], "LockPeriods", "r"),
         ("DSRA draws", X["dsra_draws_total"], "DSRADraws", "m"),
         ("NPV distributions after 2022-06-30", X["npv_distr_after_2022H1"], "NPVDistr", "m"),
         ("Fair value 2022", X["fair_value_2022"], "FV2022", "m")]
    if X["P"]["live"]:
        p += [("Sculpting divisor s", st["s"], "Sdiv", "r"), ("NILO divisor k", st["k"], "Kdiv", "r"),
              ("Capacity: gearing", st["cand"]["gearing"], "D_Gear", "m"),
              ("Capacity: interest cover", st["cand"]["interest_cover"], "D_IO", "m"),
              ("Capacity: LLCR", st["cand"]["llcr"], "D_LLCR", "m"),
              ("Capacity: downside", st["cand"]["downside"], "D_Down", "m")]
    if scn == 5:
        r = d["restructuring"]
        p += [("Notes divisor", st["sn"], "SNdiv", "r"), ("NILO restructured divisor", st["kr"], "Kdiv", "r"),
              ("Claims net", r["claims_net"], "ClaimsNetV", "m"), ("Notes face", r["notes_issue"], "NotesFace", "m"),
              ("Swap MTM", r["swap_mtm"], "SwapMTMv", "m"), ("New equity value", r["equity_value_total"], "EqV", "m"),
              ("Warrants", r["warrant_value"], "WarrV", "m"), ("Notes market value", r["notes_market_value"], "NotesMV", "m"),
              ("Senior recovery", r["senior_recovery_pct_net_claims"], "SenRec", "r"),
              ("State implied grant", r["state_capital_grant_implied"], "StateGrant", "m"),
              ("NILO PV", r["nilo_pv_at_3.06"], "NiloPV", "m"), ("NILO recovery", r["nilo_recovery_pv_pct"], "NiloRec", "r"),
              ("Min notes DSCR", d["post_restructuring"]["min_notes_dscr"], "MinNotes", "r")]
    if scn == 1:
        ps = M.psc_vfm()
        p += [("PSC total", ps["psc_total"], "PSCTot", "m"), ("PPP reference", ps["ppp_reference_total"], "PPPRef", "m"),
              ("PPP bid", ps["ppp_bid_total"], "PPPBid", "m"), ("VfM reference", ps["vfm_reference"], "VfMRef", "m"),
              ("VfM bid", ps["vfm_bid"], "VfMBid", "m"), ("VfM reference %", ps["vfm_reference_pct"], "VfMRefPct", "r"),
              ("VfM bid %", ps["vfm_bid_pct"], "VfMBidPct", "r")]
    if scn == 4:
        t = d["termination_2022"]
        p += [("Senior claims 2022-06-30", t["senior_claims"], "SenClaim22", "m"),
              ("NILO 2022-06-30", t["nilo_outstanding"], "Nilo22", "m"),
              ("FM compensation", t["fm_total"], "FMComp", "m")]
    if scn == 6:
        p += [("Concessionaire-default compensation", X["fair_value_2022"] - M.RETENDER_COST, "CDComp", "m")]
    return p


def main():
    os.makedirs(os.path.join(HERE, "recalc"), exist_ok=True)
    B = BX.build(os.path.join(HERE, "Case_T_Model.xlsx"))
    runs, out = M.main()
    d = out["derived"]
    rows = B.rows
    report = []
    allpass = True
    summary = []
    cases = [(s, 1) for s in range(1, 14)] + [(2, 2)]
    for scn, co in cases:
        dst = os.path.join(HERE, "recalc", f"Case_T_s{scn}_c{co}.xlsx")
        passes = recalc(os.path.join(HERE, "Case_T_Model.xlsx"), dst, scn, co)
        wb = openpyxl.load_workbook(dst, data_only=True)
        if co == 2:
            X = M.run(2, out["contribution_solved"]["banking_case_at_11.4pct"], down_cfads=np.array(out["locked"]["down_cfads"]))
        else:
            X = runs[scn]
        res = []
        for key, rn, u in SERIES:
            row = rows[rn]
            ws = wb[row.sheet]
            xs = [ws.cell(row.r, BX.FIRST_COL + k).value for k in range(M.T)]
            py = X[key]
            diffs = []
            for a, b in zip(py, xs):
                if isinstance(a, float) and math.isnan(a):
                    if b not in ("", None):
                        diffs.append(float("inf"))
                    continue
                if b in ("", None):
                    b = 0.0
                if isinstance(b, str):
                    diffs.append(float("inf"))
                    continue
                diffs.append(abs(float(a) - float(b)))
            md = max(diffs) if diffs else 0.0
            ok = md <= tol_of(u)
            res.append((f"{row.sheet}!{rn} (row {row.r}, 97 periods)", md, tol_of(u), ok))
        for lab, pv, rn, u in scalar_pairs(X, scn, d):
            row = rows[rn]
            xv = wb[row.sheet].cell(row.r, 6).value
            if pv is None:
                ok = xv in ("n/a", None, 0)
                md = 0.0 if ok else float("inf")
            elif isinstance(xv, str):
                ok, md = False, float("inf")
            else:
                md = abs(float(pv) - float(xv or 0.0))
                ok = md <= tol_of(u)
            res.append((f"{lab} ({row.sheet}!F{row.r})", md, tol_of(u), ok))
        ck = wb["Checks"].cell(rows["CK_All"].r, 6).value
        res.append(("Checks!CK_All (all integrity checks)", abs(ck or 0.0), 0.0, abs(ck or 0.0) <= 1e-6))
        npass = sum(1 for r_ in res if r_[3])
        name = M.SCENARIOS[scn][0] + (" (contribution solved)" if co == 2 else "")
        summary.append((scn, co, name, passes, npass, len(res), max(r_[1] for r_ in res if math.isfinite(r_[1]))))
        allpass &= npass == len(res)
        report.append((scn, co, name, passes, res))
        print(scn, co, name, passes, f"{npass}/{len(res)}")
    write(report, summary, allpass)
    return allpass


def write(report, summary, allpass):
    L = ["# Case T verification: workbook against the Python mirror\n",
         f"Generated by `model/verify_t.py` on {dt.date.today().isoformat()} (model version 1.0). The workbook "
         "`model/Case_T_Model.xlsx` was recalculated in headless LibreOffice 24.2 through UNO (`model/lo_recalc.py`: "
         "iterative calculation on, repeated full recalculation until the Outputs and Checks cells change by less "
         "than 1e-9) once per scenario, read back with openpyxl (`data_only=True`) and compared with `model/case_t.py`.\n",
         "Tolerance: 0.01 in displayed units, applied as an absolute difference of at most 0.01 for ARD m and traffic "
         "series and at most 0.0001 for ratios, rates, factors, flags and IRRs (stricter than the displayed two "
         "decimals). Every series is compared in all 97 periods; the column 'Max abs diff' is the largest "
         "difference in any period.\n",
         f"**Overall result: {'PASS' if allpass else 'FAIL'}.**\n",
         "## Summary by scenario\n",
         "| Scenario | Contribution | Name | LibreOffice passes | Items passed | Max abs diff (finite) |",
         "|---|---|---|---|---|---|"]
    for scn, co, name, passes, npass, n, md in summary:
        L.append(f"| {scn} | {'bid 287.4' if co == 1 else 'solved'} | {name} | {passes} | {npass}/{n} | {md:.2e} |")
    for scn, co, name, passes, res in report:
        L.append(f"\n## Scenario {scn}: {name}{' (contribution solved)' if co == 2 else ''}\n")
        L.append("| Item | Max abs diff | Tolerance | Result |\n|---|---|---|---|")
        for lab, md, tol, ok in res:
            L.append(f"| {lab} | {md:.2e} | {tol:g} | {'pass' if ok else 'FAIL'} |")
    open(os.path.join(HERE, "case_t_verification.md"), "w").write("\n".join(L) + "\n")


if __name__ == "__main__":
    ok = main()
    print("ALL PASS" if ok else "FAILURES")
