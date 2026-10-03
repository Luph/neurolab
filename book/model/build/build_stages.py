#!/usr/bin/env python3
"""Build the cumulative build-along series (u09 Section 0.1 item 3, Section 0.5; D-051, D-124, D-125)

    model/build/Ch39_skeleton.xlsx   Ch40_funding.xlsx   Ch41_operations.xlsx
    model/build/Ch42_waterfall.xlsx  Ch43_outputs.xlsx

by FILTERING the output of the master build script (model/build_excel_p.py, Case P v1.4): no formula
is retyped.  Each stage keeps the sheets and rows that the u09 Section 0.4 row map assigns to that
chapter and every earlier chapter, with the modeler's confirmed assignments (case_p_report.md 8c):

  * Inputs row 226 (months per period) stays in Chapter 39;
  * provisional rows (D-051 (c)): Chapter 40 pastes Construction row 38; Chapter 41 pastes Debt rows
    103, 104, 110, 112 and Waterfall rows 33, 43; Chapters 39, 42 and 43 paste nothing.  A provisional
    row holds the FC base (Scenario 1) values of the recalculated master, as numbers in input color
    (blue on pale yellow), and its label carries the suffix
    "(provisional: pasted FC base values; replaced by formulas in Chapter N)";
  * D-125: the v1.4 Monte Carlo hooks (Time F16, Operations rows 31, 37 and 39, Checks F8, F13 and F14)
    are Chapter 43 changes.  Stages 39 to 42 carry these cells as they stood before the hooks
    (formulas and labels read from the v1.3 workbook in git, commit 8addab2); Chapter 43 replaces
    them with the v1.4 formulas.  The whole Monte Carlo block on Inputs (F311, F312, draw table and
    pasted per-run results in rows 311 to 1312, F1321) and the Cover's Monte Carlo line (row 29) are
    therefore Chapter 43 rows (8c had listed F311, F312 and F1321 under Chapter 39; under D-125 they
    would be inert and unreferenced there);
  * the same applies to the one v1.4 appended row of an earlier chapter that carries a hook: the
    ECA-test check Checks F23 (Chapter 42) reads AND(Scenario=1, Inputs F312=0) in v1.4; the
    Chapter 42 file holds it as Scenario=1 only, and Chapter 43 adds the Monte Carlo condition;
  * Checks F19 (master check) is restated in each stage as the sum of the absolute values of the
    checks present in that stage (in Chapter 43 it equals the master formula);
  * Cover rows 31 to 35 carry the stage note required by u09 R7 (stage, provisional rows, scenarios
    that reconcile).  This is the only intended difference between Ch43_outputs.xlsx and
    Case_P_Model.xlsx.

Usage: python3 build_stages.py  (writes the five files; needs LibreOffice for the master recalc that
supplies the provisional values).
"""
import os, sys, subprocess, shutil, json, copy
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter as L

HERE = os.path.dirname(os.path.abspath(__file__))
MODEL = os.path.dirname(HERE)
sys.path.insert(0, MODEL)
import build_excel_p as bx          # noqa: E402  (Case P v1.4 master build script)
import case_p as cp                 # noqa: E402

WORK = os.environ.get('BUILDALONG_WORK', '/tmp/claude-0/buildalong_work')
LOPROF = '-env:UserInstallation=file:///tmp/lo_profile_buildalong'
V13_XLSX = os.path.join(WORK, 'case_p_v13.xlsx')      # git show 8addab2:book/model/Case_P_Model.xlsx
BLUE = Font(color='0000FF'); YEL = PatternFill('solid', fgColor='FFF2CC')

STAGES = [(39, 'Ch39_skeleton.xlsx'), (40, 'Ch40_funding.xlsx'), (41, 'Ch41_operations.xlsx'),
          (42, 'Ch42_waterfall.xlsx'), (43, 'Ch43_outputs.xlsx')]

# ------------------------------------------------------------------------------------------------
# Row map: chapter that builds each row (u09 Section 0.4, case_p_report.md 8c, D-125)
# ------------------------------------------------------------------------------------------------
def _ranges(*spec):
    out = {}
    for ch, rows in spec:
        for r in rows:
            out[r] = ch
    return out

R = lambda a, b: range(a, b + 1)
INPUTS = _ranges(
    (39, R(7, 55)), (40, R(56, 69)), (41, R(70, 72)), (40, [73]), (40, R(74, 105)), (40, [107]),
    (42, [106]), (42, R(108, 111)), (40, R(112, 114)), (41, R(115, 183)), (42, [184]), (39, R(185, 208)),
    # structural constants 209 to 260
    (39, [209, 210, 211, 213, 214, 226]),
    (40, [212] + list(R(237, 245)) + [247, 248, 249, 251, 253, 254, 259]),
    (41, list(R(215, 225)) + list(R(227, 236)) + [246, 256, 258]),
    (42, [250, 257, 260]), (43, [252, 255]),
    (40, R(261, 280)), (40, R(281, 294)), (41, R(295, 297)), (40, R(298, 301)), (41, R(302, 303)),
    (40, R(304, 309)),
    (43, R(311, 1312)),                    # Monte Carlo run selector, flag, draw table, pasted results (D-125)
    (42, R(1314, 1320)), (43, [1321]),     # ECA limits and day_yr (Ch 42); mc_ny (Ch 43)
)
CONSTRUCTION = _ranges((39, R(7, 15)), (40, R(16, 34)), (41, R(35, 37)), (41, [38]))
OPERATIONS = _ranges((40, R(7, 15)), (41, R(16, 99)))
DEBT = _ranges((40, R(7, 19)), (42, R(20, 142)))
WATERFALL = _ranges((42, R(7, 44)), (41, [32]))
CHECKS = _ranges((40, [7, 8, 9, 16, 17]), (41, [15]), (42, [10, 11, 12, 13, 14, 18, 22, 23, 24]),
                 (39, [19, 20, 21]), (43, [25]))
COVER = _ranges((39, R(5, 28)), (43, [29]))
SHEET_CH = {'Time': 39, 'Tax': 41, 'Funding': 40, 'Reserves': 42, 'Financials': 42,
            'Ratios': 43, 'Returns': 43, 'Outputs': 43}
ROWMAP = {'Inputs': INPUTS, 'Construction': CONSTRUCTION, 'Operations': OPERATIONS, 'Debt': DEBT,
          'Waterfall': WATERFALL, 'Checks': CHECKS, 'Cover': COVER}
# provisional rows: (sheet, row) -> (chapter whose file first holds it pasted, chapter that writes the formula)
STUBS = {('Construction', 38): (40, 41),
         ('Debt', 103): (41, 42), ('Debt', 104): (41, 42), ('Debt', 110): (41, 42), ('Debt', 112): (41, 42),
         ('Waterfall', 33): (41, 42), ('Waterfall', 43): (41, 42)}
# D-125: cells whose v1.4 Monte Carlo hooks are Chapter 43 changes (pre-hook state read from v1.3)
V13_ROWS = [('Time', 16, 'F'), ('Operations', 31, 'row'), ('Operations', 37, 'row'), ('Operations', 39, 'row'),
            ('Checks', 8, 'F'), ('Checks', 13, 'F'), ('Checks', 14, 'F')]
STAGE_NOTE_ROW = 31


def row_chapter(sheet, r):
    """Chapter that builds (sheet, row) with its final formula; None = not part of the model."""
    if r <= 4:
        return 39                                  # header rows of every sheet
    if sheet in SHEET_CH:
        return SHEET_CH[sheet]
    return ROWMAP[sheet].get(r)


def present(sheet, r, stage):
    if (sheet, r) in STUBS:
        return STUBS[(sheet, r)][0] <= stage
    ch = row_chapter(sheet, r)
    return ch is not None and ch <= stage


def is_stub(sheet, r, stage):
    s = STUBS.get((sheet, r))
    return bool(s) and s[0] <= stage < s[1]


def sheets_in(stage):
    out = []
    for name in bx.ORDER:
        S = bx.SHEETS[name]
        rows = [r for _, _, r in S.rows] + ([22] if name == 'Cover' else [])
        if name == 'Cover' or any(present(name, r, stage) for r in rows if r >= 7):
            out.append(name)
    return out


# ------------------------------------------------------------------------------------------------
def soffice_recalc(paths, outdir):
    os.makedirs(outdir, exist_ok=True)
    subprocess.run(['soffice', LOPROF, '--headless', '--calc', '--convert-to', 'xlsx', '--outdir', outdir] + list(paths),
                   check=True, capture_output=True, timeout=3600)
    return [os.path.join(outdir, os.path.basename(p)) for p in paths]


def master_files():
    """Build the master (Scenario 1) with the v1.4 script and recalculate it once with LibreOffice."""
    os.makedirs(WORK, exist_ok=True)
    m = os.path.join(WORK, 'master_s1.xlsx')
    bx.build(m, 1)
    mv = os.path.join(WORK, 'master_out', 'master_s1.xlsx')
    if not os.path.exists(mv) or os.path.getmtime(mv) < os.path.getmtime(m):
        soffice_recalc([m], os.path.join(WORK, 'master_out'))
    if not os.path.exists(V13_XLSX):
        blob = subprocess.run(['git', '-C', MODEL, 'show', '8addab2:book/model/Case_P_Model.xlsx'],
                              check=True, capture_output=True).stdout
        open(V13_XLSX, 'wb').write(blob)
    return m, mv


def clear_row(ws, r):
    """Remove every cell of row r (value and style): the row is not built yet in this stage."""
    for (rr, cc) in [k for k in ws._cells if k[0] == r]:
        del ws._cells[(rr, cc)]


def stage_note(stage):
    prov = {40: 'Construction row 38 (VAT facility interest), replaced by its formula in Chapter 41.',
            41: 'Debt rows 103, 104, 110 and 112 and Waterfall rows 33 and 43, replaced by their formulas in Chapter 42.'}
    nxt = {39: 'Ch40_funding.xlsx', 40: 'Ch41_operations.xlsx', 41: 'Ch42_waterfall.xlsx', 42: 'Ch43_outputs.xlsx', 43: None}
    fname = dict(STAGES)[stage]
    lines = [
        f"Build-along file {fname}: the companion model at the end of Chapter {stage} (u09 Section 0.4 rows of Chapters 39 to {stage}; model version 1.4).",
        "Provisional rows: " + (prov[stage] + " They hold FC base (Scenario 1) values pasted in input color." if stage in prov else "none."),
        ("Reconciles on Scenario 1 only (the provisional rows hold Scenario 1 values); other scenarios reconcile from Ch42_waterfall.xlsx."
         if stage in prov else "Reconciles in every scenario (1 to 15)." if stage in (39, 42, 43) else ""),
        (f"Next file: {nxt[stage]}." if nxt[stage] else "Final stage: equal to Case_P_Model.xlsx except this note (rows 31 to 35)."),
        "Verification: model/build/verification.md.",
    ]
    return lines


def build_stage(stage, master_path, master_vals, v13):
    wb = load_workbook(master_path)
    keep = sheets_in(stage)
    for name in list(wb.sheetnames):
        if name not in keep:
            wb.remove(wb[name])
    for name in keep:
        ws = wb[name]
        maxr = ws.max_row
        for r in range(5, maxr + 1):
            if r == 5 and name != 'Cover':
                continue                       # period-label header row
            if not present(name, r, stage):
                clear_row(ws, r)
        # provisional rows
        for (sh, r), _ in STUBS.items():
            if sh != name or not is_stub(sh, r, stage):
                continue
            S = bx.SHEETS[sh]; mv = master_vals[sh]
            ws.cell(r, 4).value = ws.cell(r, 4).value + f" (provisional: pasted FC base values; replaced by formulas in Chapter {STUBS[(sh, r)][1]})"
            for i in range(S.n):
                c = ws.cell(r, bx.FC0 + i)
                v = mv.cell(r, bx.FC0 + i).value
                c.value = float(v) if v is not None else 0.0
                c.font = BLUE; c.fill = YEL
    # D-125: pre-hook cells in stages 39 to 42
    if stage < 43:
        for sh, r, kind in V13_ROWS:
            if sh not in keep or not present(sh, r, stage):
                continue
            ws = wb[sh]; wo = v13[sh]
            ws.cell(r, 4).value = wo.cell(r, 4).value
            cols = [6] if kind == 'F' else range(bx.FC0, bx.FC0 + bx.SHEETS[sh].n)
            for c in cols:
                ws.cell(r, c).value = wo.cell(r, c).value
    # D-125 applied to a v1.4 appended row: the ECA-test check (Checks F23, Chapter 42) carries the
    # Monte Carlo condition AND(Scenario=1, Inputs F312=0); before Chapter 43 it reads Scenario=1 only.
    if 42 <= stage < 43:
        c = wb['Checks']['F23']
        hook = 'AND(Inputs!$F$8=1,Inputs!$F$312=0)'
        assert hook in c.value, c.value
        c.value = c.value.replace(hook, 'Inputs!$F$8=1')
    # master check restated over the checks present
    ck = wb['Checks']
    present_checks = [r for r in list(range(7, 19)) + list(range(20, 26)) if present('Checks', r, stage)]
    ck['F19'].value = "=" + "+".join(f"ABS($F${r})" for r in present_checks) if present_checks else "=0"
    # stage note on the Cover (u09 R7)
    cv = wb['Cover']
    for j, t in enumerate(stage_note(stage)):
        c = cv.cell(STAGE_NOTE_ROW + j, 4, t)
        if j == 0:
            c.font = Font(bold=True)
    wb.calculation.fullCalcOnLoad = True
    out = os.path.join(HERE, dict(STAGES)[stage])
    wb.save(out)
    return out


def main():
    m, mv = master_files()
    master_vals = load_workbook(mv, data_only=True)
    v13 = load_workbook(V13_XLSX)
    # sanity: the master F19 restated over all checks must equal the master's own F19
    mf = load_workbook(m)
    allr = list(range(7, 19)) + list(range(20, 26))
    assert mf['Checks']['F19'].value == "=" + "+".join(f"ABS($F${r})" for r in allr), mf['Checks']['F19'].value
    outs = []
    for stage, fname in STAGES:
        outs.append(build_stage(stage, m, master_vals, v13))
        print('written', outs[-1])
    json.dump({'master': m, 'master_recalc': mv, 'stages': outs}, open(os.path.join(WORK, 'stages.json'), 'w'), indent=1)


if __name__ == '__main__':
    main()
