"""Verifies Case_R_Model.xlsx against the Python mirror.

For each scenario: build the workbook with that scenario's switches, recalculate it with LibreOffice
headless, read values (openpyxl data_only=True) and compare every mapped series cell and scalar with
outputs_case_r.json. Tolerance: 0.01 in displayed units (absolute). Writes case_r_verification.md.
"""
import json, os, subprocess, sys, shutil, tempfile
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import build_excel_r as BX
import case_r as M

TOL = 0.01
O = json.load(open(os.path.join(HERE, "outputs_case_r.json")))
COLS = BX.COLS


def lookup(key, sc):
    """Python value for a map key in scenario sc."""
    S = O["scenarios"][sc]
    if key.startswith("sz."):
        k = key[3:]
        if k.startswith("u_series_size.") or k.startswith("u_series_share."):
            a, b = k.split(".")
            return O["sizing"][a][b]
        return O["sizing"][k]
    if key.startswith("sc."):
        return S["scalars"][key[3:]]
    if key.startswith("val."):
        parts = key[4:].split(".")
        v = O["valuations"][sc] if sc in O["valuations"] else None
        if v is None:
            return None
        for p in parts:
            v = v[p]
        return v
    if key.startswith("su."):
        _, deal, k = key.split(".")
        return S["sources_uses"][deal][k]
    if key.startswith("refi."):
        return S["refinancing_cash"][key[5:]]
    if key.startswith("uri."):
        return O["uri_stress"][key[4:]]
    if key.startswith("div."):
        _, g, k = key.split(".")
        return O["diversification"][g][k]
    if key.startswith("ys."):
        _, a, k = key.split(".")
        return O["diversification"]["assets"][a][k]
    if key.startswith("dba."):
        _, f, a = key.split(".")
        return O["debt_by_asset"][f][a]
    return S["series"][key]


def num(v):
    if v is None or v == "":
        return None
    return float(v)


def compare(sc, wb, mp):
    rows = []
    worst = 0.0
    n = 0
    fails = 0
    for key, (sh, r) in mp["series"].items():
        if key.startswith("sz.") and sc != "base":
            continue
        py = lookup(key, sc)
        ws = wb[sh]
        maxd = 0.0
        for i, c in enumerate(COLS):
            xv = num(ws["%s%d" % (c, r)].value)
            pv = py[i] if isinstance(py, list) else None
            if pv is None and xv is None:
                continue
            if pv is None or xv is None:
                maxd = float("inf")
                continue
            maxd = max(maxd, abs(pv - xv))
            n += 1
        ok = maxd <= TOL
        fails += (not ok)
        worst = max(worst, maxd if maxd != float("inf") else worst)
        rows.append((key, sh, "row %d" % r, maxd, ok))
    for key, (sh, cell) in mp["scalars"].items():
        if (key.startswith("sz.") or key.startswith("uri.") or key.startswith("div.") or key.startswith("ys.") or key.startswith("dba.")) and sc != "base":
            continue
        py = lookup(key, sc)
        if py is None and key.startswith("val."):
            continue
        xv = num(wb[sh][cell].value)
        pv = None if py is None else float(py)
        if pv is None and xv is None:
            dd = 0.0
        elif pv is None or xv is None:
            dd = float("inf")
        else:
            dd = abs(pv - xv)
        n += 1
        ok = dd <= TOL
        fails += (not ok)
        worst = max(worst, dd if dd != float("inf") else worst)
        rows.append((key, sh, cell, dd, ok))
    return rows, n, fails, worst


def main():
    tmp = tempfile.mkdtemp(dir=os.environ.get("TMPDIR", None))
    summary = []
    detail = {}
    keyouts = {}
    for sc in M.SCENARIOS:
        src = os.path.join(tmp, "Case_R_%s.xlsx" % sc)
        BX.build(sc, src)
        mp = json.load(open(os.path.join(HERE, "case_r_excel_map.json")))
        outdir = os.path.join(tmp, "recalc_" + sc)
        subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", outdir, src], check=True, capture_output=True)
        wb = load_workbook(os.path.join(outdir, os.path.basename(src)), data_only=True)
        rows, n, fails, worst = compare(sc, wb, mp)
        chk = wb["Outputs"]
        summary.append((sc, n, fails, worst))
        detail[sc] = rows
        keyouts[sc] = rows
        print(sc, n, fails, worst)
    # rebuild the base workbook as the delivered file
    BX.build("base")
    os.makedirs(os.path.join(HERE, "recalc"), exist_ok=True)
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", os.path.join(HERE, "recalc"), os.path.join(HERE, "Case_R_Model.xlsx")], check=True, capture_output=True)
    L = ["# Case R verification: Excel workbook against the Python mirror\n",
         "Method: `python3 model/verify_case_r.py` builds `Case_R_Model.xlsx` once per scenario (switches set on the Inputs sheet), recalculates each copy with LibreOffice headless "
         "(`soffice --headless --convert-to xlsx`), reads the values with openpyxl (data_only=True) and compares every mapped cell with `outputs_case_r.json`. "
         "Series are compared in every one of the 38 annual columns. Tolerance: absolute difference of 0.01 in displayed units (USD m, GWh, USD/MWh, %, x, fractions). "
         "Sizing, Uri, yield-statistics and debt-by-asset items do not depend on the scenario switches and are compared in the base run only; valuations are compared in base, low and high. "
         "The delivered workbook is saved with the base scenario and refinancing on; its recalculated copy is `model/recalc/Case_R_Model.xlsx`.\n",
         "## Summary\n", "| Scenario | Values compared | Failures | Largest difference | Result |", "|---|---|---|---|---|"]
    allpass = True
    for sc, n, f, w in summary:
        L.append("| %s | %d | %d | %.2e | %s |" % (sc, n, f, w, "PASS" if f == 0 else "FAIL"))
        allpass &= (f == 0)
    L.append("\nOverall: **%s**.\n" % ("PASS" if allpass else "FAIL"))
    L.append("## Key outputs, base scenario (Python against Excel)\n")
    L.append("| Item | Python | Excel | Difference | Pass |")
    L.append("|---|---|---|---|---|")
    mp = json.load(open(os.path.join(HERE, "case_r_excel_map.json")))
    wb = load_workbook(os.path.join(HERE, "recalc", "Case_R_Model.xlsx"), data_only=True)
    keys = ["sz.tl_debt", "sz.hc_face", "sz.rf_debt", "sz.hi_face", "sz.u_size", "sz.u_series_size.A", "sz.u_series_size.B", "sz.u_series_size.C", "sz.u_coupon",
            "sz.hn_face", "refi.opco_net", "refi.mtm_opco_receivable", "refi.mtm_redfern_payable", "su.A1.equity", "su.A2.equity", "su.A3.equity",
            "sc.tl_dscr_min_2022_2025", "sc.uspp_dscr_min_2026_2043", "sc.uspp_dscr_avg_2026_2043", "sc.holdco_cov_min_2023_2031",
            "sc.fund_irr_life_pct", "sc.fund_irr_2025_pct", "sc.fund_nav_2025", "sc.fund_moic_life_x", "val.A1.ev", "val.A1.breakeven_price", "val.A2.ev", "val.A3.ev",
            "uri.net_cash", "div.A1.p90_1yr_gwh", "div.A1.p99_1yr_gwh", "div.all_generation.p90_10yr_gwh"]
    for k in keys:
        sh, cell = mp["scalars"][k]
        xv = num(wb[sh][cell].value)
        pv = float(lookup(k, "base"))
        L.append("| %s | %.4f | %.4f | %.1e | %s |" % (k, pv, xv, abs(pv - xv), "yes" if abs(pv - xv) <= TOL else "NO"))
    L.append("\n## Detail (all comparisons, maximum absolute difference across columns)\n")
    for sc, rows in detail.items():
        L.append("\n### %s\n" % sc)
        L.append("| Key | Sheet | Location | Max diff | Pass |")
        L.append("|---|---|---|---|---|")
        for key, sh, loc, dd, ok in rows:
            L.append("| %s | %s | %s | %.1e | %s |" % (key, sh, loc, dd, "yes" if ok else "NO"))
    open(os.path.join(HERE, "case_r_verification.md"), "w").write("\n".join(L) + "\n")
    shutil.rmtree(tmp, ignore_errors=True)
    return allpass


if __name__ == "__main__":
    sys.exit(0 if main() else 1)
