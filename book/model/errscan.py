import sys
from openpyxl import load_workbook
wb=load_workbook(sys.argv[1],data_only=True)
wf=load_workbook(sys.argv[1].replace('/out',''),data_only=False)
for ws in wb:
    seen=0
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value,str) and c.value.startswith(('#','Err')):
                print(ws.title,c.coordinate,c.value, ws.cell(c.row,4).value, str(wf[ws.title][c.coordinate].value)[:300]); seen+=1
                break
        if seen>=6: break
